import json
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from core.office.word_design import new_document, add_table, table_geometry, PRESETS
from core.office.word_audit import audit_word
from core.office.ppt_design import new_deck, add_slide, build_deck, layout_catalog
from core.office.ppt_template import inspect_template, follow_template, check_fidelity
from core.office.ppt_audit import audit_ppt
from core.office.excel_design import format_report, audit_report
from core.office.validate import validate


@pytest.mark.parametrize("preset", list(PRESETS))
def test_word_presets_have_real_styles_page_geometry_and_consistent_merged_tables(tmp_path, preset):
    path = tmp_path / "中文 报告.docx"
    doc = new_document("公司报告", preset=preset, subtitle="经营分析", metadata={"期间": "2026 Q3"})
    doc.add_heading("一、经营结论", 1)
    doc.add_paragraph("内容保留，并以语义样式排版。")
    table = add_table(doc, ["指标", "本期", "同比"], [["收入", 120, "18%"]], preset=preset)
    table.cell(1, 0).merge(table.cell(1, 1))
    table_geometry(table, [7, 5, 4])
    doc.save(path)
    read = Document(path)
    assert read.styles["Normal"].font.size.pt == PRESETS[preset]["body"]
    assert read.styles["Title"].element.find('.//' + qn('w:bottom')).get(qn('w:color')) == PRESETS[preset]["accent"]
    assert read.styles["Heading 1"].paragraph_format.keep_with_next
    assert read.sections[0].page_width.cm == pytest.approx(21, abs=.01)
    assert audit_word(path)["issues"] == []
    assert validate(path, {"contains": ["经营结论", "内容保留"]})["status"] == "passed"


def test_word_audit_finds_seeded_defects_without_mutating_file(tmp_path):
    path = tmp_path / "broken.docx"
    doc = new_document("故障夹具")
    doc.add_paragraph("一、假标题")
    doc.add_heading("跨级标题", 3)
    doc.add_paragraph("正文错误字号").runs[0].font.size = Pt(25)
    table = add_table(doc, ["A", "B"], [[1, 2]])
    table.cell(1, 0)._tc.tcPr.find(qn("w:tcW")).set(qn("w:w"), "99999")
    doc.save(path)
    before = path.read_bytes()
    report = audit_word(path)
    codes = {i["code"] for i in report["issues"]}
    assert {"WORD_POSSIBLE_FAKE_HEADING", "WORD_HEADING_LEVEL_GAP", "WORD_DIRECT_SIZE_OVERRIDE", "WORD_TABLE_CELL_MISMATCH"} <= codes
    assert path.read_bytes() == before


def slot_content(slots, image):
    values = {"table": [["指标", "数值"], ["收入", 120], ["利润", 25]],
              "chart": {"categories": ["六月", "七月", "八月"], "series": {"收入": [100, 120, 150]}},
              "image": str(image), "process": ["确认范围", "执行方案", "检查结果"],
              "list": ["核心结论", "数据分析", "行动计划"], "metric": "18%\n同比增长",
              "lead": "聚焦核心问题\n用证据支撑结论", "body": "经营持续改善\n保持投入节奏", "panel": "方案与证据\n目标明确，步骤可验证"}
    return {s["name"]: values[s["kind"]] for s in slots}


@pytest.mark.parametrize("layout", list(layout_catalog()))
def test_all_original_layouts_are_editable_and_inside_canvas(tmp_path, layout):
    image = tmp_path / "image.png"
    Image.new("RGB", (600, 400), "#24546A").save(image)
    prs = new_deck()
    add_slide(prs, layout, "经营分析与行动建议", slot_content(layout_catalog()[layout], image))
    path = tmp_path / f"{layout}.pptx"
    prs.save(path)
    reopened = Presentation(path)
    assert reopened.slides[0].shapes[0].text == "经营分析与行动建议"
    assert not [i for i in audit_ppt(path)["issues"] if i["severity"] == "error"]
    assert validate(path, {"slides": 1})["status"] == "passed"
    for shape in reopened.slides[0].shapes:
        if shape.has_chart:
            assert tuple(shape.chart.series[0].values) == (100, 120, 150)
        if shape.has_table:
            assert shape.table.cell(1, 0).text == "收入"


def template_fixture(tmp_path):
    prs = new_deck()
    slide = add_slide(prs, "conclusion", "公司模板", {"message": "核心结论", "evidence": "依据", "action": "下一步"})
    slide.notes_slide.notes_text_frame.text = "原讲稿必须保留"
    add_slide(prs, "table-wide", "经营表", {"table": [["指标", "本期"], ["收入", "100"]], "note": "单位：万元"})
    path = tmp_path / "品牌 模板.pptx"
    prs.save(path)
    return path


def test_template_mapping_retains_notes_masters_shapes_and_source(tmp_path):
    source = template_fixture(tmp_path)
    info = inspect_template(source)
    before = source.read_bytes()
    output = tmp_path / "填充.pptx"
    table_id = next(s["id"] for s in info["slides"][1]["shapes"] if s["kind"] == "table")
    table_info = next(s['table'] for s in info['slides'][1]['shapes'] if s['id'] == table_id)
    assert table_info == {'rows': 2, 'columns': 2, 'cells': [['指标', '本期'], ['收入', '100']], 'merged': False}
    report = follow_template(source, output, [{"slide": 1, "texts": {2: "第三季度经营汇报"}},
                                            {"slide": 2, "tables": {table_id: [["指标", "本期"], ["收入", "150"]]}}], source_sha256=info["sha256"])
    assert report["status"] == "passed"
    assert source.read_bytes() == before
    deck = Presentation(output)
    assert deck.slides[0].notes_slide.notes_text_frame.text == "原讲稿必须保留"
    assert deck.slides[0].shapes[0].text == "第三季度经营汇报"
    assert json.loads(output.with_suffix('.fidelity.json').read_text())["output_sha256"] == report["output_sha256"]
    deck.slides[0].shapes[1].left += Inches(.1)
    deck.save(output)
    assert check_fidelity(source, output, [(1, 2), (2, table_id)])["status"] == "issues"


@pytest.mark.parametrize("plan", [[{"slide": 99, "texts": {2: "x"}}], [{"slide": 1, "texts": {999: "x"}}],
                                  [{"slide": 1, "texts": {2: "too\nmany\nparagraphs"}}]])
def test_invalid_template_plan_does_not_write(tmp_path, plan):
    source = template_fixture(tmp_path)
    output = tmp_path / "result.pptx"
    with pytest.raises(ValueError):
        follow_template(source, output, plan)
    assert not output.exists()
    with pytest.raises(ValueError):
        follow_template(source, source, [{"slide": 1, "texts": {2: "x"}}])


def test_deck_plan_is_traceable_and_text_overflow_is_advisory(tmp_path):
    output = tmp_path / "plan.pptx"
    plan = {"audience": "管理层", "purpose": "决策", "message": "批准试点", "slides": [
        {"layout": "conclusion", "title": "批准小范围试点", "takeaway": "先验证回报再扩展", "notes": "请批准试点预算。",
         "content": {"message": "结论", "evidence": "过长内容" * 500, "action": "批准预算"}, "source": "用户提供的数据"}]}
    report = build_deck(plan, output)
    assert report["plan"] == plan
    assert any(i["code"] == "PPT_TEXT_FIT_ESTIMATE" for i in report["design_audit"]["issues"])
    assert Presentation(output).slides[0].notes_slide.notes_text_frame.text == "请批准试点预算。"


def test_excel_styling_preserves_values_formulas_ids_and_outside_region(tmp_path):
    book = Workbook()
    sheet = book.active
    for row in [["编号", "收入", "成本", "利润"], ["00123", 120, 70, "=B2-C2"], ["00456", 100, 65, "=B3-C3"]]:
        sheet.append(row)
    sheet["Z20"] = "不要修改"
    before = {c.coordinate: (c.value, c.number_format, c.style_id) for c in [sheet["Z20"]]}
    values = [[c.value for c in row] for row in sheet.iter_rows(min_row=1, max_row=3, max_col=4)]
    format_report(sheet, "A1:D3", columns={"A": "text", "B": "amount", "C": "amount", "D": "amount"},
                  roles={"B2:C3": "input", "D2:D3": "formula"}, financial=True)
    path = tmp_path / "report.xlsx"
    book.save(path)
    read = load_workbook(path)
    sheet = read.active
    assert [[c.value for c in row] for row in sheet.iter_rows(min_row=1, max_row=3, max_col=4)] == values
    assert {c.coordinate: (c.value, c.number_format, c.style_id) for c in [sheet["Z20"]]} == before
    assert sheet["A2"].value == "00123" and sheet["A2"].number_format == '@'
    assert sheet.freeze_panes == "A2"
    assert audit_report(sheet)["issues"] == []
    read.close()


def test_excel_bad_role_range_leaves_formatting_untouched():
    ws = Workbook().active
    ws.append(["指标", "数值"])
    ws.append(["收入", 100])
    before = ws["A1"].style_id
    with pytest.raises(ValueError):
        format_report(ws, "A1:B2", roles={"A1:A2": "input"})
    assert ws["A1"].style_id == before


def test_excel_long_chinese_notes_get_room_and_fixed_height_clipping_is_reported(tmp_path):
    ws = Workbook().active
    ws.append(["项目", "说明"])
    note = "所有数据均为虚构，仅用于 Excel 排版与公式验收。\n编号前导零必须保留。"
    ws.append(["数据性质", note])
    format_report(ws, "A1:B2")
    path = tmp_path / "notes.xlsx"
    ws.parent.save(path)
    book = load_workbook(path)
    read = book.active
    assert read["B2"].value == note
    assert read.row_dimensions[2].height >= 60
    assert audit_report(read)["issues"] == []
    read.row_dimensions[2].height = 25
    assert any(i["code"] == "EXCEL_TEXT_FIT_ESTIMATE" and i["cell"] == "B2" for i in audit_report(read)["issues"])
    book.close()


@pytest.mark.parametrize('area', ['D4:A1', 'A0:B2', 'XFE1:XFE2'])
def test_excel_rejects_reversed_or_out_of_sheet_range(area):
    ws = Workbook().active
    ws['A1'] = '保留'
    with pytest.raises(ValueError):
        format_report(ws, area)
    assert ws['A1'].style_id == 0 and not ws.print_area


def test_template_rejects_stale_source_and_detects_shape_reordering(tmp_path):
    source = template_fixture(tmp_path)
    output = tmp_path / "result.pptx"
    with pytest.raises(ValueError, match="模板已更新"):
        follow_template(source, output, [{"slide": 1, "texts": {2: "更新"}}], source_sha256="stale")
    assert not output.exists()
    prs = Presentation(source)
    tree = prs.slides[0].shapes._spTree
    tree.append(prs.slides[0].shapes[0]._element)
    prs.save(output)
    assert any(i['code'] == 'PPT_TEMPLATE_SHAPES_CHANGED' for i in check_fidelity(source, output)['issues'])


def test_template_accepts_blank_placeholder_without_losing_style(tmp_path):
    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    ident = slide.shapes.title.shape_id
    source, output = tmp_path / 'blank.pptx', tmp_path / 'filled.pptx'
    prs.save(source)
    assert follow_template(source, output, [{'slide': 1, 'texts': {ident: '中文标题'}}])['status'] == 'passed'
    assert Presentation(output).slides[0].shapes.title.text == '中文标题'


def test_visual_defaults_do_not_inherit_centered_panels_or_stretched_tables(tmp_path):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    prs = new_deck()
    slide = add_slide(prs, 'action-plan', '行动计划', {'actions': [['行动', '负责人'], ['试点', '运营']], 'decision': '请审议\n试点范围'})
    path = tmp_path / 'layout.pptx'
    prs.save(path)
    shapes = {s.name: s for s in Presentation(path).slides[0].shapes}
    assert shapes['actions'].height / Inches(1) == pytest.approx(1.3)
    assert shapes['decision'].text_frame.vertical_anchor == MSO_ANCHOR.TOP
    assert all(p.alignment == PP_ALIGN.LEFT for p in shapes['decision'].text_frame.paragraphs)
