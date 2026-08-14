#!/usr/bin/env python3
"""End-to-end dry run for the cloud approval workbench."""
from __future__ import annotations

import argparse
import json
import sys
import tempfile
import time
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any, Callable, Mapping
from uuid import uuid4

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from core.cloud_approval_client import CloudApprovalClient
from core.cloud_batch_sync import sync_local_approval_batch


ADMIN_TOKEN = "fake-admin-session"
MACHINE_ID = "dry-run-machine-1"
DEFAULT_BATCH_ID = "cloud-dry-run-batch"
MACHINE_TOKEN = "csr_machine_dry_run"
CAPABILITIES = ["generate_ai_image", "regenerate_ai_image", "submit_tmall_material_test", "crawl_tmall_material_test_data"]
PROMPT_HEADERS = {
    "字段名": "field_name",
    "字段 ID": "source_field_id",
    "字段ID": "source_field_id",
    "字段顺序": "field_order",
    "在当前视图": "visible",
    "尺寸": "size_label",
    "格式": "output_format",
    "引用字段": "reference_fields",
    "描述内容": "prompt_text",
    "字数": "word_count",
    "字段类型": "field_type",
    "女性优先度": "female_priority",
    "男性/中性优先度": "male_neutral_priority",
    "男性优先度": "male_neutral_priority",
    "中性优先度": "male_neutral_priority",
}
MATERIAL_OVERVIEW_REQUIRED = {"记录类型", "商品ID", "任务ID", "测试状态", "测试渠道", "测试素材数", "最优素材"}
MATERIAL_DETAIL_REQUIRED = {"记录类型", "商品ID", "任务ID", "统计口径", "统计日期", "图片类型", "素材ID", "素材URL"}
WORKFLOW_REQUIRED = {"款号", "ID（用于测图的ID）", "品类（后期匹配）", "模特性别"}


class FakeResponse:
    def __init__(self, status: int, body: Mapping[str, Any] | None = None):
        self.status = status
        self._body = json.dumps(dict(body or {}), ensure_ascii=False).encode("utf-8")

    def read(self) -> bytes:
        return self._body

    def getcode(self) -> int:
        return self.status


class CookieTransport:
    def __init__(self, cookie: str):
        self.cookie = str(cookie or "").strip()

    def __call__(self, request, timeout):
        if self.cookie:
            request.add_header("Cookie", self.cookie)
        return urllib.request.urlopen(request, timeout=timeout)


class FakeCloudApprovalTransport:
    """Deterministic in-memory transport that mirrors the worker API shape."""

    def __init__(self):
        self.calls: list[dict[str, Any]] = []
        self.uploads: list[dict[str, Any]] = []
        self.enrollment_tokens: dict[str, dict[str, Any]] = {}
        self.machines: dict[str, dict[str, Any]] = {}
        self.batches: dict[str, dict[str, Any]] = {}
        self.jobs: list[dict[str, Any]] = []
        self.prompt_libraries: list[dict[str, Any]] = []
        self.material_imports: list[dict[str, Any]] = []
        self.material_schedules: list[dict[str, Any]] = []
        self._next_token_id = 1
        self._next_style_id = 101
        self._next_job_id = 1

    def __call__(self, request, timeout):
        method = request.get_method()
        parsed = urllib.parse.urlparse(request.full_url)
        path = parsed.path
        body = _json_from_request(request)
        token_type = _token_type(request)
        self.calls.append({"method": method, "path": path, "body": body, "token_type": token_type})
        if method == "PUT" and path.startswith("/api/assets/upload/"):
            asset_uid = path.rsplit("/", 1)[-1]
            self.uploads.append({
                "asset_uid": asset_uid,
                "content_type": request.headers.get("Content-type") or request.headers.get("Content-Type") or "",
                "size": len(request.data or b""),
            })
            return FakeResponse(200, {"ok": True, "asset_uid": asset_uid})
        status, payload = self._handle_json(method, path, body, token_type)
        return FakeResponse(status, payload)

    def _handle_json(self, method: str, path: str, body: Mapping[str, Any], token_type: str) -> tuple[int, dict]:
        if method == "POST" and path == "/api/admin/machine-enrollment-tokens":
            token = f"csr_enroll_dry_run_{self._next_token_id}"
            row = {
                "id": self._next_token_id,
                "label": str(body.get("label") or "Dry run enrollment token"),
                "allowed_capabilities_json": json.dumps(list(body.get("allowed_capabilities") or [])),
                "require_approval": 0 if body.get("require_approval") is False else 1,
                "status": "issued",
                "used_by_machine_id": None,
            }
            self._next_token_id += 1
            self.enrollment_tokens[token] = row
            return 201, {"token": token, "enrollment_token": row}

        if method == "POST" and path == "/api/machines/enroll":
            token = str(body.get("enrollment_token") or "")
            row = self.enrollment_tokens.get(token)
            if not row or row["status"] != "issued":
                return 400, {"error": "Enrollment token is invalid"}
            allowed = set(json.loads(row["allowed_capabilities_json"]))
            capabilities = [str(item) for item in body.get("capabilities") or []]
            invalid = [item for item in capabilities if item not in allowed]
            if invalid:
                return 400, {"error": f"capability not allowed: {', '.join(invalid)}"}
            machine_id = str(body.get("machine_id") or MACHINE_ID)
            auth_status = "pending_approval" if row["require_approval"] else "active"
            row["status"] = "used"
            row["used_by_machine_id"] = machine_id
            self.machines[machine_id] = {
                "machine_id": machine_id,
                "machine_name": str(body.get("machine_name") or machine_id),
                "auth_status": auth_status,
                "health": "online_idle",
                "capabilities": capabilities,
                "machine_token": MACHINE_TOKEN,
            }
            return 201, {"machine_id": machine_id, "auth_status": auth_status, "machine_token": MACHINE_TOKEN}

        if method == "POST" and path == "/api/machines/heartbeat":
            machine = self._machine_for_token()
            if not machine:
                return 401, {"error": "Invalid machine token"}
            machine["health"] = str(body.get("health") or "online_idle")
            machine["capabilities"] = list(body.get("capabilities") or machine["capabilities"])
            return 200, {"ok": True, "machine_id": machine["machine_id"], "auth_status": machine["auth_status"], "health": machine["health"]}

        if method == "POST" and path == "/api/prompt-libraries/import":
            templates = list(body.get("templates") or [])
            if not templates:
                return 400, {"error": "templates are required"}
            library = {
                "id": len(self.prompt_libraries) + 1,
                "name": str(body.get("name") or "AI 测图提示词库 默认版"),
                "scenario": str(body.get("scenario") or "裂变图"),
                "status": "draft",
                "templates": templates,
            }
            self.prompt_libraries.append(library)
            return 201, {"library": library}

        if method == "POST" and path == "/api/material-test/import":
            overview_rows = list(body.get("overview_rows") or [])
            detail_rows = list(body.get("detail_rows") or [])
            source = dict(body.get("source") or {})
            source_uid = str(source.get("source_uid") or f"material-dry-run-{len(self.material_imports) + 1}")
            result = {
                "overview_rows": len(overview_rows),
                "detail_rows": len(detail_rows),
                "inserted_or_updated": len(overview_rows) + len(detail_rows),
                "source_uid": source_uid,
            }
            self.material_imports.append({**result, "source": source})
            return 200, result

        if method == "POST" and path == "/api/material-test/crawl-jobs":
            machine_id = str(body.get("machine_id") or "")
            machine = self.machines.get(machine_id) if machine_id else None
            if machine_id and (not machine or "crawl_tmall_material_test_data" not in machine["capabilities"]):
                return 400, {"error": "selected machine lacks crawl_tmall_material_test_data capability"}
            job = self._job(
                batch_uid="material-test",
                job_type="crawl_tmall_material_test_data",
                assigned_machine_id=machine_id or None,
                required_capabilities=["crawl_tmall_material_test_data"],
                priority=60,
                idempotency_key=str(body.get("idempotency_key") or f"crawl_tmall_material_test_data:{uuid4()}"),
                payload={"run_params": dict(body.get("run_params") or {}), "source": dict(body.get("source") or {})},
            )
            return 201, {"job": job}

        if method == "POST" and path == "/api/material-test/schedules":
            schedule = {
                "schedule_uid": str(body.get("schedule_uid") or f"mts_dry_run_{len(self.material_schedules) + 1}"),
                "label": str(body.get("label") or "天猫测图数据抓取"),
                "statistic_type": str(body.get("statistic_type") or "ACCUMULATE_30_DAYS"),
                "schedule_time": str(body.get("schedule_time") or "08:30"),
                "timezone": str(body.get("timezone") or "Asia/Shanghai"),
                "status": str(body.get("status") or "active"),
                "target_machine_id": body.get("machine_id"),
            }
            self.material_schedules.append(schedule)
            return 201, {"schedule": schedule}

        if method == "POST" and path == "/api/ai-image-batches/sync":
            batch_uid = str(body.get("batch_uid") or "")
            if not batch_uid:
                return 400, {"error": "batch_uid and title are required"}
            styles = []
            assets = {}
            for style in body.get("styles") or []:
                style_id = self._next_style_id
                self._next_style_id += 1
                style_copy = dict(style)
                style_copy["id"] = style_id
                style_copy["style_id"] = style_id
                styles.append(style_copy)
                for asset in style.get("assets") or []:
                    asset_copy = dict(asset)
                    asset_copy["style_id"] = style_id
                    asset_copy["status"] = str(asset_copy.get("status") or "uploaded")
                    assets[str(asset_copy["asset_uid"])] = asset_copy
            self.batches[batch_uid] = {"batch": dict(body), "status": "syncing", "styles": styles, "assets": assets}
            return 201, {"batch": {"batch_uid": batch_uid}, "styles": styles}

        if method == "POST" and path == "/api/assets/presign":
            asset_uid = str(body.get("asset_uid") or "")
            return 200, {
                "asset_uid": asset_uid,
                "object_key": f"batches/{body.get('batch_uid')}/{body.get('kind')}/{asset_uid}-{body.get('filename')}",
                "upload_url": f"/api/assets/upload/{asset_uid}",
                "method": "PUT",
                "headers": {},
            }

        if method == "POST" and path.endswith("/sync-complete"):
            batch_uid = _batch_uid(path, "/sync-complete")
            batch = self.batches.get(batch_uid)
            if not batch:
                return 404, {"error": "Not found"}
            if not any(asset.get("kind") == "ai" for asset in batch["assets"].values()):
                return 400, {"error": "sync-complete requires at least one style and one AI asset"}
            batch["status"] = "pending_review"
            return 200, {"ok": True, "status": "pending_review"}

        if method == "PATCH" and "/assets/" in path and path.endswith("/decision"):
            batch_uid, asset_uid = _decision_parts(path)
            batch = self.batches.get(batch_uid)
            if not batch or asset_uid not in batch["assets"]:
                return 404, {"error": "Not found"}
            decision = str(body.get("decision") or "")
            if decision not in {"approved", "rejected", "pending"}:
                return 400, {"error": "decision must be approved, rejected, or pending"}
            batch["assets"][asset_uid]["status"] = decision
            batch["status"] = "ready_to_submit" if _has_approved_ai(batch) else "pending_review"
            return 200, {"ok": True, "decision": decision, "batch_status": batch["status"]}

        if method == "POST" and path.endswith("/regenerate"):
            batch_uid = _batch_uid(path, "/regenerate")
            batch = self.batches.get(batch_uid)
            selected = [str(item) for item in body.get("asset_uids") or []]
            if not batch or not selected:
                return 400, {"error": "batch_uid and asset_uids are required"}
            jobs = []
            for asset_uid in selected:
                asset = batch["assets"].get(asset_uid)
                if not asset or asset.get("kind") != "ai":
                    return 400, {"error": f"asset is not in batch: {asset_uid}"}
                if asset.get("status") != "rejected":
                    return 409, {"error": "regeneration requires selected rejected assets"}
                jobs.append(self._job(
                    batch_uid=batch_uid,
                    job_type="regenerate_ai_image",
                    assigned_machine_id=None,
                    required_capabilities=["regenerate_ai_image"],
                    priority=50,
                    idempotency_key=f"regenerate_ai_image:{batch_uid}:{asset_uid}",
                    payload={
                        "batch_uid": batch_uid,
                        "style_id": asset["style_id"],
                        "asset_uid": asset_uid,
                        "prompt_text": asset.get("prompt_text") or asset.get("prompt") or "",
                        "reference_asset_uids": [
                            uid for uid, candidate in batch["assets"].items()
                            if candidate.get("style_id") == asset["style_id"] and candidate.get("kind") in {"source", "reference"}
                        ],
                        "parent_asset_uid": asset.get("parent_asset_uid"),
                    },
                ))
            return 201, {"jobs": jobs}

        if method == "POST" and path.endswith("/mark-ready"):
            batch_uid = _batch_uid(path, "/mark-ready")
            batch = self.batches.get(batch_uid)
            if not batch:
                return 404, {"error": "Not found"}
            if not _has_approved_ai(batch):
                return 409, {"error": "every non-skipped style must have at least one approved AI asset"}
            batch["status"] = "ready_to_submit"
            return 200, {"ok": True, "status": "ready_to_submit"}

        if method == "POST" and path.endswith("/submit"):
            batch_uid = _batch_uid(path, "/submit")
            machine_id = str(body.get("machine_id") or "")
            batch = self.batches.get(batch_uid)
            machine = self.machines.get(machine_id)
            if not batch or not machine:
                return 400, {"error": "batch_uid and machine_id are required"}
            if batch["status"] != "ready_to_submit":
                return 409, {"error": "submit requires batch status ready_to_submit"}
            if "submit_tmall_material_test" not in machine["capabilities"]:
                return 400, {"error": "selected machine lacks submit_tmall_material_test capability"}
            submit_plan = _submit_plan(batch_uid, batch)
            job = self._job(
                batch_uid=batch_uid,
                job_type="submit_tmall_material_test",
                assigned_machine_id=machine_id,
                required_capabilities=["submit_tmall_material_test"],
                priority=40,
                idempotency_key=f"submit_tmall_material_test:{batch_uid}:{machine_id}",
                payload={"submit_plan": submit_plan},
            )
            return 201, {"job": job, "submit_plan": submit_plan}

        if method == "POST" and path == "/api/machines/jobs/claim":
            machine = self._machine_for_token()
            if not machine:
                return 401, {"error": "Invalid machine token"}
            if machine.get("auth_status") != "active" or machine.get("health") not in {"online_idle", "online_busy"}:
                return 403, {"error": "Machine is not available for claims"}
            machine_capabilities = set(machine.get("capabilities") or [])
            claimable = [
                job for job in self.jobs
                if job.get("status") == "queued"
                and (not job.get("assigned_machine_id") or job.get("assigned_machine_id") == machine["machine_id"])
                and set(json.loads(job.get("required_capabilities_json") or "[]")).issubset(machine_capabilities)
            ]
            if not claimable:
                return 200, {"job": None, "next_poll_after_seconds": 10}
            selected = sorted(claimable, key=lambda job: (int(job.get("priority") or 100), int(job.get("id") or 0)))[0]
            lease_id = f"lease_dry_run_{selected['id']}"
            selected["status"] = "leased"
            selected["assigned_machine_id"] = machine["machine_id"]
            selected["attempt_count"] = int(selected.get("attempt_count") or 0) + 1
            selected["lease_id"] = lease_id
            selected["lease_expires_at"] = "2099-01-01T00:00:00Z"
            machine["health"] = "online_busy"
            machine["current_job_id"] = selected["job_uid"]
            return 200, {
                "job": {
                    "job_uid": selected["job_uid"],
                    "batch_uid": selected["batch_uid"],
                    "job_type": selected["job_type"],
                    "lease_id": lease_id,
                    "lease_expires_at": selected["lease_expires_at"],
                    "payload": selected["payload"],
                    "required_capabilities": json.loads(selected["required_capabilities_json"]),
                    "attempt_count": selected["attempt_count"],
                },
                "next_poll_after_seconds": 0,
            }

        if method == "POST" and path.startswith("/api/jobs/") and path.endswith("/complete"):
            machine = self._machine_for_token()
            if not machine:
                return 401, {"error": "Invalid machine token"}
            job_uid = path.removesuffix("/complete").rsplit("/", 1)[-1]
            lease_id = str(body.get("lease_id") or "")
            for job in self.jobs:
                if (
                    job["job_uid"] == job_uid
                    and job.get("lease_id") == lease_id
                    and job.get("assigned_machine_id") == machine["machine_id"]
                    and job.get("status") in {"leased", "running", "uploading_results"}
                ):
                    job["status"] = "succeeded"
                    job["result"] = dict(body.get("result") or {})
                    machine["health"] = "online_idle"
                    machine["current_job_id"] = None
                    return 200, {"ok": True, "status": "succeeded"}
            return 403, {"error": "Stale lease"}

        if method == "GET" and path.endswith("/submit-result"):
            batch_uid = _batch_uid(path, "/submit-result")
            jobs = [
                {
                    **job,
                    "required_capabilities": json.loads(job.get("required_capabilities_json") or "[]"),
                    "result": dict(job.get("result") or {}),
                }
                for job in self.jobs
                if job.get("batch_uid") == batch_uid and job.get("job_type") == "submit_tmall_material_test"
            ]
            return 200, {"jobs": jobs}

        return 404, {"error": f"Unhandled fake route: {method} {path}"}

    def _machine_for_token(self) -> dict[str, Any] | None:
        for machine in self.machines.values():
            if machine.get("machine_token") == MACHINE_TOKEN:
                return machine
        return None

    def _job(
        self,
        *,
        batch_uid: str,
        job_type: str,
        assigned_machine_id: str | None,
        required_capabilities: list[str],
        priority: int,
        idempotency_key: str,
        payload: Mapping[str, Any],
    ) -> dict:
        for job in self.jobs:
            if job["job_type"] == job_type and job["idempotency_key"] == idempotency_key:
                return job
        job = {
            "id": self._next_job_id,
            "job_uid": f"job_dry_run_{self._next_job_id}",
            "batch_uid": batch_uid,
            "job_type": job_type,
            "status": "queued",
            "assigned_machine_id": assigned_machine_id,
            "required_capabilities_json": json.dumps(required_capabilities),
            "priority": priority,
            "attempt_count": 0,
            "idempotency_key": idempotency_key,
            "lease_id": None,
            "lease_expires_at": None,
            "payload": dict(payload),
            "result": {},
        }
        self._next_job_id += 1
        self.jobs.append(job)
        return job


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run a cloud approval workbench dry run.")
    parser.add_argument(
        "--scenario",
        default="default",
        choices=["default", "embedded-login", "prompt-import", "material-data-import", "mvp-ai-test"],
        help="Dry-run scenario to execute. The default keeps the original fake approval loop.",
    )
    parser.add_argument("--prompt-file", default="", help="Workbook for --scenario prompt-import.")
    parser.add_argument("--data-file", default="", help="Workbook for --scenario material-data-import.")
    parser.add_argument("--workflow-file", default="", help="Workbook for --scenario mvp-ai-test.")
    parser.add_argument("--live-cloud-url", default="", help="Opt in to real cloud API calls against this Worker URL.")
    parser.add_argument("--admin-cookie", default="", help="Admin cs_session cookie for live admin routes. Never store this in files.")
    parser.add_argument("--machine-id", default="", help="Override task machine id; live mode defaults to a unique dry-run id.")
    parser.add_argument("--batch-id", default="", help="Override approval batch id; live mode defaults to a unique dry-run id.")
    return parser


def main(argv: list[str] | None = None) -> int:
    """Run a local fake cloud approval loop without Tmall side effects."""
    args = build_parser().parse_args(argv)
    if args.scenario == "embedded-login":
        return run_embedded_login_scenario(args, print)
    if args.scenario == "prompt-import":
        return run_prompt_import_scenario(Path(args.prompt_file), print)
    if args.scenario == "material-data-import":
        return run_material_data_import_scenario(Path(args.data_file), print)
    if args.scenario == "mvp-ai-test":
        return run_mvp_ai_test_scenario(Path(args.workflow_file), args, print)
    if args.live_cloud_url:
        transport = CookieTransport(args.admin_cookie)
        client = CloudApprovalClient(args.live_cloud_url, user_token=ADMIN_TOKEN, transport=transport)
        suffix = int(time.time())
        machine_id = args.machine_id or f"{MACHINE_ID}-{suffix}"
        batch_id = args.batch_id or f"{DEFAULT_BATCH_ID}-{suffix}"
        return run_dry_run(client=client, transport=None, batch_factory=build_fake_local_batch, printer=print, machine_id=machine_id, batch_id=batch_id)
    return run_dry_run(transport=FakeCloudApprovalTransport(), batch_factory=build_fake_local_batch, printer=print, machine_id=args.machine_id or MACHINE_ID, batch_id=args.batch_id or DEFAULT_BATCH_ID)


def run_embedded_login_scenario(args: argparse.Namespace, printer: Callable[[str], None]) -> int:
    app_source = REPO_ROOT / "cloud/approval-workbench/src/app/App.vue"
    text = app_source.read_text(encoding="utf-8")
    try:
        _assert("get('embed') === '1'" in text, "embed=1 query handling is missing")
        _assert(".app-shell.embedded" in text and ".embedded .workspace" in text, "embedded shell CSS is missing")
        _assert("v-if=\"!isEmbedded\"" in text, "embedded mode must hide standalone navigation")
        if args.live_cloud_url:
            if not args.admin_cookie:
                raise AssertionError("live embedded-login requires --admin-cookie from the current shell")
            url = f"{args.live_cloud_url.rstrip('/')}/?embed=1"
            request = urllib.request.Request(url)
            CookieTransport(args.admin_cookie)(request, timeout=20).read()
            printer(f"embedded-login live URL accepted admin cookie: {url}")
        else:
            printer("embedded-login fake check: embed=1 hides standalone chrome and uses the existing admin cookie session surface")
    except Exception as exc:
        printer(f"DRY RUN FAIL: {type(exc).__name__}: {exc}")
        return 1
    printer("DRY RUN PASS")
    return 0


def run_prompt_import_scenario(prompt_file: Path, printer: Callable[[str], None]) -> int:
    try:
        rows = parse_prompt_workbook(prompt_file)
        _assert(rows, "prompt workbook produced no templates")
        groups = sorted({row["group_name"] for row in rows})
        _assert(any(row.get("female_priority") is not None or row.get("male_neutral_priority") is not None for row in rows), "priority columns were not mapped")
        transport = FakeCloudApprovalTransport()
        client = CloudApprovalClient("https://approval.example.test", user_token=ADMIN_TOKEN, transport=transport)
        imported = client.request_json("POST", "/api/prompt-libraries/import", {
            "name": prompt_file.stem,
            "scenario": "裂变图",
            "templates": rows,
        }, token_type="user")
        library = imported.get("library") or {}
        _assert(len(library.get("templates") or []) == len(rows), "fake prompt import did not preserve every row")
        printer(f"prompt-import workbook={prompt_file} sheets={len(groups)} templates={len(rows)} groups={','.join(groups)}")
    except Exception as exc:
        printer(f"DRY RUN FAIL: {type(exc).__name__}: {exc}")
        return 1
    printer("DRY RUN PASS")
    return 0


def run_material_data_import_scenario(data_file: Path, printer: Callable[[str], None]) -> int:
    try:
        parsed = parse_material_workbook(data_file)
        _assert(parsed["overview_rows"], "material workbook produced no overview rows")
        _assert(parsed["detail_rows"], "material workbook produced no detail rows")
        transport = FakeCloudApprovalTransport()
        client = CloudApprovalClient("https://approval.example.test", user_token=ADMIN_TOKEN, transport=transport)
        imported = client.request_json("POST", "/api/material-test/import", {
            "source": {"filename": data_file.name, "source_uid": f"dry-run-{int(time.time())}"},
            "overview_rows": parsed["overview_rows"],
            "detail_rows": parsed["detail_rows"],
        }, token_type="user")
        _assert(imported.get("overview_rows") == len(parsed["overview_rows"]), "overview import count mismatch")
        _assert(imported.get("detail_rows") == len(parsed["detail_rows"]), "detail import count mismatch")
        printer(
            f"material-data-import workbook={data_file} overview_rows={len(parsed['overview_rows'])} "
            f"detail_rows={len(parsed['detail_rows'])} sheets={','.join(parsed['sheet_names'])}"
        )
    except Exception as exc:
        printer(f"DRY RUN FAIL: {type(exc).__name__}: {exc}")
        return 1
    printer("DRY RUN PASS")
    return 0


def run_mvp_ai_test_scenario(workflow_file: Path, args: argparse.Namespace, printer: Callable[[str], None]) -> int:
    try:
        rows = parse_workflow_workbook(workflow_file)
        _assert(rows, "workflow workbook produced no rows")
        transport = FakeCloudApprovalTransport()
        client = CloudApprovalClient("https://approval.example.test", user_token=ADMIN_TOKEN, transport=transport)
        batch_factory = lambda root: build_workflow_batch(root, rows)
        exit_code = run_dry_run(
            transport=transport,
            batch_factory=batch_factory,
            printer=printer,
            client=client,
            machine_id=args.machine_id or MACHINE_ID,
            batch_id=args.batch_id or DEFAULT_BATCH_ID,
        )
        if exit_code != 0:
            return exit_code
        crawl = client.request_json("POST", "/api/material-test/crawl-jobs", {
            "machine_id": args.machine_id or MACHINE_ID,
            "run_params": {"statistic_type": "ACCUMULATE_30_DAYS"},
            "source": {"workflow_file": str(workflow_file), "rows": len(rows)},
            "idempotency_key": f"crawl_tmall_material_test_data:{workflow_file.name}:{len(rows)}",
        }, token_type="user")
        job = crawl.get("job") or {}
        _assert(job.get("job_type") == "crawl_tmall_material_test_data", "crawl job type mismatch")
        _assert(job.get("required_capabilities_json") == json.dumps(["crawl_tmall_material_test_data"]), "crawl capability mismatch")
        printer(f"mvp-ai-test workflow={workflow_file} workflow_rows={len(rows)} crawl_job={job.get('job_uid')}")
    except Exception as exc:
        printer(f"DRY RUN FAIL: {type(exc).__name__}: {exc}")
        return 1
    return 0


def run_fake_dry_run(printer: Callable[[str], None] = print) -> dict:
    summary: dict[str, Any] = {}
    _run_flow(
        client=CloudApprovalClient("https://approval.example.test", user_token=ADMIN_TOKEN, transport=FakeCloudApprovalTransport()),
        batch_factory=build_fake_local_batch,
        printer=printer,
        summary=summary,
        machine_id=MACHINE_ID,
        batch_id=DEFAULT_BATCH_ID,
    )
    return summary


def run_dry_run(
    *,
    transport: FakeCloudApprovalTransport | None,
    batch_factory: Callable[[Path], dict],
    printer: Callable[[str], None],
    client: CloudApprovalClient | None = None,
    machine_id: str = MACHINE_ID,
    batch_id: str = DEFAULT_BATCH_ID,
) -> int:
    if client is None:
        client = CloudApprovalClient("https://approval.example.test", user_token=ADMIN_TOKEN, transport=transport)
    summary: dict[str, Any] = {}
    try:
        _run_flow(client=client, batch_factory=batch_factory, printer=printer, summary=summary, machine_id=machine_id, batch_id=batch_id)
    except AssertionError as exc:
        printer(f"DRY RUN FAIL: {exc}")
        return 1
    except Exception as exc:
        printer(f"DRY RUN FAIL: {type(exc).__name__}: {exc}")
        return 1
    printer("DRY RUN PASS")
    return 0


def _run_flow(
    *,
    client: CloudApprovalClient,
    batch_factory: Callable[[Path], dict],
    printer: Callable[[str], None],
    summary: dict[str, Any],
    machine_id: str,
    batch_id: str,
) -> None:
    printer("Phase 1: seed first admin")
    instructions = seed_admin_instructions()
    _assert(instructions.get("seed_admin"), "seed admin instructions are missing")
    summary["instructions"] = instructions

    printer("Phase 2: create machine enrollment token")
    enrollment = client.request_json("POST", "/api/admin/machine-enrollment-tokens", {
        "label": "Dry run task machine",
        "allowed_capabilities": CAPABILITIES,
        "require_approval": False,
        "expires_in_seconds": 3600,
    }, token_type="user")
    _assert(str(enrollment.get("token") or "").startswith("csr_enroll"), "enrollment token was not created")
    summary["enrollment_token"] = {"enrollment_token": dict(enrollment.get("enrollment_token") or {})}

    printer("Phase 3: enroll fake task machine")
    machine = client.request_json("POST", "/api/machines/enroll", {
        "enrollment_token": enrollment["token"],
        "machine_id": machine_id,
        "machine_name": "Dry Run Task Machine",
        "fingerprint": "dry-run-fingerprint",
        "app_version": "dry-run",
        "capabilities": CAPABILITIES,
    })
    _assert(machine.get("auth_status") == "active", "fake task machine did not become active")
    _assert(machine.get("machine_token"), "fake task machine did not receive a machine token")
    client.machine_token = str(machine["machine_token"])
    summary["machine"] = {key: value for key, value in machine.items() if key != "machine_token"}
    heartbeat = client.request_json("POST", "/api/machines/heartbeat", {
        "health": "online_idle",
        "capabilities": CAPABILITIES,
        "app_version": "dry-run",
    })
    _assert(heartbeat.get("health") == "online_idle", "fake task machine heartbeat failed")

    printer("Phase 4: sync fake local AI batch")
    with tempfile.TemporaryDirectory(prefix="crawshrimp-cloud-approval-dry-run-") as temp:
        batch = batch_factory(Path(temp))
        batch["batch_id"] = batch_id
        batch["title"] = f"Cloud Approval Dry Run Batch {batch_id}"
        sync_result = sync_local_approval_batch(batch, client)
    _assert(sync_result.get("status") == "pending_review", "fake local batch did not sync to pending_review")
    summary["sync"] = sync_result
    batch_uid = batch_id

    printer("Phase 5: reject image and create regeneration job")
    rejected = client.request_json("PATCH", f"/api/ai-image-batches/{batch_uid}/assets/ai-reject-1/decision", {
        "decision": "rejected",
        "note": "dry run rejection",
    }, token_type="user")
    _assert(rejected.get("decision") == "rejected", "rejected image decision was not saved")
    regen = client.request_json("POST", f"/api/ai-image-batches/{batch_uid}/regenerate", {
        "asset_uids": ["ai-reject-1"],
    }, token_type="user")
    regen_job = (regen.get("jobs") or [{}])[0]
    _assert(regen_job.get("job_type") == "regenerate_ai_image", "rejected image did not create regeneration job")
    _assert(regen_job.get("status") == "queued", "regeneration job is not queued")
    summary["regeneration_job"] = regen_job

    printer("Phase 6: approve image and create submit job")
    approved = client.request_json("PATCH", f"/api/ai-image-batches/{batch_uid}/assets/ai-approve-1/decision", {
        "decision": "approved",
        "note": "dry run approval",
    }, token_type="user")
    _assert(approved.get("decision") == "approved", "approved image decision was not saved")
    ready = client.request_json("POST", f"/api/ai-image-batches/{batch_uid}/mark-ready", {}, token_type="user")
    _assert(ready.get("status") == "ready_to_submit", "approved image did not make batch ready")
    submit = client.request_json("POST", f"/api/ai-image-batches/{batch_uid}/submit", {
        "machine_id": machine_id,
    }, token_type="user")
    submit_job = submit.get("job") or {}
    _assert(submit_job.get("job_type") == "submit_tmall_material_test", "approved image did not create submit job")
    _assert(submit_job.get("status") == "queued", "submit job is not queued")
    _assert(submit_job.get("assigned_machine_id") == machine_id, "submit job was not assigned to selected task machine")
    summary["submit_job"] = submit_job

    printer("Phase 7: task machine claims submit job")
    claim = client.request_json("POST", "/api/machines/jobs/claim", {})
    claimed_job = claim.get("job") or {}
    _assert(claimed_job.get("job_uid") == submit_job.get("job_uid"), "task machine did not claim the submit job")
    _assert(claimed_job.get("job_type") == "submit_tmall_material_test", "claimed job is not a submit job")
    _assert(claimed_job.get("lease_id"), "claimed submit job did not include a lease")
    summary["claimed_job"] = claimed_job

    printer("Phase 8: task machine completes submit job")
    completed = client.request_json("POST", f"/api/jobs/{claimed_job['job_uid']}/complete", {
        "lease_id": claimed_job["lease_id"],
        "result": {"dry_run": True, "batch_uid": batch_uid},
    })
    _assert(completed.get("status") == "succeeded", "submit job completion was not accepted")
    submit_job["status"] = "succeeded"
    summary["completed_job"] = {
        "job_uid": claimed_job["job_uid"],
        "status": completed.get("status"),
    }

    printer("Phase 9: verify submit result is visible to reviewers")
    submit_result = client.request_json("GET", f"/api/ai-image-batches/{batch_uid}/submit-result", token_type="user")
    submit_result_jobs = submit_result.get("jobs") or []
    _assert(any(job.get("job_uid") == submit_job.get("job_uid") and job.get("status") == "succeeded" for job in submit_result_jobs), "completed submit job is not visible in submit result")
    summary["submit_result_jobs"] = submit_result_jobs


def seed_admin_instructions() -> dict:
    return {
        "seed_admin": (
            "Create the first admin directly in D1 during deployment using "
            "cloud/approval-workbench/migrations/0001_init.sql roles plus the users, user_roles, and password_hash columns; "
            "after that, manage accounts only through /api/admin/users."
        ),
        "no_public_registration": "The workbench has /api/auth/login only; public self-registration is intentionally absent.",
    }


def build_fake_local_batch(root: Path) -> dict:
    root.mkdir(parents=True, exist_ok=True)
    source = root / "source-main-1.jpg"
    reject = root / "ai-reject-1.jpg"
    approve = root / "ai-approve-1.jpg"
    source.write_bytes(b"source-image")
    reject.write_bytes(b"rejected-ai-image")
    approve.write_bytes(b"approved-ai-image")
    return {
        "batch_id": DEFAULT_BATCH_ID,
        "title": "Cloud Approval Dry Run Batch",
        "status": "pending_approval",
        "created_at": "2026-07-07T10:00:00+08:00",
        "adapter_id": "tmall-ops-assistant",
        "task_id": "tmall_ai_image_test_chain",
        "task_run_uid": "dry-run-local-task",
        "items": [{
            "id": "style-dry-run-1",
            "row_no": 1,
            "style_code": "208326100202",
            "item_id": "1002178235142",
            "category": "长袖T恤",
            "gender": "中性",
            "skc_code": "208326100202-00482",
            "origin_path": str(source),
            "assets": [
                {"id": "source-main-1", "kind": "origin", "path": str(source), "filename": source.name, "label": "原图/主图"},
                {
                    "id": "ai-reject-1",
                    "kind": "ai",
                    "path": str(reject),
                    "filename": reject.name,
                    "label": "AI 图-退回",
                    "prompt": "保留主商品并更换童装场景",
                    "prompt_index": 1,
                    "generation_row": {"prompt_version": "dry-run-v1", "提示词版本": "Dry Run"},
                },
                {
                    "id": "ai-approve-1",
                    "kind": "ai",
                    "path": str(approve),
                    "filename": approve.name,
                    "label": "AI 图-通过",
                    "prompt": "保留主商品并突出面料细节",
                    "prompt_index": 2,
                    "generation_row": {"prompt_version": "dry-run-v1", "提示词版本": "Dry Run"},
                },
            ],
        }],
    }


def build_workflow_batch(root: Path, workflow_rows: list[dict[str, Any]]) -> dict:
    row = workflow_rows[0]
    root.mkdir(parents=True, exist_ok=True)
    source = root / f"{row['style_code']}-source.jpg"
    reject = root / f"{row['style_code']}-ai-reject.jpg"
    approve = root / f"{row['style_code']}-ai-approve.jpg"
    source.write_bytes(b"workflow-source-image")
    reject.write_bytes(b"workflow-rejected-ai-image")
    approve.write_bytes(b"workflow-approved-ai-image")
    return {
        "batch_id": DEFAULT_BATCH_ID,
        "title": "Cloud Approval MVP AI Test Dry Run",
        "status": "pending_approval",
        "created_at": "2026-07-08T10:00:00+08:00",
        "adapter_id": "tmall-ops-assistant",
        "task_id": "tmall_ai_image_test_chain",
        "task_run_uid": "mvp-ai-test-dry-run",
        "items": [{
            "id": f"style-{row['style_code']}",
            "row_no": row["row_no"],
            "style_code": row["style_code"],
            "item_id": row["item_id"],
            "category": row["category"],
            "gender": row["gender"],
            "skc_code": row["style_code"],
            "origin_path": str(source),
            "assets": [
                {"id": "source-main-1", "kind": "origin", "path": str(source), "filename": source.name, "label": "原图/主图"},
                {
                    "id": "ai-reject-1",
                    "kind": "ai",
                    "path": str(reject),
                    "filename": reject.name,
                    "label": "AI 图-退回",
                    "prompt": f"{row['category']} {row['gender']} 测图干跑重生图",
                    "prompt_index": 1,
                    "generation_row": {"prompt_version": "mvp-dry-run", "提示词版本": "MVP Dry Run"},
                },
                {
                    "id": "ai-approve-1",
                    "kind": "ai",
                    "path": str(approve),
                    "filename": approve.name,
                    "label": "AI 图-通过",
                    "prompt": f"{row['category']} {row['gender']} 测图干跑提交",
                    "prompt_index": 2,
                    "generation_row": {"prompt_version": "mvp-dry-run", "提示词版本": "MVP Dry Run"},
                },
            ],
        }],
    }


def parse_prompt_workbook(path: Path) -> list[dict[str, Any]]:
    _assert(path and path.exists(), f"prompt workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    templates: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header_index = next((index for index, row in enumerate(rows) if any(_cell_text(cell) == "字段名" for cell in row)), -1)
        if header_index < 0:
            continue
        headers = [_cell_text(cell) for cell in rows[header_index]]
        index_by_key = {PROMPT_HEADERS[header]: index for index, header in enumerate(headers) if header in PROMPT_HEADERS}
        for row in rows[header_index + 1:]:
            template = {
                "group_name": sheet.title.strip(),
                "field_name": _value_at(row, index_by_key, "field_name"),
                "source_field_id": _value_at(row, index_by_key, "source_field_id"),
                "field_order": _int_or_none(_value_at(row, index_by_key, "field_order")),
                "visible": _value_at(row, index_by_key, "visible").lower() not in {"否", "false", "0", "no"},
                "size_label": _value_at(row, index_by_key, "size_label") or "960x1280",
                "output_format": _value_at(row, index_by_key, "output_format") or "jpeg",
                "reference_fields": _value_at(row, index_by_key, "reference_fields"),
                "prompt_text": _value_at(row, index_by_key, "prompt_text"),
                "word_count": _int_or_none(_value_at(row, index_by_key, "word_count")),
                "field_type": _value_at(row, index_by_key, "field_type"),
                "female_priority": _int_or_none(_value_at(row, index_by_key, "female_priority")),
                "male_neutral_priority": _int_or_none(_value_at(row, index_by_key, "male_neutral_priority")),
                "enabled": True,
            }
            if template["field_name"] and template["prompt_text"]:
                templates.append(template)
    return templates


def parse_material_workbook(path: Path) -> dict[str, Any]:
    _assert(path and path.exists(), f"material data workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    _assert("概览" in workbook.sheetnames and "明细" in workbook.sheetnames, "material workbook must contain 概览 and 明细 sheets")
    overview_rows = _rows_from_header_sheet(workbook["概览"], MATERIAL_OVERVIEW_REQUIRED)
    detail_rows = _rows_from_header_sheet(workbook["明细"], MATERIAL_DETAIL_REQUIRED)
    return {
        "overview_rows": overview_rows,
        "detail_rows": detail_rows,
        "sheet_names": workbook.sheetnames,
    }


def parse_workflow_workbook(path: Path) -> list[dict[str, Any]]:
    _assert(path and path.exists(), f"workflow workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    _assert(rows, "workflow workbook is empty")
    headers = [_cell_text(cell) for cell in rows[0]]
    missing = WORKFLOW_REQUIRED - set(headers)
    _assert(not missing, f"workflow workbook missing headers: {', '.join(sorted(missing))}")
    index = {header: headers.index(header) for header in WORKFLOW_REQUIRED}
    workflow_rows = []
    for row_no, row in enumerate(rows[1:], start=2):
        style_code = _cell_text(row[index["款号"]])
        item_id = _cell_text(row[index["ID（用于测图的ID）"]])
        if not style_code or not item_id:
            continue
        workflow_rows.append({
            "row_no": row_no,
            "style_code": style_code,
            "item_id": item_id,
            "category": _cell_text(row[index["品类（后期匹配）"]]),
            "gender": _cell_text(row[index["模特性别"]]),
        })
    return workflow_rows


def _rows_from_header_sheet(sheet, required_headers: set[str]) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_cell_text(cell) for cell in rows[0]]
    missing = required_headers - set(headers)
    _assert(not missing, f"{sheet.title} sheet missing headers: {', '.join(sorted(missing))}")
    parsed = []
    for row in rows[1:]:
        item = {header: _cell_text(row[index]) for index, header in enumerate(headers) if header}
        if any(item.values()):
            parsed.append(item)
    return parsed


def _value_at(row: tuple[Any, ...], index_by_key: Mapping[str, int], key: str) -> str:
    index = index_by_key.get(key, -1)
    return _cell_text(row[index]) if index >= 0 and index < len(row) else ""


def _int_or_none(value: str) -> int | None:
    try:
        number = int(float(str(value).strip()))
    except Exception:
        return None
    return number


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _json_from_request(request) -> dict:
    data = getattr(request, "data", None)
    if not data:
        return {}
    try:
        payload = json.loads(data.decode("utf-8"))
    except Exception:
        return {}
    return payload if isinstance(payload, dict) else {}


def _token_type(request) -> str:
    header = request.headers.get("Authorization") or request.headers.get("authorization") or ""
    if ADMIN_TOKEN in header:
        return "user"
    if MACHINE_TOKEN in header:
        return "machine"
    return ""


def _batch_uid(path: str, suffix: str) -> str:
    return urllib.parse.unquote(path.removesuffix(suffix).rsplit("/", 1)[-1])


def _decision_parts(path: str) -> tuple[str, str]:
    match = path.split("/api/ai-image-batches/", 1)[-1]
    batch_uid, rest = match.split("/assets/", 1)
    asset_uid = rest.split("/decision", 1)[0]
    return urllib.parse.unquote(batch_uid), urllib.parse.unquote(asset_uid)


def _has_approved_ai(batch: Mapping[str, Any]) -> bool:
    return any(asset.get("kind") == "ai" and asset.get("status") == "approved" for asset in batch["assets"].values())


def _submit_plan(batch_uid: str, batch: Mapping[str, Any]) -> dict:
    approved_style_ids = {
        asset["style_id"]
        for asset in batch["assets"].values()
        if asset.get("kind") == "ai" and asset.get("status") == "approved"
    }
    styles = [style for style in batch["styles"] if style["id"] in approved_style_ids]
    assets = [
        asset for asset in batch["assets"].values()
        if asset.get("style_id") in approved_style_ids and (asset.get("kind") in {"source", "reference"} or asset.get("status") == "approved")
    ]
    return {"batch_uid": batch_uid, "styles": styles, "assets": assets}


def _assert(condition: Any, message: str) -> None:
    if not condition:
        raise AssertionError(message)


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
