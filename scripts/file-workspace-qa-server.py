"""Local, isolated P1 acceptance server. No model requests or user-data writes."""
import asyncio
import json
import os
import subprocess
import tempfile
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
ROOT = Path(__file__).resolve().parents[1]
QA_OUT = Path(os.environ.get('FILE_WORKSPACE_QA_OUTPUT') or ROOT / 'docs/crawshrimp-harness/qa-2026-09-30-deepseek/harness-p1')
QA_OUT.mkdir(parents=True, exist_ok=True)
DATA = Path(tempfile.mkdtemp(prefix='harness-file-workspace-p1-')).resolve()
os.environ['CRAWSHRIMP_DATA'] = str(DATA)
os.environ['CRAWSHRIMP_PYTHON_EXECUTABLE'] = str(ROOT / 'app/python-dist/mac-arm64/bin/python3.12')
os.environ['CRAWSHRIMP_OFFICE_ROOT'] = str(ROOT / 'build-staging/office/mac-arm64')
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from core import runtime_paths
from core.agent import api, db, mcp_gateway as gw, office_tools
from core.agent.artifact_delivery import collect_tool_deliveries
from core.agent.cdp import CdpClient
from core.cdp_bridge import get_bridge
runtime_paths.reset_runtime_data_root_cache()
db.init_agent_db(); db.create_session('qa-session-a','qa-a'); db.create_session('qa-session-b','qa-b')
db.create_run('qa-run-p1','qa-session-a',None,'qa-no-model','deterministic-tools')
CONTEXT = {'active_run': {'run_id':'qa-run-p1','session_id':'qa-session-a','runtime_session_id':'qa-a'}, 'current_tool_call_id':'qa-run-p1:browser', 'grant':None}
gw.ctx.workspace_root = DATA
activity = None
owned_tab = ''
trace = []
async def approve(*args): return 'approved'
gw.ctx.request_approval = approve

def emit(name, payload):
    global activity
    payload = {**payload, 'run_id':'qa-run-p1', 'runtime_session_id':'qa-a'}
    if name == 'browser.activity':
        payload.update(source='agent', tool_call_id=gw.ctx.current_tool_call_id)
        activity = payload
    db.append_event('qa-session-a', 'qa-run-p1', name, payload)
    trace.append({'event': name, 'payload':payload})
gw.ctx.emit_event = emit
app = FastAPI()
app.add_middleware(CORSMiddleware, allow_origins=['http://127.0.0.1:5173'], allow_methods=['*'], allow_headers=['*'])

@app.on_event('startup')
async def setup():
    token = gw.bind_tool_context(CONTEXT)
    try:
        # Execute a real shell command and feed its explicit output through the same delivery collector.
        code = "from pathlib import Path; import json; p=Path("+repr(str(DATA/'中文材料'))+"); p.mkdir(); (p/'销售.csv').write_text('商品,销量,单价,金额\\n童装A,12,99,1188\\n童装B,8,129,1032\\n童装C,5,159,795\\n', encoding='utf-8'); (p/'摘要.md').write_text('# 童装销售摘要\\n\\n总销售额 **3015 元**。\\n\\n| 商品 | 金额 |\\n|---|---:|\\n|童装A|1188|\\n|童装B|1032|\\n|童装C|795|\\n\\n'+('滚动保留验收\\n\\n'*100), encoding='utf-8'); print(json.dumps({'artifacts':[{'path':str(p/'销售.csv')},{'path':str(p/'摘要.md')}]},ensure_ascii=False))"
        html = '<!doctype html><html lang="zh"><head><meta charset="utf-8"><style>body{font-family:system-ui;padding:28px;background:#f6f7fa}article{padding:24px;background:white;border-radius:12px}h1{color:#ed6b42}td{padding:12px;border-bottom:1px solid #ddd}</style></head><body><article><h1>童装销售 HTML 交付</h1><p>总销售额 3015 元</p><table><tr><td>童装A</td><td>1188</td></tr><tr><td>童装B</td><td>1032</td></tr></table></article></body></html>'
        code += ";(p/'页面预览.html').write_text(" + repr(html) + ", encoding='utf-8')"
        proc = subprocess.run([sys.executable, '-c', code], capture_output=True, text=True, check=True)
        for item in collect_tool_deliveries({'tool_name':'fs_exec','run_id':'qa-run-p1','tool_call_id':'qa-shell'}, json.dumps({'ok':True,'data':{'exit_code':0,'stdout':proc.stdout}})): emit('artifact.created',item)
        material = DATA/'中文材料'
        (material/'销售图.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="720" height="450"><rect width="720" height="450" fill="#faf7f2"/><text x="40" y="65" font-size="28">童装销售 · 3015 元</text><rect x="75" y="130" width="125" height="245" fill="#ed6b42"/><rect x="280" y="163" width="125" height="212" fill="#9cb7ae"/><rect x="485" y="211" width="125" height="164" fill="#c2a97e"/><text x="82" y="410" font-size="20">童装A 1188</text><text x="287" y="410" font-size="20">童装B 1032</text><text x="492" y="410" font-size="20">童装C 795</text></svg>')
        (material/'大文件.log').write_text('中文边界验收\n'*70000)
        gw.tool_artifact_present([str(material/'销售图.svg'),str(material/'大文件.log'),str(material/'页面预览.html')])
        # Existing adapter collector and metadata shape, with real file readback.
        from core.agent.service import AgentService
        adapter = material/'适配器交付.txt'; adapter.write_text('适配器路径回读：3015 元')
        call = db.upsert_tool_call('qa-run-p1','adapter-call','task_run',{})
        db.update_tool_call(call['tool_call_id'],status='succeeded',result_json={'text':json.dumps({'evidence':{'task_instance_uid':'qa-adapter'}})})
        gw.ctx.list_task_artifacts = lambda uid: [{'id':'qa-adapter-file','path':str(adapter),'label':adapter.name,'kind':'file'}] if uid == 'qa-adapter' else []
        for item in AgentService._collect_run_artifacts('qa-run-p1'): emit('artifact.created',item)
        # Office generation and rendering go through original job contracts and bundled LibreOffice.
        office_code = """from pathlib import Path
import os
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
p=Path(os.environ['CRAWSHRIMP_OFFICE_OUTPUT'])
d=Document(); d.add_heading('童装销售摘要',0); d.add_paragraph('总销售额：3015 元'); d.add_page_break(); d.add_heading('第二页 · 明细核对',0); d.add_paragraph('连续滚动查看文件全部页面'); d.save(p/'办公摘要.docx')
r=Presentation(); slide=r.slides.add_slide(r.slide_layouts[1]); slide.shapes.title.text='童装销售摘要'; slide.placeholders[1].text='总销售额：3015 元'; second=r.slides.add_slide(r.slide_layouts[1]); second.shapes.title.text='第二页 · 销售明细'; second.placeholders[1].text='连续滚动展示下一页'; r.save(p/'办公演示.pptx')
w=Workbook(); s=w.active; s.append(['商品','金额']); s.append(['童装A',1188]); s.append(['童装B',1032]); s.append(['童装C',795]); second=w.create_sheet('第二页明细'); second.append(['第二页','连续滚动展示']); w.save(p/'办公销售.xlsx')
"""
        started = await office_tools.office_run(office_code)
        if not started['ok']: raise RuntimeError(started)
        async def wait_job(job):
            for _ in range(120):
                result = office_tools.office_job(job)
                if result.get('status') not in ('queued','running'): return result
                await asyncio.sleep(.5)
            raise RuntimeError('Office QA timeout')
        generated = await wait_job(started['data']['job_id'])
        if generated.get('status') != 'completed': raise RuntimeError(generated)
        for filename in ['办公摘要.docx','办公演示.pptx','办公销售.xlsx']:
            rendered = await office_tools.office_render(started['data']['job_id'],filename)
            if not rendered['ok']: raise RuntimeError(rendered)
            result = await wait_job(rendered['data']['job_id'])
            if result.get('status') != 'completed': raise RuntimeError(result)
            pdf = result['data']['result'].get('pdf')
            if pdf: gw.tool_artifact_present([pdf])
        trace.append({'office':'all three generated/rendered','data':str(DATA)})
        (QA_OUT/'backend-trace.json').write_text(json.dumps(trace,ensure_ascii=False,indent=2))
        print('QA_READY '+str(DATA),flush=True)
    finally: gw.reset_tool_context(token)

@app.get('/agent/session-resources')
def resources(runtime_session_id:str): return api.session_resources(runtime_session_id)

@app.get('/qa/media')
def media(path:str):
    target=Path(path).resolve()
    if not target.is_relative_to(DATA) or not target.is_file(): raise HTTPException(404,'QA file missing')
    return FileResponse(target, headers={'Access-Control-Expose-Headers':'Content-Range, Content-Length'})

@app.post('/agent/session-resources/pages/{tab_id}/select')
async def select_page(tab_id:str, request:api.SessionPageRequest):
    class Service:
        async def broadcast(self, session_id, seq, event, payload):
            db.append_event(session_id,None,event,payload)
    original=api.get_agent_service
    api.get_agent_service=lambda:Service()
    try: return await api.select_session_page(tab_id,request)
    finally: api.get_agent_service=original

@app.get('/agent/approvals')
def approvals(): return {'approvals':[]}

@app.get('/qa/state')
def state(): return {'data':str(DATA), 'activity':activity, 'trace':trace[-6:]}

@app.post('/qa/browser/{operation}')
async def browser(operation:str):
    global owned_tab
    token=gw.bind_tool_context(CONTEXT)
    try:
        if operation == 'navigate':
            result=await gw.tool_browser_navigate('https://example.com',new_tab=not owned_tab)
            owned_tab=(result.get('data') or {}).get('tab_id') or owned_tab
            grant=gw.ctx.grant or {}; grant['toolset_json']='["act","observe","eval","verify"]'; CONTEXT['grant']=grant
        elif operation == 'click': result=await gw.tool_browser_act('click',selector='a',text='Learn more')
        elif operation == 'verify':
            result=await gw.tool_browser_eval('({url:location.href,title:document.title})')
            result['verification']=await gw.tool_browser_verify('location.hostname.endsWith("iana.org")')
        else: raise HTTPException(400,'Unknown QA operation')
        trace.append({'operation':operation,'result':result,'target_id':owned_tab})
        (QA_OUT/'browser-trace.json').write_text(json.dumps(trace[-12:],ensure_ascii=False,indent=2))
        return {'result':result,'activity':activity}
    finally: gw.reset_tool_context(token)

@app.on_event('shutdown')
async def cleanup():
    if owned_tab: await get_bridge().close_tab_async(owned_tab)

if __name__ == '__main__':
    import uvicorn
    uvicorn.run(app,host='127.0.0.1',port=int(os.environ.get('FILE_WORKSPACE_QA_PORT', '5189')))
