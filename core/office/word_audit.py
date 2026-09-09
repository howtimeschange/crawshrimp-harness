"""Read-only Word layout diagnostics; heuristics are warnings, never visual proof."""
from __future__ import annotations

import re
from docx import Document
from docx.oxml.ns import qn
from .word_design import PRESETS


def audit_word(path) -> dict:
    doc = Document(path)
    issues = []
    def issue(code, location, message, severity="warning"):
        issues.append(dict(code=code, location=location, message=message, severity=severity))
    previous_level = 0
    paragraphs = [(f"正文段落 {i}", p) for i, p in enumerate(doc.paragraphs, 1)]
    for t, table in enumerate(doc.tables, 1):
        seen = set()
        for r, row in enumerate(table.rows, 1):
            for c, cell in enumerate(row.cells, 1):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                paragraphs.extend((f"表 {t} 行 {r} 列 {c}", p) for p in cell.paragraphs)
    for location, paragraph in paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style
        heading = re.fullmatch(r"Heading (\d+)", style.name)
        if heading and location.startswith("正文"):
            level = int(heading[1])
            if level > previous_level + 1:
                issue("WORD_HEADING_LEVEL_GAP", location, f"标题从 {previous_level} 级跳到 {level} 级")
            previous_level = level
        elif style.name not in ("Title", "Subtitle", "Caption") and len(text) < 65:
            looks_numbered = re.match(r"^(第[一二三四五六七八九十\d]+[章节]|[一二三四五六七八九十]+、|\d+(?:\.\d+)+\s)", text)
            large_bold = any(r.bold and r.font.size and r.font.size.pt >= 14 for r in paragraph.runs)
            if looks_numbered or large_bold:
                issue("WORD_POSSIBLE_FAKE_HEADING", location, "疑似标题使用正文样式，请确认语义层级")
        for run in paragraph.runs:
            if run.text.strip() and run.font.size and style.font.size and abs(run.font.size.pt - style.font.size.pt) > .5:
                issue("WORD_DIRECT_SIZE_OVERRIDE", location, "局部字号覆盖样式，确认是否为有意强调")
                break
        if paragraph.paragraph_format.keep_with_next and not heading and len(text) > 150:
            issue("WORD_LONG_KEEP_WITH_NEXT", location, "长正文与下段绑定可能产生大块空白")
    marker = re.search(r"\bharness-word:(\w+)\b", doc.core_properties.keywords or "")
    preset = marker[1] if marker else None
    if preset in PRESETS:
        p = PRESETS[preset]
        for name, expected in [("Normal", p["body"]), ("Heading 1", p["h1"]), ("Heading 2", p["h2"])]:
            size = doc.styles[name].font.size
            if size is None or abs(size.pt - expected) > .5:
                issue("WORD_PRESET_STYLE_DRIFT", name, f"字号应为 {expected} pt")
    # Validate native width declarations, independent of python-docx's computed cell widths.
    max_page_width = max(s.page_width - s.left_margin - s.right_margin for s in doc.sections) / 635
    for index, table in enumerate(doc.tables, 1):
        location = f"表 {index}"
        grid = [int(c.get(qn("w:w"), 0)) for c in table._tbl.tblGrid]
        node = table._tbl.tblPr.find(qn("w:tblW"))
        if node is not None and node.get(qn("w:type")) == "dxa":
            declared = int(node.get(qn("w:w"), 0))
            if abs(declared - sum(grid)) > 3:
                issue("WORD_TABLE_GRID_MISMATCH", location, "表宽与列网格总宽不一致", "error")
        if sum(grid) > max_page_width + 5:
            issue("WORD_TABLE_TOO_WIDE", location, "表格宽度超过文档最大正文宽度", "error")
        mismatch = False
        for row in table._tbl.tr_lst:
            col = row.grid_before
            for cell in row.tc_lst:
                span = cell.grid_span
                width = cell.tcPr.find(qn("w:tcW")) if cell.tcPr is not None else None
                if width is not None and width.get(qn("w:type")) == "dxa":
                    mismatch |= abs(int(width.get(qn("w:w"), 0)) - sum(grid[col:col + span])) > 3
                col += span
        if mismatch:
            issue("WORD_TABLE_CELL_MISMATCH", location, "单元格宽度与对应网格列不一致", "error")
        if len(table.rows) > 10 and table.rows[0]._tr.trPr is not None:
            header = table.rows[0]._tr.trPr.find(qn("w:tblHeader"))
            if header is None or header.get(qn("w:val")) in ("false", "0", "off"):
                issue("WORD_TABLE_HEADER_NOT_REPEATED", location, "长表未设置重复表头")
        elif len(table.rows) > 10:
            issue("WORD_TABLE_HEADER_NOT_REPEATED", location, "长表未设置重复表头")
    return dict(preset=preset, issues=issues, status="issues" if issues else "passed", visual="not_run")
