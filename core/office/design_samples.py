"""Original, explicitly labeled fixtures for native rendering and client acceptance."""
from __future__ import annotations

import json
from pathlib import Path
from openpyxl import Workbook
from .word_design import PRESETS, new_document, add_table
from .ppt_design import build_deck, new_deck, add_slide, layout_catalog
from .ppt_template import inspect_template, follow_template
from .excel_design import format_report


def generate_samples(output, *, atlas=False):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    files = []
    for preset, values in PRESETS.items():
        title = f"{values['label']} · 中文排版验收"
        doc = new_document(title, preset=preset, subtitle="第三季度经营与试点计划", metadata={"性质": "示例数据", "日期": "2026-09-09"})
        doc.add_heading("一、摘要与目标", 1)
        doc.add_paragraph("本文件使用 Harness 内置中文字体、语义标题和原生表格。所有数字仅用于验证排版与公式，不代表实际业务。建议先进行小范围试点，再根据收入、成本和利润变化决定后续投入。")
        doc.add_heading("二、经营数据", 1)
        add_table(doc, ["指标", "本期", "说明"], [["收入", "150 万元", "示例：较基期增长 50%"], ["成本", "90 万元", "含试点期间投入"], ["利润", "60 万元", "收入减成本"]], preset=preset)
        doc.add_heading("三、执行安排", 1)
        doc.add_paragraph("运营负责人确认范围，财务核对数据口径，项目组按周复盘。风险与未解决事项应在下一次会议中明确责任人和完成时间。")
        path = output / f"{values['label']}.docx"
        doc.save(path)
        files.append(path)
    chart = {"categories": ["六月", "七月", "八月"], "series": {"收入": [100, 120, 150]}, "unit": "万元"}
    pages = [
        dict(layout="cover", title="以小范围试点验证增长质量", takeaway="先验证，再扩展", content={"subtitle": "第三季度经营汇报", "context": "管理层审议 · 示例数据"}),
        dict(layout="metric-three", title="收入提升，仍需关注利润与投入", takeaway="同时看增长和回报", content={"metric_a": "150 万元\n本期收入", "metric_b": "60 万元\n本期利润", "metric_c": "40%\n利润率", "explanation": "本期数据均为示例；利润为收入减成本，不代表完整财务口径。"}),
        dict(layout="chart-insight", title="收入连续增长，投入应分阶段验证", takeaway="增长需要证据支持", content={"chart": chart, "insight": "六月至八月收入由 100 增至 150 万元。\n\n下一步同时跟踪渠道成本和留存。"}),
        dict(layout="compare-two", title="比较两种投入方式，优先选择可验证方案", takeaway="先试点能控制不确定性", content={"option_a": "方案 A：全面扩张\n\n覆盖范围大\n投入集中\n验证周期较长", "option_b": "方案 B：分阶段试点\n\n明确试点边界\n按周检查数据\n达到目标再扩大"}),
        dict(layout="process", title="按三个阶段推进，并保留复盘节点", takeaway="每一步都有检查结果", content={"steps": ["确认范围与口径", "执行四周试点", "复盘投入与回报"], "outcome": "产出：试点结果、风险清单和是否扩大投入的建议。"}),
        dict(layout="action-plan", title="明确负责人和验收标准", takeaway="把决策落实到具体行动", content={"actions": [["行动", "负责人", "验收"], ["确认口径", "财务", "数据可核对"], ["执行试点", "运营", "周报完整"], ["复盘决策", "项目组", "结果有依据"]], "decision": "本次需审议\n\n试点范围\n预算上限\n复盘时间"}),
    ]
    for page in pages:
        page.update(source="来源：Harness 排版验收示例数据", notes=page["takeaway"])
    plan = dict(audience="经营管理层", purpose="审议试点方案", message="先验证增长质量，再扩展投入", slides=pages)
    deck = output / "经营汇报.pptx"
    build_deck(plan, deck)
    files.append(deck)
    info = inspect_template(deck)
    title_id = next(s["id"] for s in info["slides"][0]["shapes"] if s["name"] == "title")
    filled = output / "公司模板填充.pptx"
    follow_template(deck, filled, [{"slide": 1, "texts": {title_id: "公司模板填充：试点投入审议"}}], source_sha256=info["sha256"])
    files.append(filled)
    book = Workbook()
    detail = book.active
    detail.title = "数据"
    for row in [["商品编号", "收入", "成本", "利润"], ["00123", 120, 70, "=B2-C2"], ["00456", 100, 65, "=B3-C3"], ["00789", 150, 90, "=B4-C4"]]:
        detail.append(row)
    format_report(detail, "A1:D4", columns={"A": "text", "B": "amount", "C": "amount", "D": "amount"},
                  roles={"B2:C4": "input", "D2:D4": "formula"}, title="经营明细（示例）", financial=True)
    summary = book.create_sheet("汇总")
    for row in [["指标", "结果"], ["收入", "=SUM(数据!B2:B4)"], ["成本", "=SUM(数据!C2:C4)"], ["利润", "=B2-B3"], ["利润率", "=IF(B2=0,0,B4/B2)"]]:
        summary.append(row)
    format_report(summary, "A1:B5", columns={"B": "amount"}, roles={"B2:B5": "summary"}, title="经营汇总（示例）")
    summary["B5"].number_format = "0.0%"
    note = book.create_sheet("说明")
    for row in [["项目", "约定"], ["数据", "示例数据，不代表实际业务"], ["币种与单位", "人民币，万元"], ["颜色", "蓝色为输入，灰色为公式，绿色为汇总"], ["缺失值", "缺失不等于零，需单独核对"], ["口径", "利润仅为收入减成本的示例计算"]]:
        note.append(row)
    format_report(note, "A1:B6", title="口径说明")
    note.column_dimensions["B"].width = 48
    excel = output / "经营报表.xlsx"
    book.save(excel)
    files.append(excel)
    if atlas:
        from .runtime import configure_fonts
        configure_fonts()
        import matplotlib.pyplot as plt
        fig, ax = plt.subplots(figsize=(6, 4))
        ax.bar(["六月", "七月", "八月"], [100, 120, 150], color="#24546A")
        ax.set_title("示例收入（万元）")
        fig.tight_layout()
        image = output / "示例图.png"
        fig.savefig(image)
        plt.close(fig)
        atlas_deck = new_deck()
        values = {"table": [["指标", "数值"], ["收入", "150 万元"], ["利润", "60 万元"]], "chart": chart,
                  "image": str(image), "process": ["确认目标", "执行方案", "验证结果"], "list": ["核心结论", "数据证据", "行动建议"],
                  "metric": "18%\n同比增长", "lead": "聚焦结论\n让证据支持行动", "body": "经营持续改善，保持验证节奏。", "panel": "目标清楚\n证据充分\n行动可执行"}
        for name, slots in layout_catalog().items():
            add_slide(atlas_deck, name, f"版式验证 · {name}", {s["name"]: values[s["kind"]] for s in slots}, source="原创版式 · 示例内容")
        path = output / "26种原生版式.pptx"
        atlas_deck.save(path)
        files.append(path)
    return files
