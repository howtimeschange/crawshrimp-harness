"""Read back real Office documents; keep cached formula values distinct from formulas."""
from __future__ import annotations

from pathlib import Path
from .runtime import OfficeError, file_hash

ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}


def validate(path: Path, expected: dict | None = None) -> dict:
    expected = expected or {}
    issues, texts = [], []
    result = {"path": str(path), "revision": file_hash(path), "issues": issues}
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(path)
        texts = [p.text for p in doc.paragraphs]
        texts.extend(c.text for t in doc.tables for r in t.rows for c in r.cells)
        result.update(paragraphs=len(doc.paragraphs), tables=len(doc.tables))
    elif path.suffix.lower() == ".pptx":
        from pptx import Presentation
        doc = Presentation(path)
        result["slides"] = len(doc.slides)
        for index, slide in enumerate(doc.slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
                if (shape.left < 0 or shape.top < 0 or
                    shape.left + shape.width > doc.slide_width + 100 or
                    shape.top + shape.height > doc.slide_height + 100):
                    issues.append({"code": "SHAPE_OUTSIDE_SLIDE", "page": index, "shape": shape.name})
    elif path.suffix.lower() == ".xlsx":
        import openpyxl
        formulas = openpyxl.load_workbook(path, data_only=False)
        cached = openpyxl.load_workbook(path, data_only=True)
        try:
            result["sheets"] = [{"name": s.title, "rows": s.max_row, "columns": s.max_column,
                                 "state": s.sheet_state, "print_area": str(s.print_area)} for s in formulas]
            result["formula_count"] = 0
            for sheet in formulas:
                for row in sheet:
                    for cell in row:
                        value = cell.value
                        if value is not None:
                            texts.append(str(value))
                        if cell.data_type == "f":
                            result["formula_count"] += 1
                            cv = cached[sheet.title][cell.coordinate].value
                            if cv is None:
                                issues.append({"code": "FORMULA_CACHE_MISSING", "cell": f"{sheet.title}!{cell.coordinate}"})
                            elif str(cv) in ERROR_VALUES:
                                issues.append({"code": "FORMULA_ERROR", "cell": f"{sheet.title}!{cell.coordinate}", "value": cv})
            for key, value in expected.get("cells", {}).items():
                sheet, address = key.rsplit("!", 1)
                actual = cached[sheet][address].value
                if actual != value:
                    issues.append({"code": "CELL_MISMATCH", "cell": key, "expected": value, "actual": actual})
        finally:
            formulas.close()
            cached.close()
    else:
        raise OfficeError("OFFICE_FORMAT_UNSUPPORTED", "仅支持 docx、pptx、xlsx。")
    text = "\n".join(texts)
    for needle in expected.get("contains", []):
        if needle not in text:
            issues.append({"code": "MISSING_TEXT", "text": needle})
    if "slides" in expected and result.get("slides") != expected["slides"]:
        issues.append({"code": "SLIDE_COUNT", "expected": expected["slides"], "actual": result.get("slides")})
    result["status"] = "issues" if issues else "passed"
    return result
