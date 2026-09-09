"""Conservative template filling with explicit slide/shape mapping and fidelity readback."""
from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
from .runtime import file_hash

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _canonical(data):
    try:
        return etree.tostring(etree.fromstring(data), method="c14n")
    except etree.XMLSyntaxError:
        return data


def inspect_template(path):
    prs = Presentation(path)
    def inspect_shape(sh):
        result = dict(id=sh.shape_id, name=sh.name,
                      kind="text" if sh.has_text_frame else "table" if sh.has_table else "other",
                      text=sh.text if sh.has_text_frame else None,
                      paragraphs=len(sh.text_frame.paragraphs) if sh.has_text_frame else 0,
                      placeholder=sh.is_placeholder,
                      box=[sh.left, sh.top, sh.width, sh.height])
        if sh.has_table:
            result["table"] = dict(rows=len(sh.table.rows), columns=len(sh.table.columns),
                                   cells=[[c.text for c in row.cells] for row in sh.table.rows],
                                   merged=any(c.is_merge_origin or c.is_spanned for row in sh.table.rows for c in row.cells))
        return result
    return dict(source=str(path), sha256=file_hash(Path(path)),
                size_inches=[round(prs.slide_width / 914400, 4), round(prs.slide_height / 914400, 4)],
                slides=[dict(slide=i, layout=s.slide_layout.name,
                             shapes=[inspect_shape(sh) for sh in s.shapes])
                        for i, s in enumerate(prs.slides, 1)])


def _replace_frame(frame, value):
    lines = str(value).split("\n")
    if len(lines) > len(frame.paragraphs):
        raise ValueError("替换段落多于模板段落，请缩短内容或明确调整版式")
    for i, p in enumerate(frame.paragraphs):
        text = lines[i] if i < len(lines) else ""
        if p.runs:
            p.runs[0].text = text
            for r in p.runs[1:]:
                r.text = ""
        elif text:
            p.add_run().text = text


def _shape_signature(shape):
    tree = etree.fromstring(etree.tostring(shape._element))
    for node in tree.findall(".//a:t", NS):
        node.text = ""
    # A blank placeholder may acquire its first unformatted text run.
    for run in tree.findall(".//a:r", NS):
        if len(run) == 1 and run[0].tag == f"{{{NS['a']}}}t":
            run.getparent().remove(run)
    return etree.tostring(tree, method="c14n")


def check_fidelity(source, output, edited_shapes=None):
    """Check untouched package parts plus geometry/formatting of filled shapes."""
    allowed = {(int(page), int(shape)) for page, shape in (edited_shapes or [])}
    before, after = Presentation(source), Presentation(output)
    issues = []
    changed_parts = set()
    if (before.slide_width, before.slide_height, len(before.slides)) != (after.slide_width, after.slide_height, len(after.slides)):
        issues.append(dict(code="PPT_TEMPLATE_PAGE_STRUCTURE_CHANGED"))
    for page, (a, b) in enumerate(zip(before.slides, after.slides), 1):
        am, bm = {s.shape_id: s for s in a.shapes}, {s.shape_id: s for s in b.shapes}
        if list(am) != list(bm):
            issues.append(dict(code="PPT_TEMPLATE_SHAPES_CHANGED", page=page))
        for ident in am.keys() & bm.keys():
            planned = (page, ident) in allowed
            a_sig = _shape_signature(am[ident]) if planned else _canonical(etree.tostring(am[ident]._element))
            b_sig = _shape_signature(bm[ident]) if planned else _canonical(etree.tostring(bm[ident]._element))
            if a_sig != b_sig:
                issues.append(dict(code="PPT_TEMPLATE_SHAPE_DRIFT", page=page, shape_id=ident))
            if planned:
                changed_parts.add(str(a.part.partname).lstrip("/"))
    with ZipFile(source) as a, ZipFile(output) as b:
        if set(a.namelist()) != set(b.namelist()):
            issues.append(dict(code="PPT_TEMPLATE_PACKAGE_PARTS_CHANGED"))
        preserved = 0
        for name in set(a.namelist()) & set(b.namelist()):
            if name in changed_parts:
                # Slide background, transitions, notes relationships and shape tree order remain fixed.
                x, y = etree.fromstring(a.read(name)), etree.fromstring(b.read(name))
                for tree in (x, y):
                    for st in tree.findall(".//{http://schemas.openxmlformats.org/presentationml/2006/main}spTree"):
                        for child in list(st)[2:]:
                            st.remove(child)
                same = etree.tostring(x, method="c14n") == etree.tostring(y, method="c14n")
            else:
                same = _canonical(a.read(name)) == _canonical(b.read(name))
            if not same:
                issues.append(dict(code="PPT_TEMPLATE_PART_CHANGED", part=name))
            else:
                preserved += 1
    return dict(status="issues" if issues else "passed", issues=issues,
                source_sha256=file_hash(Path(source)), output_sha256=file_hash(Path(output)),
                checked_parts=preserved, visual="not_run")


def follow_template(source, output, plan, *, source_sha256=None):
    """plan=[{slide:1,texts:{shape_id:text},tables:{shape_id:rows}}]; preserve every page."""
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve():
        raise ValueError("模板原件不能覆盖，请指定新的输出文件")
    if source_sha256 and file_hash(source) != source_sha256:
        raise ValueError("模板已更新，请重新检查和映射")
    prs = Presentation(source)
    edits, used, changes = [], set(), []
    if not plan:
        raise ValueError("页面映射不能为空")
    for item in plan:
        page = item["slide"]
        if not isinstance(page, int) or not 1 <= page <= len(prs.slides) or page in used:
            raise ValueError("页码须存在且不能重复")
        used.add(page)
        shapes = {s.shape_id: s for s in prs.slides[page - 1].shapes}
        ids = set()
        for kind in ("texts", "tables"):
            for raw, value in item.get(kind, {}).items():
                ident = int(raw)
                if ident in ids or ident not in shapes:
                    raise ValueError("形状 ID 无效或重复；先 inspect_template 获取准确 ID")
                ids.add(ident)
                shape = shapes[ident]
                if kind == "texts":
                    if not shape.has_text_frame:
                        raise ValueError("目标不是文本形状")
                    _replace_frame(shape.text_frame, value)
                else:
                    if not shape.has_table or len(value) != len(shape.table.rows) or any(len(r) != len(shape.table.columns) for r in value):
                        raise ValueError("替换表格须与模板行列数一致")
                    if any(c.is_merge_origin or c.is_spanned for r in shape.table.rows for c in r.cells):
                        raise ValueError("合并表格需单独映射，当前不自动替换")
                    for row, values in zip(shape.table.rows, value):
                        for cell, text in zip(row.cells, values):
                            _replace_frame(cell.text_frame, text)
                edits.append((page, ident))
                changes.append(dict(slide=page, shape_id=ident, operation=kind))
    if not edits:
        raise ValueError("映射没有指定替换内容")
    prs.save(output)
    report = check_fidelity(source, output, edits)
    report["changes"] = changes
    report["limitations"] = ["保留页面与段落结构；不自动复制幻灯片或修改图表/图片/SmartArt", "替换文本沿用首个 run 的样式，新的字词强调需人工指定", "结构保真不能代替渲染后的文本溢出检查"]
    output.with_suffix(".fidelity.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if report["issues"]:
        raise ValueError("模板保真检查未通过，请读取 .fidelity.json")
    return report
