"""Explicit-range report formatting without changing workbook values or formulas."""
from __future__ import annotations

import math
import unicodedata

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries

ROLES = {"input": "E9F1FA", "formula": "F2F5F7", "summary": "DDEDE8"}
FORMATS = {"amount": '#,##0.00;[Red](#,##0.00);"—"', "integer": '#,##0',
           "percent": '0.0%', "date": 'yyyy-mm-dd', "text": '@'}


def _text_height(cell, width):
    """Conservative estimate in points; formulas need rendered cached-value review."""
    if not isinstance(cell.value, str) or cell.data_type == "f":
        return 0
    size = cell.font.sz or 11
    capacity = max(1, (width - 2) * 11 / size)
    lines = 0
    for line in cell.value.split("\n"):
        units = sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line)
        lines += max(1, math.ceil(units / capacity)) if cell.alignment.wrap_text else 1
    return lines * size * 1.4 + 8


def format_report(ws, area: str, *, columns=None, roles=None, title="", financial=False):
    """Format only area. columns={'B':'amount'}, roles={'B2:B5':'input'}; no merges."""
    min_col, min_row, max_col, max_row = range_boundaries(area)
    if (None in (min_col, min_row, max_col, max_row)
            or not (1 <= min_col <= max_col <= 16384 and 1 <= min_row <= max_row <= 1048576)
            or (max_col - min_col + 1) * (max_row - min_row + 1) > 100000):
        raise ValueError("请提供有界矩形区域，最多 100000 个单元格")
    column_formats = {}
    for col, name in (columns or {}).items():
        from openpyxl.utils.cell import column_index_from_string
        index = column_index_from_string(col)
        if not min_col <= index <= max_col:
            raise ValueError("格式列不在报表范围内")
        column_formats[index] = FORMATS[name]
    role_regions = []
    for region, role in (roles or {}).items():
        bounds = range_boundaries(region)
        if None in bounds or not (min_col <= bounds[0] <= bounds[2] <= max_col and min_row < bounds[1] <= bounds[3] <= max_row):
            raise ValueError("角色范围必须位于正文数据区内")
        role_regions.append((bounds, ROLES[role]))
    # Validate first, then mutate formatting. Cell values/formulas remain untouched.
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            header = cell.row == min_row
            cell.font = Font(name="思源黑体", size=11, bold=header, color="FFFFFF" if header else "263340")
            cell.fill = PatternFill("solid", fgColor="24546A" if header else "FFFFFF")
            cell.border = Border(bottom=Side(style="thin", color="D5DEE3") if header else Side())
            cell.alignment = Alignment(horizontal="left" if header or not isinstance(cell.value, (int, float)) and cell.data_type != "f" else "right", vertical="center", wrap_text=True)
            if not header and cell.column in column_formats:
                cell.number_format = column_formats[cell.column]
            for (left, top, right, bottom), color in role_regions:
                if left <= cell.column <= right and top <= cell.row <= bottom:
                    cell.fill = PatternFill("solid", fgColor=color)
        ws.row_dimensions[row[0].row].height = 28 if row[0].row == min_row else 25
    for col in range(min_col, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 26 if col == min_col else 20
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        required = max((_text_height(c, ws.column_dimensions[c.column_letter].width)
                        for c in row if c.__class__.__name__ != "MergedCell"), default=0)
        # Calc normalizes Excel column widths on round-trip. Reserve extra wrapping
        # room for that change and for font-specific word boundaries.
        if required > 25:
            required += max(18, required * .2)
        ws.row_dimensions[row[0].row].height = min(409, max(ws.row_dimensions[row[0].row].height, required))
    ws.freeze_panes = f"{get_column_letter(min_col)}{min_row + 1}"
    # Existing filters/tables are structural user data, preserve them.
    if not ws.auto_filter.ref and not ws.tables:
        ws.auto_filter.ref = area
    ws.sheet_view.showGridLines = False
    ws.print_area = area
    ws.print_title_rows = f"{min_row}:{min_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape" if max_col - min_col >= 4 else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    if title:
        ws.oddHeader.center.text = title
    ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"
    if financial:
        # This is a presentation hint; no business formula or sign convention is inferred.
        ws.oddFooter.left.text = "金额口径以说明页为准"
    return ws


def audit_report(ws) -> dict:
    issues = []
    def issue(code, cell, message):
        issues.append(dict(code=code, cell=cell, message=message, severity="warning"))
    if ws.max_row > 15 and not ws.freeze_panes:
        issue("EXCEL_HEADER_NOT_FROZEN", "", "长表未冻结表头")
    if not ws.print_area:
        issue("EXCEL_PRINT_AREA_MISSING", "", "未指定打印区域")
    # Do not traverse phantom formatted ranges from imported workbooks.
    for cell in ws._cells.values():
        if cell.value is None:
            continue
        if isinstance(cell.value, str) and len(cell.value) > 40 and not cell.alignment.wrap_text:
            issue("EXCEL_LONG_TEXT_NOT_WRAPPED", cell.coordinate, "长文字未启用换行")
        if cell.alignment.wrap_text and cell.coordinate not in ws.merged_cells:
            height = ws.row_dimensions[cell.row].height
            width = ws.column_dimensions[cell.column_letter].width
            if height is not None and _text_height(cell, width) > height + 3:
                issue("EXCEL_TEXT_FIT_ESTIMATE", cell.coordinate, "估算换行文字超出行高；请加宽列或增加行高后渲染确认")
        if cell.data_type == "f" and any(cell.coordinate in region for region in ws.merged_cells.ranges):
            issue("EXCEL_FORMULA_IN_MERGED_RANGE", cell.coordinate, "计算区域存在合并单元格")
        if cell.data_type == "e":
            issue("EXCEL_ERROR_VALUE", cell.coordinate, str(cell.value))
    return dict(sheet=ws.title, issues=issues, status="issues" if issues else "passed", visual="not_run")
