"""26 original Harness slide recipes using editable python-pptx objects."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

THEMES = {
    "business": dict(font="思源黑体", ink="243746", accent="24546A", pale="EDF3F6", background="FFFFFF"),
    "proposal": dict(font="思源黑体", ink="243B38", accent="25635E", pale="EEF5F2", background="FFFFFF"),
    "mono": dict(font="思源黑体", ink="252B33", accent="414E60", pale="F0F2F5", background="FFFFFF"),
}


def slot(name, kind, x, y, w, h):
    return dict(name=name, kind=kind, box=[x, y, w, h])


# Every box is in inches on a 13.3333 × 7.5 canvas; title/footer are shared anchors.
LAYOUTS = {
    "cover": [slot("subtitle", "lead", .7, 2.7, 9, 1.4), slot("context", "body", .7, 4.7, 8, 1.2)],
    "section": [slot("number", "metric", .7, 2.1, 2.1, 2), slot("summary", "lead", 3.2, 2.1, 9.3, 2)],
    "agenda": [slot("items", "list", .7, 1.8, 8, 4.7), slot("context", "body", 9.1, 1.8, 3.5, 4.7)],
    "conclusion": [slot("message", "lead", .7, 1.8, 11.9, 1.5), slot("evidence", "body", .7, 3.7, 7.3, 2.5), slot("action", "body", 8.5, 3.7, 4.1, 2.5)],
    "two-column": [slot("left", "body", .7, 1.8, 5.7, 4.7), slot("right", "body", 6.9, 1.8, 5.7, 4.7)],
    "three-column": [slot("first", "body", .7, 1.8, 3.6, 4.7), slot("second", "body", 4.85, 1.8, 3.6, 4.7), slot("third", "body", 9, 1.8, 3.6, 4.7)],
    "compare-two": [slot("option_a", "panel", .7, 1.8, 5.7, 4.7), slot("option_b", "panel", 6.9, 1.8, 5.7, 4.7)],
    "compare-three": [slot("option_a", "panel", .7, 1.8, 3.6, 4.7), slot("option_b", "panel", 4.85, 1.8, 3.6, 4.7), slot("option_c", "panel", 9, 1.8, 3.6, 4.7)],
    "before-after": [slot("before", "panel", .7, 1.8, 5.7, 3.5), slot("after", "panel", 6.9, 1.8, 5.7, 3.5), slot("change", "body", .7, 5.7, 11.9, .8)],
    "problem-solution": [slot("problem", "lead", .7, 1.8, 4, 4.7), slot("solution", "body", 5.3, 1.8, 7.3, 4.7)],
    "quote": [slot("quote", "lead", 1.2, 2, 10.9, 2.6), slot("attribution", "body", 1.2, 5, 10.9, 1)],
    "image-left": [slot("image", "image", .7, 1.8, 6.3, 4.7), slot("explanation", "body", 7.5, 1.8, 5.1, 4.7)],
    "image-right": [slot("explanation", "body", .7, 1.8, 5.1, 4.7), slot("image", "image", 6.3, 1.8, 6.3, 4.7)],
    "image-wide": [slot("image", "image", .7, 1.7, 11.9, 3.9), slot("caption", "body", .7, 5.9, 11.9, .6)],
    "two-images": [slot("image_a", "image", .7, 1.8, 5.7, 3.5), slot("image_b", "image", 6.9, 1.8, 5.7, 3.5), slot("caption", "body", .7, 5.7, 11.9, .8)],
    "metric-focus": [slot("value", "metric", .7, 2, 4.2, 2.2), slot("explanation", "body", 5.5, 2, 7.1, 3.8)],
    "metric-three": [slot("metric_a", "metric", .7, 2, 3.6, 2), slot("metric_b", "metric", 4.85, 2, 3.6, 2), slot("metric_c", "metric", 9, 2, 3.6, 2), slot("explanation", "body", .7, 4.6, 11.9, 1.7)],
    "chart-wide": [slot("chart", "chart", .7, 1.7, 11.9, 4), slot("insight", "body", .7, 6, 11.9, .55)],
    "chart-insight": [slot("chart", "chart", .7, 1.8, 7.7, 4.7), slot("insight", "body", 8.9, 1.8, 3.7, 4.7)],
    "chart-pair": [slot("chart_a", "chart", .7, 1.8, 5.7, 3.8), slot("chart_b", "chart", 6.9, 1.8, 5.7, 3.8), slot("insight", "body", .7, 5.9, 11.9, .65)],
    "table-wide": [slot("table", "table", .7, 1.8, 11.9, 4.1), slot("note", "body", .7, 6.1, 11.9, .45)],
    "table-insight": [slot("table", "table", .7, 1.8, 8, 4.7), slot("insight", "body", 9.2, 1.8, 3.4, 4.7)],
    "process": [slot("steps", "process", .7, 2.3, 11.9, 2.7), slot("outcome", "body", .7, 5.5, 11.9, 1)],
    "timeline": [slot("stages", "process", .7, 2, 11.9, 3.8), slot("note", "body", .7, 6.1, 11.9, .45)],
    "four-quadrant": [slot("top_left", "panel", .7, 1.8, 5.7, 2.1), slot("top_right", "panel", 6.9, 1.8, 5.7, 2.1), slot("bottom_left", "panel", .7, 4.35, 5.7, 2.1), slot("bottom_right", "panel", 6.9, 4.35, 5.7, 2.1)],
    "action-plan": [slot("actions", "table", .7, 1.8, 8.3, 4.7), slot("decision", "panel", 9.5, 1.8, 3.1, 4.7)],
}


def layout_catalog():
    return deepcopy(LAYOUTS)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    return prs


def _text(slide, name, text, box, theme, *, size=22, accent=False, panel=False):
    x, y, w, h = box
    if panel:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *map(Inches, box))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(theme["pale"])
        shape.line.fill.background()
        shape._element.spPr.append(OxmlElement("a:effectLst"))
    else:
        shape = slide.shapes.add_textbox(*map(Inches, box))
    shape.name = name
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(.16 if panel else .02)
    tf.margin_top = tf.margin_bottom = Inches(.12 if panel else .02)
    for i, line in enumerate(str(text).splitlines() or [""]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8 if h > 1 else 0)
        p.line_spacing = 1.15
        for run in p.runs:
            run.font.name, run.font.size = theme["font"], Pt(size)
            run.font.bold = accent
            run.font.color.rgb = RGBColor.from_string(theme["accent"] if accent else theme["ink"])
            ea = OxmlElement("a:ea")
            ea.set("typeface", theme["font"])
            run._r.get_or_add_rPr().append(ea)
    return shape


def add_slide(prs, layout: str, title: str, content: dict, *, theme="business", source=""):
    """Strict named slots avoid accidental placeholders; overflow is reported, not shrunk."""
    slots, colors = LAYOUTS[layout], THEMES[theme]
    if abs(prs.slide_width / Inches(1) - 13.333333) > .01 or abs(prs.slide_height / Inches(1) - 7.5) > .01:
        raise ValueError("版式库仅用于 16:9 标准画布；公司模板请使用模板跟随")
    expected = {s["name"] for s in slots}
    if set(content) != expected:
        raise ValueError(f"{layout} 内容字段应为 {sorted(expected)}")
    # Validate all slot data before appending a slide.
    for s in slots:
        value = content[s["name"]]
        if s["kind"] == "image" and not Path(value).is_file():
            raise ValueError(f"图片不存在: {value}")
        if s["kind"] == "table" and (not value or not value[0] or len(value) > 9 or any(len(row) != len(value[0]) for row in value)):
            raise ValueError("表格须为矩形，最多 9 行（含表头），过多请拆页")
        if s["kind"] in ("list", "process") and (not isinstance(value, list) or not 1 <= len(value) <= (5 if s["kind"] == "process" else 7)):
            raise ValueError("列表须为 1–7 项，流程须为 1–5 项")
        if s["kind"] == "chart":
            if not value.get("categories") or not value.get("series") or any(len(v) != len(value["categories"]) for v in value["series"].values()):
                raise ValueError("图表必须提供分类及长度一致的数值系列")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(colors["background"])
    _text(slide, "title", title, [.7, .5, 11.9, 1], colors, size=30, accent=True)
    for s in slots:
        name, kind, box = s["name"], s["kind"], s["box"]
        value = content[name]
        if kind == "image":
            from PIL import Image
            with Image.open(value) as img:
                ratio = img.width / img.height
            x, y, w, h = box
            iw, ih = min(w, h * ratio), min(h, w / ratio)
            shape = slide.shapes.add_picture(str(value), Inches(x + (w - iw) / 2), Inches(y + (h - ih) / 2), Inches(iw), Inches(ih))
            shape.name = name
        elif kind == "table":
            table_box = [*box[:3], min(box[3], .65 * len(value))]
            shape = slide.shapes.add_table(len(value), len(value[0]), *map(Inches, table_box))
            shape.name = name
            for r, row in enumerate(value):
                for c, text in enumerate(row):
                    cell = shape.table.cell(r, c)
                    cell.text = str(text)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(colors["accent"] if r == 0 else colors["pale"] if r % 2 else "FFFFFF")
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.name, run.font.size = colors["font"], Pt(16)
                            run.font.bold = r == 0
                            run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else colors["ink"])
                            ea = OxmlElement("a:ea")
                            ea.set("typeface", colors["font"])
                            run._r.get_or_add_rPr().append(ea)
        elif kind == "chart":
            data = CategoryChartData()
            data.categories = value["categories"]
            for label, numbers in value["series"].items():
                data.add_series(label, numbers)
            chart_type = XL_CHART_TYPE.LINE if value.get("type") == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
            shape = slide.shapes.add_chart(chart_type, *map(Inches, box), data)
            shape.name = name
            chart = shape.chart
            chart.font.name = colors["font"]
            chart.has_legend = len(value["series"]) > 1
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.font.size = Pt(12)
            chart.category_axis.tick_labels.font.size = Pt(13)
            chart.value_axis.tick_labels.font.size = Pt(12)
            if value.get("unit"):
                chart.value_axis.has_title = True
                chart.value_axis.axis_title.text_frame.text = value["unit"]
            if chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED and all(n >= 0 for values in value["series"].values() for n in values):
                chart.value_axis.minimum_scale = 0
            palette = [colors["accent"], "B16B36", "6B7387", "438F8B"]
            for index, series in enumerate(chart.series):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = RGBColor.from_string(palette[index % len(palette)])
                series.format.line.color.rgb = RGBColor.from_string(palette[index % len(palette)])
        elif kind == "process":
            x, y, w, h = box
            step_w = (w - .35 * (len(value) - 1)) / len(value)
            for index, text in enumerate(value):
                _text(slide, f"{name}-{index + 1}", f"{index + 1:02d}\n{text}", [x + index * (step_w + .35), y, step_w, h], colors, size=20, panel=True)
        else:
            text = "\n".join(f"{i + 1}. {item}" for i, item in enumerate(value)) if kind == "list" else value
            size = 36 if kind == "metric" else 26 if kind == "lead" else 18 if box[3] < 1 else 22
            _text(slide, name, text, box, colors, size=size, accent=kind == "metric", panel=kind == "panel")
    _text(slide, "page-number", str(len(prs.slides)), [12, 6.95, .6, .3], colors, size=10)
    if source:
        _text(slide, "source", source, [.7, 6.95, 10.8, .3], colors, size=10)
    return slide


def build_deck(plan: dict, output):
    """Materialize a traceable deck plan; content, native objects and notes share one source."""
    import json
    from .runtime import file_hash
    from .ppt_audit import audit_ppt
    if not all(str(plan.get(k, "")).strip() for k in ("audience", "purpose", "message")) or not plan.get("slides"):
        raise ValueError("计划须包含 audience、purpose、message 和 slides")
    prs = new_deck()
    for item in plan["slides"]:
        if not str(item.get("takeaway", "")).strip():
            raise ValueError("每页需要 takeaway，说明该页要让读者理解什么")
        slide = add_slide(prs, item["layout"], item["title"], item["content"],
                          theme=plan.get("theme", "business"), source=item.get("source", ""))
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = item["notes"]
    output = Path(output)
    prs.save(output)
    report = dict(plan=plan, output_sha256=file_hash(output), design_audit=audit_ppt(output))
    output.with_suffix(".design.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report
