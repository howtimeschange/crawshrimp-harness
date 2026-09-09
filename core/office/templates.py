"""Editable starter templates and a reproducible Chinese acceptance fixture."""
from pathlib import Path

FONT = "思源黑体"


def word_template(path: Path, title: str = "业务报告") -> None:
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for style in (doc.styles["Normal"], doc.styles["Title"], doc.styles["Heading 1"], doc.styles["Heading 2"]):
        style.font.name = FONT
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        for key in list(style.element.rPr.rFonts.attrib):
            if "theme" in key.lower():
                del style.element.rPr.rFonts.attrib[key]
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    doc.add_heading(title, 0)
    doc.add_paragraph("在此填写摘要与业务结论。")
    doc.add_heading("一、数据与发现", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ["指标", "本期", "说明"]):
        cell.text = value
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in [["销售额", "120,000.00", "示例数据"], ["退货率", "3.5%", "示例数据"]]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.save(path)


def ppt_template(path: Path, title: str = "业务汇报") -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    for heading, body in [(title, "结论 · 数据 · 下一步"), ("目录", "核心结论\n数据分析\n行动计划"),
                          ("核心结论", "在此填写有数据依据的发现"), ("数据分析", "在此插入原生图表或表格"),
                          ("行动计划", "行动 / 负责人 / 时间 / 验收标准")]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for x, y, w, h, text, size in [(0.65, 0.65, 12, 1, heading, 32), (0.7, 2, 11.8, 4.5, body, 22)]:
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            box.text_frame.word_wrap = True
            for i, line in enumerate(text.splitlines()):
                p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
                p.text = line
                for run in p.runs:
                    run.font.name, run.font.size = FONT, Pt(size)
                    run.font.color.rgb = RGBColor.from_string("183544")
                    ea = OxmlElement("a:ea")
                    ea.set("typeface", FONT)
                    run._r.get_or_add_rPr().append(ea)
    prs.save(path)


def excel_template(path: Path, title: str = "销售报表") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    for row in [["商品编号", "数量", "单价", "金额"], ["00123", 2, 19.5, "=B2*C2"], ["00456", 3, 10, "=B3*C3"]]:
        ws.append(row)
    summary = wb.create_sheet("汇总")
    summary.append([title, "数值"])
    summary.append(["销售额", "=SUM(数据!D2:D3)"])
    summary.append(["状态", '=IF(B2>0,"有销售","无销售")'])
    note = wb.create_sheet("说明")
    note.append(["口径", "示例数据；商品编号为文本；金额为数量乘单价。"])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = "1:1"
        sheet.print_area = sheet.dimensions
        for column in "ABCD":
            sheet.column_dimensions[column].width = 24
        for row in sheet:
            for cell in row:
                cell.font = Font(name=FONT, size=11, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "183544")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="183544")
        sheet.row_dimensions[1].height = 30
    ws.auto_filter.ref = ws.dimensions
    for row in (2, 3):
        ws[f"A{row}"].number_format = "@"
        ws[f"C{row}"].number_format = ws[f"D{row}"].number_format = "#,##0.00"
    summary["B2"].number_format = "#,##0.00"
    wb.save(path)


def acceptance_samples(output: Path) -> list[Path]:
    output.mkdir(parents=True, exist_ok=True)
    from core.office.runtime import configure_fonts
    configure_fonts()
    import matplotlib.pyplot as plt
    import pandas as pd
    data = pd.DataFrame({"月份": ["六月", "七月", "八月"], "销售额": [100, 120, 150]})
    fig, ax = plt.subplots(figsize=(7, 3))
    bars = ax.bar(data["月份"], data["销售额"], color="#236f85")
    ax.bar_label(bars, padding=3)
    ax.set_ylim(0, 180)
    ax.set_title("销售趋势（示例数据）")
    fig.tight_layout()
    image = output / "chart.png"
    fig.savefig(image, dpi=160)
    plt.close(fig)
    word, ppt, excel = output / "中文报告.docx", output / "中文汇报.pptx", output / "销售报表.xlsx"
    word_template(word)
    from docx import Document
    from docx.shared import Inches
    doc = Document(word)
    doc.add_page_break()
    doc.add_heading("二、趋势分析", 1)
    doc.add_picture(str(image), width=Inches(6))
    doc.add_paragraph("示例销售额由 100 增长至 150；金额 ¥69.00，负数 −1.25。")
    doc.add_page_break()
    doc.add_heading("三、行动计划", 1)
    doc.add_paragraph("2026-09-09：核对来源、负责人及下次复盘时间。")
    doc.save(word)
    ppt_template(ppt)
    from pptx import Presentation
    from pptx.util import Inches as PptInches
    prs = Presentation(ppt)
    slide = prs.slides[3]
    slide.shapes[1].text_frame.clear()
    slide.shapes.add_picture(str(image), PptInches(1), PptInches(2), width=PptInches(10))
    prs.save(ppt)
    excel_template(excel)
    return [word, ppt, excel]
