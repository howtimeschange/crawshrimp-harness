"""Native shape checks and advisory text-fit estimates; actual rendering stays required."""
from __future__ import annotations

import math
import unicodedata
from pptx import Presentation


def _frame_height(frame, width_pt):
    total = 0
    for p in frame.paragraphs:
        if not p.text:
            continue
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
        units = sum(1 if unicodedata.east_asian_width(c) in "WF" else .55 for c in p.text)
        lines = max(1, math.ceil(units * size / max(width_pt, 1)))
        spacing = p.line_spacing
        line_pt = spacing.pt if hasattr(spacing, "pt") else size * (spacing or 1.2)
        total += lines * line_pt + (p.space_after.pt if p.space_after else 0) + (p.space_before.pt if p.space_before else 0)
    return total


def audit_ppt(path):
    prs = Presentation(path)
    issues = []
    for page, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 100 or shape.top + shape.height > prs.slide_height + 100:
                issues.append(dict(code="PPT_OUTSIDE_CANVAS", page=page, shape=shape.name, severity="error"))
            frames = [(shape.name, shape.text_frame, shape.width, shape.height)] if shape.has_text_frame else []
            if shape.has_table:
                frames += [(f"{shape.name} R{r + 1}C{c + 1}", cell.text_frame, shape.table.columns[c].width, row.height)
                           for r, row in enumerate(shape.table.rows) for c, cell in enumerate(row.cells) if not cell.is_spanned]
            for name, frame, width, height in frames:
                available_w = (width - frame.margin_left - frame.margin_right) / 12700
                available_h = (height - frame.margin_top - frame.margin_bottom) / 12700
                if _frame_height(frame, available_w) > available_h + 3:
                    issues.append(dict(code="PPT_TEXT_FIT_ESTIMATE", page=page, shape=name, severity="warning",
                                       message="估算文本高度超过可用空间；请渲染确认，优先拆页或扩大区域"))
    return dict(status="issues" if issues else "passed", issues=issues, visual="not_run", estimator="advisory")
