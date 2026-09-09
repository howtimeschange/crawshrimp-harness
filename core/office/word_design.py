"""Harness-owned Word presets and native OOXML table geometry."""
from __future__ import annotations

from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Twips
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PRESETS = {
    "report": dict(label="商务报告", font="思源黑体", body=11, title=26, h1=18, h2=14,
                   accent="24546A", margin=2.2, line=1.4, after=7, header="业务报告"),
    "proposal": dict(label="项目方案", font="思源黑体", body=11, title=28, h1=19, h2=14,
                     accent="25635E", margin=2.3, line=1.45, after=8, header="项目方案"),
    "policy": dict(label="制度规范", font="思源宋体", body=12, title=24, h1=17, h2=14,
                   accent="243746", margin=2.5, line=1.5, after=6, header="制度规范"),
}


def presets() -> dict:
    return deepcopy(PRESETS)


def _set(parent, tag, **attrs):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    for key, value in attrs.items():
        child.set(qn(f"w:{key}"), str(value))
    return child


def set_font(font, rpr, family: str):
    font.name = family
    fonts = _set(rpr, "w:rFonts", eastAsia=family, ascii=family, hAnsi=family)
    for attr in list(fonts.attrib):
        if "theme" in attr.lower():
            del fonts.attrib[attr]


def apply_preset(doc, preset="report"):
    """Set document-level styles; do not discard paragraphs, fields or user content."""
    p = PRESETS[preset]
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(p["margin"])
        section.left_margin = section.right_margin = Cm(p["margin"])
        section.header_distance = section.footer_distance = Cm(1.1)
    for name, size in [("Normal", p["body"]), ("Title", p["title"]), ("Subtitle", 12),
                       ("Heading 1", p["h1"]), ("Heading 2", p["h2"]), ("Heading 3", 12),
                       ("Caption", 10), ("Header", 9), ("Footer", 9),
                       ("List Bullet", p["body"]), ("List Number", p["body"])]:
        style = doc.styles[name]
        heading = name.startswith("Heading") or name == "Title"
        set_font(style.font, style.element.get_or_add_rPr(), "思源黑体" if heading else p["font"])
        style.font.size = Pt(size)
        style.font.bold = heading
        style.font.color.rgb = RGBColor.from_string(p["accent"] if heading else "263340")
        fmt = style.paragraph_format
        fmt.line_spacing = 1.2 if heading else p["line"]
        fmt.space_before, fmt.space_after = Pt(14 if heading else 0), Pt(p["after"])
        fmt.keep_with_next = heading
        fmt.widow_control = True
        if name == "Title":
            borders = _set(style.element.get_or_add_pPr(), "w:pBdr")
            _set(borders, "w:bottom", val="single", sz=6, space=8, color=p["accent"])
    # The marker is opt-in: arbitrary imported files are not judged against a new preset.
    import re
    keywords = re.sub(r"\bharness-word:\w+\b", "", doc.core_properties.keywords or "").strip()
    doc.core_properties.keywords = f"{keywords} harness-word:{preset}".strip()
    return doc


def new_document(title: str, *, preset="report", subtitle="", metadata=None):
    doc = apply_preset(Document(), preset)
    p = PRESETS[preset]
    doc.sections[0].header.paragraphs[0].text = p["header"]
    doc.add_paragraph(title, "Title")
    if subtitle:
        doc.add_paragraph(subtitle, "Subtitle")
    if metadata:
        doc.add_paragraph("   ·   ".join(f"{k}：{v}" for k, v in metadata.items()), "Caption")
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def table_geometry(table, widths_cm, *, repeat_header=True):
    """Set table/grid/cell widths consistently, including horizontal merged cells."""
    widths = [int(Cm(float(w)).twips) for w in widths_cm]
    if len(widths) != len(table.columns) or any(w <= 0 for w in widths):
        raise ValueError("列宽必须为每个网格列提供一个正数")
    table.autofit = False
    props = table._tbl.tblPr
    _set(props, "w:tblW", type="dxa", w=sum(widths))
    _set(props, "w:tblInd", type="dxa", w=0)
    margins = _set(props, "w:tblCellMar")
    for edge in ("top", "bottom", "left", "right"):
        _set(margins, f"w:{edge}", type="dxa", w=100)
    for column, width in zip(table.columns, widths):
        column.width = Twips(width)
    for row in table._tbl.tr_lst:
        col = row.grid_before
        for cell in row.tc_lst:
            span = cell.grid_span
            _set(cell.get_or_add_tcPr(), "w:tcW", type="dxa", w=sum(widths[col:col + span]))
            col += span
    if repeat_header and table.rows:
        _set(table.rows[0]._tr.get_or_add_trPr(), "w:tblHeader", val="true")
    return table


def add_table(doc, headers, rows, *, widths_cm=None, preset="report"):
    if not headers or any(len(row) != len(headers) for row in rows):
        raise ValueError("表头和每行列数必须一致")
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
        _set(cell._tc.get_or_add_tcPr(), "w:shd", fill=PRESETS[preset]["accent"])
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = str(value)
    section = doc.sections[-1]
    width = (section.page_width - section.left_margin - section.right_margin) / 360000
    table_geometry(table, widths_cm or [width / len(headers)] * len(headers))
    return table
