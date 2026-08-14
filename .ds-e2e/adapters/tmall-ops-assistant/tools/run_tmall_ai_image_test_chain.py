#!/usr/bin/env python3
"""Run the Semir cloud -> 1XM image -> Tmall material-test chain.

The script keeps live Tmall mutations behind explicit flags. It reads 1XM keys
from Crawshrimp settings or environment variables, never from CLI args.
"""
from __future__ import annotations

import argparse
import asyncio
import base64
import html
import json
import math
import os
import re
import secrets
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping


REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from core import data_sink  # noqa: E402
from core.cdp_bridge import CDPBridge  # noqa: E402
from core.js_runner import JSRunner  # noqa: E402
from core.one_xm_image import file_to_data_url  # noqa: E402


ADAPTER_ID = "tmall-ops-assistant"
TASK_ID = "tmall_ai_image_test_chain"
SEMIR_URL = "https://fmp.semirapp.com/web/index#/home/file"
TMALL_URL = "https://myseller.taobao.com/home.htm/material-center/material-test/common_test?testStatus=1&testChannel=common_search"
DEFAULT_WORKFLOW_FILE = "/Users/xingyicheng/Downloads/测图工作流导入模板表.xlsx"
DEFAULT_PROMPT_FILE = "/Users/xingyicheng/Downloads/AI 测图提示词库.xlsx"
DEFAULT_CLOUD_PATH = "巴拉营运BU-商品//巴拉货控/02 产品上新模块/2-2 巴拉产品上新/2026年巴拉秋/"
PROMPT_SHEETS = {"上装", "连衣裙", "下装", "短裙", "鞋品", "竞品模版"}
IMAGE_EXT_RE = re.compile(r"\.(jpe?g|png|webp|bmp|gif|tiff?)($|\?)", re.IGNORECASE)
REFERENCE_IMAGE_MAX_BYTES = 19 * 1024 * 1024
RISKY_DEFAULT_PROMPT_RE = re.compile(r"多件衣服|衣服[1-9]|效果图|空字段|多件商品|多款", re.IGNORECASE)
TMALL_UPLOAD_IMAGE_SIZE = (960, 1280)
CHAIN_EXECUTION_MODES = ("approval_then_create", "direct_create")
APPROVAL_BATCH_FILENAME_PREFIX = "tmall-ai-image-approval-batch"
APPROVAL_BOARD_FILENAME_PREFIX = "tmall-ai-image-approval-board"
APPROVED_ASSET_STATUSES = {"approved", "selected", "confirmed", "通过", "已确认", "确认"}
GENERATION_CONFIRMATION_STATUS = "pending_generation_confirmation"
TMALL_LOGIN_READY_TIMEOUT_MS = 300_000
TMALL_READY_POLL_INTERVAL_SECONDS = 3.0
TMALL_AI_IMAGE_RESULT_COLUMN_ORDER = [
    "表格行号",
    "款号",
    "商品ID",
    "阶段",
    "参考图模式",
    "品类",
    "性别",
    "SKC编码",
    "云盘路径",
    "文件名",
    "主图原图文件",
    "平铺参考图路径",
    "平铺参考图文件名",
    "平铺参考图URL",
    "主参考图文件",
    "细节参考图文件",
    "提示词分组",
    "提示词字段名",
    "提示词序号",
    "尺寸",
    "格式",
    "质量",
    "参考图文件",
    "完整Prompt",
    "最终提示词",
    "1XM任务ID",
    "1XM轮询URL",
    "生成图URL",
    "生成图数量",
    "审批批次ID",
    "审批状态",
    "审批看板",
    "云端审批看板",
    "本地审批看板",
    "任务ID",
    "测图详情URL",
    "上传图数量",
    "提交图片数量",
    "提交图片文件",
    "提交模式",
    "上线结果",
    "页面回读",
    "素材URL",
    "本地文件",
    "执行结果",
    "备注",
]


def compact(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_generation_image_count(value: object, fallback: int = 1) -> int:
    count = to_int(value, fallback)
    return max(1, min(8, count))


def image_ratio_label(value: object) -> str:
    text = compact(value)
    if not text:
        return ""
    explicit = re.search(r"(\d+)\s*[:：]\s*(\d+)", text)
    if explicit:
        left = int(explicit.group(1))
        right = int(explicit.group(2))
        if left > 0 and right > 0:
            divisor = math.gcd(left, right)
            return f"{left // divisor}:{right // divisor}"
    size_match = re.search(r"(\d+)\s*[xX×*]\s*(\d+)", text)
    if size_match:
        width = int(size_match.group(1))
        height = int(size_match.group(2))
        if width > 0 and height > 0:
            divisor = math.gcd(width, height)
            return f"{width // divisor}:{height // divisor}"
    return ""


def generation_count_from_row(row: Mapping[str, Any], fallback: int = 1) -> int:
    payload = row.get("__1xm_payload") if isinstance(row.get("__1xm_payload"), Mapping) else {}
    return normalize_generation_image_count(row.get("生成数量") or payload.get("n"), fallback)


def tmall_material_test_detail_url(item_id: object) -> str:
    item = compact(item_id)
    if not item:
        return ""
    query = urllib.parse.urlencode({
        "testStatus": "1",
        "testChannel": "common_search",
        "tab": "all",
        "itemId": item,
    })
    return f"https://myseller.taobao.com/home.htm/material-center/material-test/common_test?{query}"


def normalize_chain_execution_mode(value: object) -> dict[str, bool | str]:
    mode = compact(value or "approval_then_create").lower()
    if mode not in CHAIN_EXECUTION_MODES:
        raise ValueError(f"天猫AI测图全链路只支持执行模式：{', '.join(CHAIN_EXECUTION_MODES)}")
    approval_required = mode == "approval_then_create"
    confirm_generation = True
    return {
        "mode": mode,
        "confirm_generation": confirm_generation,
        "generate": False,
        "approval_required": approval_required,
        "live_upload": not approval_required,
        "live_create": not approval_required,
        "live_online": False,
    }


def apply_chain_execution_mode(args: argparse.Namespace) -> argparse.Namespace:
    normalized = normalize_chain_execution_mode(getattr(args, "execute_mode", "approval_then_create"))
    args.execute_mode = str(normalized["mode"])
    args.confirm_generation = bool(normalized["confirm_generation"])
    args.generate = bool(normalized["generate"])
    args.approval_required = bool(normalized["approval_required"])
    args.live_upload = bool(normalized["live_upload"])
    args.live_create = bool(normalized["live_create"])
    args.live_online = bool(normalized["live_online"])
    return args


def normalize_key(value: object) -> str:
    return re.sub(r"[\s_./\\\-：:（）()]+", "", compact(value).lower())


def row_value(row: Mapping[str, Any], aliases: Iterable[str]) -> str:
    wanted = {normalize_key(alias) for alias in aliases}
    for key, value in (row or {}).items():
        if normalize_key(key) in wanted and compact(value):
            return compact(value)
    return ""


def parse_list(value: object) -> list[str]:
    if isinstance(value, list):
        return [compact(item) for item in value if compact(item)]
    return [item.strip() for item in re.split(r"[\n\r,，、；;]+", str(value or "")) if item.strip()]


def safe_token(value: object, fallback: str = "item") -> str:
    text = compact(value)
    token = re.sub(r"[^A-Za-z0-9._~-]+", "_", text).strip("_")
    return token or fallback


def stable_hash(value: object) -> str:
    text = str(value or "")
    h = 2166136261
    for char in text:
        h ^= ord(char)
        h = (h * 16777619) & 0xFFFFFFFF
    return f"{h:08x}"


def json_safe(value: Any) -> Any:
    return json.loads(json.dumps(value, ensure_ascii=False, default=str))


def workflow_to_dict(workflow: WorkflowItem | Mapping[str, Any]) -> dict[str, Any]:
    if isinstance(workflow, WorkflowItem):
        return {
            "row_no": workflow.row_no,
            "style_code": workflow.style_code,
            "item_id": workflow.item_id,
            "category": workflow.category,
            "gender": workflow.gender,
            "prompt_name": workflow.prompt_name,
            "skc_code": workflow.skc_code,
        }
    return {
        "row_no": int((workflow or {}).get("row_no") or (workflow or {}).get("表格行号") or 0),
        "style_code": compact((workflow or {}).get("style_code") or (workflow or {}).get("款号")),
        "item_id": compact((workflow or {}).get("item_id") or (workflow or {}).get("商品ID")),
        "category": compact((workflow or {}).get("category") or (workflow or {}).get("品类")),
        "gender": compact((workflow or {}).get("gender") or (workflow or {}).get("性别")),
        "prompt_name": compact((workflow or {}).get("prompt_name") or (workflow or {}).get("提示词字段名")),
        "skc_code": compact((workflow or {}).get("skc_code") or (workflow or {}).get("SKC编码")),
    }


def workflow_from_dict(value: Mapping[str, Any]) -> WorkflowItem:
    data = workflow_to_dict(value or {})
    return WorkflowItem(
        row_no=int(data.get("row_no") or 0),
        style_code=compact(data.get("style_code")),
        item_id=compact(data.get("item_id")),
        category=compact(data.get("category")),
        gender=compact(data.get("gender")),
        prompt_name=compact(data.get("prompt_name")),
        skc_code=compact(data.get("skc_code")),
    )


def approval_asset_id(style_code: str, kind: str, index: object, path: str) -> str:
    return "-".join([
        safe_token(style_code, "style"),
        safe_token(kind, "asset"),
        safe_token(index, "1"),
        stable_hash(path)[:8],
    ])


def fresh_approval_nonce(batch: Mapping[str, Any], purpose: str) -> str:
    batch_run_uid = compact(batch.get("task_run_uid"))
    batch_id = compact(batch.get("batch_id")) or "approval"
    return "-".join([
        safe_token(batch_run_uid or batch_id, "run"),
        safe_token(purpose, "manual"),
        datetime.now(timezone.utc).astimezone().strftime("%Y%m%d%H%M%S"),
        secrets.token_hex(3),
    ])


def approval_batch_run_uid(batch: Mapping[str, Any]) -> str:
    run_params = batch.get("run_params") if isinstance(batch.get("run_params"), Mapping) else {}
    return (
        compact(batch.get("task_run_uid"))
        or compact(run_params.get("task_run_uid"))
        or compact(run_params.get("__task_run_id"))
    )


def apply_fresh_approval_generation_key(
    generation_row: dict[str, Any],
    batch: Mapping[str, Any],
    purpose: str,
) -> str:
    nonce = fresh_approval_nonce(batch, purpose)
    batch_run_uid = approval_batch_run_uid(batch)
    key_parts = [
        "tmall_ai_chain",
        safe_token(batch_run_uid or nonce, "run"),
        safe_token(generation_row.get("款号")),
        safe_token(generation_row.get("提示词分组") or "approval"),
        safe_token(generation_row.get("提示词字段名") or purpose),
        safe_token(generation_row.get("提示词序号") or 1),
        safe_token(generation_row.get("尺寸") or ""),
        stable_hash(generation_row.get("完整Prompt") or generation_row.get("最终提示词") or ""),
        safe_token(nonce),
    ]
    generation_row["__1xm_idempotency_key"] = "_".join(key_parts)[:180]
    generation_row["__1xm_output_token"] = safe_token(nonce)
    if batch_run_uid:
        generation_row["__task_run_uid"] = batch_run_uid
    return nonce


def _image_asset(path: str, *, style_code: str, kind: str, label: str, index: int = 0, status: str = "reference", extra: Mapping[str, Any] | None = None) -> dict[str, Any]:
    clean_path = compact(path)
    filename = Path(clean_path).name if clean_path else ""
    return {
        "id": approval_asset_id(style_code, kind, index, clean_path or label),
        "kind": kind,
        "label": label,
        "path": clean_path,
        "filename": filename,
        "status": status,
        **dict(extra or {}),
    }


def build_approval_item(
    workflow: WorkflowItem,
    semir_data: Mapping[str, Any],
    main_origin_path: str,
    detail_reference_path: str,
    generation_rows: list[Mapping[str, Any]],
    generated_paths: list[str],
    *,
    reference_mode: str = "main_only",
) -> dict[str, Any]:
    workflow_data = workflow_to_dict(workflow)
    assets: list[dict[str, Any]] = []
    if compact(main_origin_path):
        assets.append(_image_asset(
            main_origin_path,
            style_code=workflow.style_code,
            kind="origin",
            label="原图/主图",
            index=1,
            status="reference",
            extra={"source_label": "原图/主图", "slot": "main", "use_for_generation": True},
        ))
    if compact(detail_reference_path):
        detail_paths = parse_list((semir_data or {}).get("detailReferencePaths"))
        if detail_reference_path not in detail_paths:
            detail_paths.insert(0, detail_reference_path)
        for detail_index, path in enumerate([item for item in detail_paths if compact(item)], start=1):
            assets.append(_image_asset(
                path,
                style_code=workflow.style_code,
                kind="detail_reference",
                label="款色参考图" if detail_index == 1 else f"款色参考图 {detail_index}",
                index=detail_index,
                status="reference",
                extra={"source_label": "款色参考图", "slot": "reference", "use_for_generation": False},
            ))

    selected_paths = [compact(path) for path in (generated_paths or []) if compact(path)]
    selected_path_set = set(selected_paths)
    fallback_paths = list(selected_paths)
    item_task_run_uid = ""
    for row_index, generation_row in enumerate(generation_rows or [], start=1):
        if not item_task_run_uid:
            item_task_run_uid = compact(generation_row.get("__task_run_uid"))
        paths = parse_list(generation_row.get("本地生成图文件"))
        if selected_path_set:
            paths = [path for path in paths if compact(path) in selected_path_set]
        if not paths and row_index <= len(fallback_paths):
            paths = [fallback_paths[row_index - 1]]
        for image_offset, path in enumerate(paths, start=1):
            prompt_index = int(generation_row.get("提示词序号") or row_index)
            display_index = prompt_index if len(paths) == 1 else f"{prompt_index}-{image_offset}"
            assets.append(_image_asset(
                path,
                style_code=workflow.style_code,
                kind="ai",
                label=f"AI 图 {display_index}",
                index=display_index,
                status="pending",
                extra={
                    "prompt_index": prompt_index,
                    "prompt_name": compact(generation_row.get("提示词字段名")),
                    "prompt_group": compact(generation_row.get("提示词分组")),
                    "prompt": compact(generation_row.get("完整Prompt") or generation_row.get("最终提示词")),
                    "generation_row": json_safe(generation_row),
                    "reference_paths": parse_list(generation_row.get("参考图文件")),
                    "reference_labels": [Path(path).name for path in parse_list(generation_row.get("参考图文件"))],
                    "task_run_uid": compact(generation_row.get("__task_run_uid")),
                    "custom_prompt": "",
                    "review_note": "",
                },
            ))

    return {
        "id": safe_token(f"{workflow.style_code}-{workflow.row_no}", "item"),
        "row_no": workflow.row_no,
        "style_code": workflow.style_code,
        "item_id": workflow.item_id,
        "category": workflow.category,
        "gender": workflow.gender,
        "skc_code": workflow.skc_code,
        "reference_mode": reference_mode,
        "origin_path": compact(main_origin_path),
        "detail_reference_path": compact(detail_reference_path),
        "workflow": workflow_data,
        "semir_data": json_safe(semir_data or {}),
        "task_run_uid": item_task_run_uid,
        "assets": assets,
    }


def _generation_prompt_id(style_code: str, index: object, prompt: object) -> str:
    return "-".join([
        safe_token(style_code, "style"),
        "prompt",
        safe_token(index, "1"),
        stable_hash(prompt)[:8],
    ])


def generation_prompt_from_row(
    workflow: WorkflowItem,
    generation_row: Mapping[str, Any],
    *,
    index: int,
) -> dict[str, Any]:
    prompt_text = compact(generation_row.get("完整Prompt") or generation_row.get("最终提示词"))
    reference_paths = parse_list(generation_row.get("参考图文件") or generation_row.get("__1xm_reference_paths"))
    image_count = generation_count_from_row(generation_row, 1)
    return {
        "id": _generation_prompt_id(workflow.style_code, generation_row.get("提示词序号") or index, prompt_text),
        "prompt_index": int(generation_row.get("提示词序号") or index),
        "prompt_name": compact(generation_row.get("提示词字段名")) or f"Prompt {index}",
        "prompt_group": compact(generation_row.get("提示词分组")),
        "prompt": prompt_text,
        "custom_prompt": "",
        "original_content": prompt_text,
        "modified_in_batch": False,
        "image_count": image_count,
        "reference_paths": reference_paths,
        "reference_labels": [Path(path).name for path in reference_paths],
        "reference_binding_mode": "automatic",
        "use_custom_references": False,
        "status": "pending",
        "generation_row": json_safe(generation_row),
    }


def build_generation_confirmation_item(
    workflow: WorkflowItem,
    semir_data: Mapping[str, Any],
    main_origin_path: str,
    detail_reference_path: str,
    generation_rows: list[Mapping[str, Any]],
    *,
    reference_mode: str = "main_only",
) -> dict[str, Any]:
    item = build_approval_item(
        workflow,
        semir_data,
        main_origin_path,
        detail_reference_path,
        generation_rows,
        [],
        reference_mode=reference_mode,
    )
    item["status"] = GENERATION_CONFIRMATION_STATUS
    item["generation_rows"] = json_safe(list(generation_rows or []))
    item["generation_prompts"] = [
        generation_prompt_from_row(workflow, row, index=index)
        for index, row in enumerate(generation_rows or [], start=1)
    ]
    return item


def generation_confirmation_prompt_total(items: Iterable[Mapping[str, Any]]) -> int:
    total = 0
    for item in items or []:
        total += len([
            prompt
            for prompt in item.get("generation_prompts") or []
            if compact(prompt.get("status")).lower() not in {"removed", "deleted", "disabled", "skip", "skipped"}
        ])
    return total


def write_generation_confirmation_batch(
    artifact_dir: Path,
    items: list[Mapping[str, Any]],
    *,
    run_params: Mapping[str, Any] | None = None,
    base_url: str = "",
    batch_id: str = "",
    token: str = "",
    created_at: str = "",
) -> dict[str, Any]:
    batch = write_approval_batch(
        artifact_dir,
        items,
        run_params=run_params,
        base_url=base_url,
        batch_id=batch_id,
        token=token,
        created_at=created_at,
    )
    batch["status"] = GENERATION_CONFIRMATION_STATUS
    batch["generation_status"] = "pending"
    batch["approval_status"] = "not_required_yet"
    batch["test_task_status"] = "not_started"
    batch["generation_prompt_total"] = generation_confirmation_prompt_total(batch.get("items") or [])
    batch["approval_message"] = (
        f"天猫AI测图待确认提交生图批次 {batch.get('batch_id')} 已准备完成，"
        f"共 {len(batch.get('items') or [])} 款、{batch.get('generation_prompt_total') or 0} 条 Prompt。"
        f"\n请打开确认提交看板检查素材和 Prompt：{batch.get('board_url')}"
    )
    save_approval_batch(batch)
    return batch


def generation_confirmation_result_rows(batch: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for item in batch.get("items") or []:
        prompt_total = len(item.get("generation_prompts") or [])
        rows.append({
            "表格行号": item.get("row_no", ""),
            "款号": item.get("style_code", ""),
            "商品ID": item.get("item_id", ""),
            "阶段": "确认提交生图",
            "审批批次ID": batch.get("batch_id", ""),
            "审批状态": batch.get("status", ""),
            "审批看板": batch.get("board_url", ""),
            "云端审批看板": batch.get("cloud_board_url", ""),
            "本地审批看板": batch.get("local_board_url", ""),
            "执行结果": "等待确认提交生图任务",
            "备注": f"Prompt={prompt_total}；{batch.get('approval_message', '')}",
        })
    return rows


def render_approval_message_template(template: str, batch: Mapping[str, Any]) -> str:
    item_count = len(batch.get("items") or [])
    ai_count = sum(
        1
        for item in batch.get("items") or []
        for asset in item.get("assets") or []
        if asset.get("kind") == "ai"
    )
    values = {
        "batch_id": compact(batch.get("batch_id")),
        "board_url": compact(batch.get("board_url")),
        "total_styles": str(item_count),
        "total_ai_images": str(ai_count),
        "created_at": compact(batch.get("created_at")),
    }
    text = compact(template) or (
        "天猫AI测图生图批次 {{batch_id}} 已完成，共 {{total_styles}} 款、{{total_ai_images}} 张 AI 图。"
        "\n请打开审批看板确认后创建测图任务：{{board_url}}"
    )
    for key, value in values.items():
        text = text.replace("{{" + key + "}}", value)
    return text


def approval_base_url(value: object = "") -> str:
    explicit = compact(value)
    if explicit:
        return explicit.rstrip("/")
    port = compact(os.environ.get("CRAWSHRIMP_PORT")) or "18765"
    return f"http://127.0.0.1:{port}"


def cloud_approval_board_url(base_url: object, batch_id: object) -> str:
    query = urllib.parse.urlencode({"batch_uid": compact(batch_id)})
    return f"{approval_base_url(base_url)}/?{query}"


def render_approval_board_html(batch: Mapping[str, Any]) -> str:
    payload = (
        json.dumps(json_safe(batch), ensure_ascii=False)
        .replace("&", "\\u0026")
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
        .replace("\u2028", "\\u2028")
        .replace("\u2029", "\\u2029")
    )
    is_generation_confirmation = compact(batch.get("status")) == GENERATION_CONFIRMATION_STATUS
    title = "天猫AI测图确认提交看板" if is_generation_confirmation else "天猫AI测图审批看板"
    submit_label = "确认提交生图任务" if is_generation_confirmation else "提交已确认图片并创建测图任务"
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{title}</title>
  <style>
    :root {{ color-scheme: light; --ink:#172026; --muted:#66747f; --line:#d8e0e6; --bg:#f6f8fa; --panel:#ffffff; --accent:#0f766e; --danger:#b42318; --warn:#a15c00; }}
    * {{ box-sizing:border-box; }}
    body {{ margin:0; font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif; color:var(--ink); background:var(--bg); }}
    header {{ position:sticky; top:0; z-index:2; padding:18px 24px; background:rgba(255,255,255,.94); border-bottom:1px solid var(--line); backdrop-filter:blur(12px); }}
    h1 {{ margin:0 0 6px; font-size:22px; font-weight:680; letter-spacing:0; }}
    .meta {{ display:flex; flex-wrap:wrap; gap:10px 18px; color:var(--muted); font-size:13px; }}
    main {{ padding:20px 24px 36px; }}
    .toolbar {{ display:flex; gap:10px; flex-wrap:wrap; align-items:center; margin-bottom:16px; }}
    button {{ border:1px solid var(--line); background:#fff; color:var(--ink); border-radius:7px; padding:8px 12px; cursor:pointer; font-size:13px; }}
    button.primary {{ background:var(--accent); color:#fff; border-color:var(--accent); }}
    button.danger {{ color:var(--danger); }}
    button:disabled {{ opacity:.55; cursor:not-allowed; }}
    .item {{ background:var(--panel); border:1px solid var(--line); border-radius:8px; margin:0 0 18px; overflow:hidden; }}
    .item-head {{ display:flex; justify-content:space-between; gap:12px; padding:14px 16px; border-bottom:1px solid var(--line); }}
    .item-title {{ font-size:16px; font-weight:650; }}
    .item-sub {{ color:var(--muted); font-size:12px; margin-top:3px; }}
    .grid {{ display:grid; grid-template-columns:repeat(auto-fill,minmax(190px,1fr)); gap:12px; padding:14px; }}
    .tile {{ border:1px solid var(--line); border-radius:8px; background:#fff; overflow:hidden; min-width:0; }}
    .tile.reference {{ background:#f8fbfb; }}
    .tile.prompt-card {{ background:#fffaf7; }}
    .thumb {{ width:100%; aspect-ratio:3/4; object-fit:cover; background:#eef3f6; display:block; }}
    .tile-body {{ padding:10px; }}
    .label {{ font-size:13px; font-weight:650; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }}
    .path {{ color:var(--muted); font-size:11px; margin-top:4px; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }}
    .actions {{ display:flex; gap:6px; flex-wrap:wrap; margin-top:8px; }}
    .status {{ font-size:12px; color:var(--muted); margin-top:7px; }}
    .status.approved {{ color:var(--accent); }}
    .status.rejected {{ color:var(--danger); }}
    textarea,input {{ width:100%; border:1px solid var(--line); border-radius:7px; padding:8px; font:12px/1.4 ui-monospace,SFMono-Regular,Menlo,monospace; }}
    textarea {{ min-height:86px; resize:vertical; margin-top:8px; }}
    .refs {{ margin-top:8px; display:grid; gap:6px; }}
    .count-field {{ display:grid; gap:5px; margin-top:8px; color:var(--muted); font-size:12px; }}
    .count-field input {{ max-width:110px; }}
    dialog {{ border:0; border-radius:8px; max-width:min(900px,92vw); width:900px; padding:0; box-shadow:0 24px 90px rgba(20,28,35,.25); }}
    dialog::backdrop {{ background:rgba(15,23,42,.35); }}
    .modal-head {{ padding:14px 16px; border-bottom:1px solid var(--line); display:flex; justify-content:space-between; gap:12px; }}
    .modal-body {{ padding:16px; display:grid; grid-template-columns:minmax(260px,360px) 1fr; gap:16px; }}
    .modal-body img {{ width:100%; border-radius:8px; border:1px solid var(--line); }}
    pre {{ white-space:pre-wrap; word-break:break-word; background:#f5f7f9; border:1px solid var(--line); border-radius:8px; padding:12px; margin:0; max-height:520px; overflow:auto; }}
    @media (max-width:700px) {{ main,header {{ padding-left:14px; padding-right:14px; }} .modal-body {{ grid-template-columns:1fr; }} }}
  </style>
</head>
<body>
  <header>
    <h1>{title}</h1>
    <div class="meta">
      <span>批次：{html.escape(compact(batch.get("batch_id")))}</span>
      <span>状态：<strong id="batch-status">{html.escape(compact(batch.get("status")))}</strong></span>
      <span>创建时间：{html.escape(compact(batch.get("created_at")))}</span>
    </div>
  </header>
  <main>
    <div class="toolbar">
      <button class="primary" id="submit">{submit_label}</button>
      <button id="save">保存审批状态</button>
      <span id="message" class="meta"></span>
    </div>
    <section id="items"></section>
  </main>
  <dialog id="detail">
    <div class="modal-head"><strong id="detail-title"></strong><button id="close-detail">关闭</button></div>
    <div class="modal-body"><img id="detail-image" alt="" /><pre id="detail-prompt"></pre></div>
  </dialog>
  <script>
    const batch = {payload};
    const isGenerationConfirmation = batch.status === '{GENERATION_CONFIRMATION_STATUS}';
    const token = new URLSearchParams(location.search).get('token') || batch.token || '';
    const api = (suffix) => `/tmall-ai-image-approval/api/${{encodeURIComponent(batch.batch_id)}}${{suffix}}?token=${{encodeURIComponent(token)}}`;
    const imageUrl = (asset) => `/tmall-ai-image-approval/api/${{encodeURIComponent(batch.batch_id)}}/image/${{encodeURIComponent(asset.id)}}?token=${{encodeURIComponent(token)}}`;
    const state = new Map();
    function setMessage(text, danger=false) {{
      const el = document.getElementById('message');
      el.textContent = text || '';
      el.style.color = danger ? 'var(--danger)' : 'var(--muted)';
    }}
    function allAssets() {{
      return batch.items.flatMap((item) => (item.assets || []).map((asset) => [item, asset]));
    }}
    function escapeHtml(value) {{
      return String(value || '').replace(/[&<>"']/g, (char) => ({{ '&':'&amp;', '<':'&lt;', '>':'&gt;', '"':'&quot;', "'":'&#39;' }}[char]));
    }}
    function normalizeImageCount(value) {{
      const count = Number.parseInt(String(value || '1'), 10);
      if (!Number.isFinite(count)) return 1;
      return Math.min(8, Math.max(1, count));
    }}
    function safeClassName(value) {{
      return String(value || '').replace(/[^a-zA-Z0-9_-]+/g, '-');
    }}
    function updateAsset(asset, patch) {{
      Object.assign(asset, patch);
      state.set(asset.id, {{ status: asset.status, custom_prompt: asset.custom_prompt || '', reference_paths: asset.reference_paths || [], review_note: asset.review_note || '' }});
      render();
    }}
    function openDetail(asset) {{
      document.getElementById('detail-title').textContent = asset.label + ' / ' + (asset.prompt_name || asset.filename || '');
      document.getElementById('detail-image').src = imageUrl(asset);
      document.getElementById('detail-prompt').textContent = asset.custom_prompt || asset.prompt || '无 prompt';
      document.getElementById('detail').showModal();
    }}
    function render() {{
      const root = document.getElementById('items');
      root.innerHTML = '';
      for (const item of batch.items || []) {{
        const section = document.createElement('article');
        section.className = 'item';
        section.innerHTML = `<div class="item-head"><div><div class="item-title">${{escapeHtml(item.style_code || '')}}</div><div class="item-sub">商品ID ${{escapeHtml(item.item_id || '-')}} · ${{escapeHtml(item.category || '-')}} · SKC ${{escapeHtml(item.skc_code || '-')}}</div></div><div class="item-sub">参考图模式 ${{escapeHtml(item.reference_mode || '')}}</div></div><div class="grid"></div>`;
        const grid = section.querySelector('.grid');
        for (const asset of item.assets || []) {{
          const tile = document.createElement('div');
          tile.className = 'tile ' + (asset.kind === 'ai' ? '' : 'reference');
          const canApprove = asset.kind === 'ai';
          const refsValue = escapeHtml((asset.reference_paths || []).join('\\n'));
          const statusText = asset.status || 'pending';
          tile.innerHTML = `
            <img class="thumb" src="${{imageUrl(asset)}}" alt="${{escapeHtml(asset.label || '')}}" />
            <div class="tile-body">
              <div class="label" title="${{escapeHtml(asset.label || '')}}">${{escapeHtml(asset.label || '')}}</div>
              <div class="path" title="${{escapeHtml(asset.path || '')}}">${{escapeHtml(asset.filename || asset.path || '')}}</div>
              <div class="actions">
                <button data-act="detail">Prompt</button>
                ${{canApprove ? '<button data-act="approve">确认</button><button class="danger" data-act="reject">舍弃</button><button data-act="retry">重试/改图</button>' : ''}}
              </div>
              ${{canApprove ? `<textarea data-field="prompt" placeholder="修改生图 prompt 后点击重试">${{escapeHtml(asset.custom_prompt || asset.prompt || '')}}</textarea><div class="refs"><input data-field="refs" placeholder="参考图路径，可多条用逗号/换行分隔" value="${{refsValue}}" /></div>` : ''}}
              <div class="status ${{safeClassName(statusText)}}">状态：${{escapeHtml(statusText)}}</div>
            </div>`;
          tile.querySelector('[data-act="detail"]').onclick = () => openDetail(asset);
          const promptInput = tile.querySelector('[data-field="prompt"]');
          if (promptInput) promptInput.oninput = () => {{ asset.custom_prompt = promptInput.value; }};
          const refsInput = tile.querySelector('[data-field="refs"]');
          if (refsInput) refsInput.oninput = () => {{ asset.reference_paths = refsInput.value.split(/[\\n,，、；;]+/).map((v) => v.trim()).filter(Boolean); }};
          const approve = tile.querySelector('[data-act="approve"]');
          if (approve) approve.onclick = () => updateAsset(asset, {{ status: 'approved' }});
          const reject = tile.querySelector('[data-act="reject"]');
          if (reject) reject.onclick = () => updateAsset(asset, {{ status: 'rejected' }});
          const retry = tile.querySelector('[data-act="retry"]');
          if (retry) retry.onclick = async () => {{
            setMessage('正在重新生图...');
            const res = await fetch(api('/regenerate'), {{ method:'POST', headers:{{'Content-Type':'application/json'}}, body:JSON.stringify({{ asset_id: asset.id, prompt: asset.custom_prompt || asset.prompt || '', reference_paths: asset.reference_paths || [] }}) }});
            const payload = await res.json();
            if (!res.ok || payload.error) return setMessage(payload.error || '重新生图失败', true);
            Object.assign(asset, payload.asset || {{}});
            setMessage('重新生图完成');
            render();
          }};
          grid.appendChild(tile);
        }}
        if (isGenerationConfirmation) {{
          for (const prompt of item.generation_prompts || []) {{
            const tile = document.createElement('div');
            tile.className = 'tile prompt-card';
            const refsValue = escapeHtml((prompt.reference_paths || []).join('\\n'));
            tile.innerHTML = `
              <div class="tile-body">
                <div class="label" title="${{escapeHtml(prompt.prompt_name || '')}}">${{escapeHtml(prompt.prompt_name || 'Prompt')}}</div>
                <div class="path">待提交生图 · 序号 ${{escapeHtml(prompt.prompt_index || '')}}</div>
                <textarea data-field="prompt" placeholder="确认或修改本条生图 Prompt">${{escapeHtml(prompt.custom_prompt || prompt.prompt || '')}}</textarea>
                <label class="count-field">生图张数<input data-field="image-count" type="number" min="1" max="8" step="1" value="${{escapeHtml(normalizeImageCount(prompt.image_count || 1))}}" /></label>
                <div class="refs"><input data-field="refs" placeholder="参考图路径，可多条用逗号/换行分隔" value="${{refsValue}}" /></div>
                <div class="actions"><button class="danger" data-act="remove">删除本条 Prompt</button></div>
                <div class="status ${{safeClassName(prompt.status || 'pending')}}">状态：${{escapeHtml(prompt.status || 'pending')}}</div>
              </div>`;
            const promptInput = tile.querySelector('[data-field="prompt"]');
            if (promptInput) promptInput.oninput = () => {{ prompt.custom_prompt = promptInput.value; }};
            const imageCountInput = tile.querySelector('[data-field="image-count"]');
            if (imageCountInput) imageCountInput.oninput = () => {{ prompt.image_count = normalizeImageCount(imageCountInput.value); imageCountInput.value = String(prompt.image_count); }};
            const refsInput = tile.querySelector('[data-field="refs"]');
            if (refsInput) refsInput.oninput = () => {{ prompt.reference_paths = refsInput.value.split(/[\\n,，、；;]+/).map((v) => v.trim()).filter(Boolean); }};
            const remove = tile.querySelector('[data-act="remove"]');
            if (remove) remove.onclick = () => {{ prompt.status = 'removed'; render(); }};
            if (prompt.status !== 'removed') grid.appendChild(tile);
          }}
          const add = document.createElement('button');
          add.type = 'button';
          add.className = 'tile prompt-card';
          add.innerHTML = '<div class="tile-body"><div class="label">+ 新增 Prompt</div><div class="path">增加一张待生成 AI 图</div></div>';
          add.onclick = () => {{
            item.generation_prompts = item.generation_prompts || [];
            const nextIndex = item.generation_prompts.length + 1;
            item.generation_prompts.push({{ id: `${{item.style_code || 'style'}}-manual-${{Date.now()}}`, prompt_index: nextIndex, prompt_name: `新增 Prompt ${{nextIndex}}`, prompt: '', custom_prompt: '', image_count: 1, reference_paths: (item.assets || []).filter((asset) => asset.kind !== 'ai' && asset.path).map((asset) => asset.path), status: 'pending' }});
            render();
          }};
          grid.appendChild(add);
        }}
        root.appendChild(section);
      }}
    }}
    async function saveDecisions() {{
      const decisions = Object.fromEntries(allAssets().map(([, asset]) => [asset.id, {{ status: asset.status, custom_prompt: asset.custom_prompt || '', reference_paths: asset.reference_paths || [], review_note: asset.review_note || '' }}]));
      const res = await fetch(api('/decisions'), {{ method:'POST', headers:{{'Content-Type':'application/json'}}, body:JSON.stringify({{ decisions }}) }});
      const payload = await res.json();
      if (!res.ok || payload.error) return setMessage(payload.error || '保存失败', true);
      setMessage('审批状态已保存');
    }}
    document.getElementById('save').onclick = saveDecisions;
    document.getElementById('submit').onclick = async () => {{
      if (isGenerationConfirmation) {{
        if (!confirm('确认提交当前主图、参考图和 Prompt，开始批量生图？')) return;
        setMessage('正在提交批量生图任务...');
        const items = (batch.items || []).map((item) => ({{ id: item.id, style_code: item.style_code, origin_path: item.origin_path, detail_reference_path: item.detail_reference_path, generation_prompts: item.generation_prompts || [] }}));
        const res = await fetch(api('/submit-generation'), {{ method:'POST', headers:{{'Content-Type':'application/json'}}, body:JSON.stringify({{ items }}) }});
        const payload = await res.json();
        if (!res.ok || payload.error) return setMessage(payload.error || '提交生图失败', true);
        Object.assign(batch, payload.batch || {{}});
        document.getElementById('batch-status').textContent = batch.status || payload.status || 'pending_approval';
        setMessage(`生图完成：${{payload.generated || 0}} 张`);
        render();
        return;
      }}
      await saveDecisions();
      if (!confirm('确认只上传已确认的 AI 图并创建天猫测图任务？')) return;
      setMessage('正在上传并创建测图任务...');
      const res = await fetch(api('/submit'), {{ method:'POST' }});
      const payload = await res.json();
      if (!res.ok || payload.error) return setMessage(payload.error || '提交失败', true);
      document.getElementById('batch-status').textContent = payload.status || 'submitted';
      setMessage(`提交完成：${{payload.submitted || 0}} 款`);
    }};
    document.getElementById('close-detail').onclick = () => document.getElementById('detail').close();
    render();
  </script>
</body>
</html>"""


def write_approval_batch(
    artifact_dir: Path,
    items: list[Mapping[str, Any]],
    *,
    run_params: Mapping[str, Any] | None = None,
    base_url: str = "",
    batch_id: str = "",
    token: str = "",
    created_at: str = "",
) -> dict[str, Any]:
    artifact_dir.mkdir(parents=True, exist_ok=True)
    params = json_safe(dict(run_params or {}))
    execution_mode = str(normalize_chain_execution_mode(params.get("execute_mode"))["mode"])
    batch_id = safe_token(batch_id or f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-{secrets.token_hex(3)}", "batch")
    token = compact(token) or secrets.token_urlsafe(24)
    board_url = f"{approval_base_url(base_url)}/tmall-ai-image-approval/{batch_id}?token={token}"
    batch = {
        "schema_version": 2,
        "batch_id": batch_id,
        "token": token,
        "status": "pending_approval",
        "created_at": created_at or datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "adapter_id": ADAPTER_ID,
        "task_id": TASK_ID,
        "task_instance_uid": compact((run_params or {}).get("task_instance_uid")),
        "task_run_uid": compact((run_params or {}).get("task_run_uid")),
        "execution_mode": execution_mode,
        "cloud_prompt_library": {
            "id": compact(params.get("cloud_prompt_library_id")),
            "name": compact(params.get("cloud_prompt_library_name")),
            "source": compact(params.get("cloud_prompt_library_source")) or (
                "cloud" if compact(params.get("cloud_prompt_library_id")) else ""
            ),
        },
        "generation_status": "succeeded",
        "approval_status": "pending" if execution_mode == "approval_then_create" else "not_required",
        "test_task_status": "not_started",
        "artifact_dir": str(artifact_dir),
        "board_url": board_url,
        "items": json_safe(list(items or [])),
        "run_params": params,
    }
    batch["approval_message"] = render_approval_message_template(
        str((run_params or {}).get("approval_message_template") or ""),
        batch,
    )
    json_path = artifact_dir / f"{APPROVAL_BATCH_FILENAME_PREFIX}-{batch_id}.json"
    board_path = artifact_dir / f"{APPROVAL_BOARD_FILENAME_PREFIX}-{batch_id}.html"
    batch["json_path"] = str(json_path)
    batch["board_path"] = str(board_path)
    json_path.write_text(json.dumps(batch, ensure_ascii=False, indent=2), encoding="utf-8")
    board_path.write_text(render_approval_board_html(batch), encoding="utf-8")
    return batch


def save_approval_batch(batch: Mapping[str, Any]) -> None:
    json_path_raw = compact(batch.get("json_path"))
    if not json_path_raw:
        return
    json_path = Path(json_path_raw)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(json_safe(batch), ensure_ascii=False, indent=2), encoding="utf-8")
    board_path = Path(compact(batch.get("board_path")))
    if board_path:
        board_path.write_text(render_approval_board_html(batch), encoding="utf-8")


def update_approval_decisions(batch: dict[str, Any], decisions: Mapping[str, Any]) -> dict[str, Any]:
    decision_map = decisions or {}
    for item in batch.get("items") or []:
        for asset in item.get("assets") or []:
            patch = decision_map.get(asset.get("id"))
            if not isinstance(patch, Mapping):
                continue
            if compact(patch.get("status")):
                asset["status"] = compact(patch.get("status"))
            if "custom_prompt" in patch:
                asset["custom_prompt"] = compact(patch.get("custom_prompt"))
            if "reference_paths" in patch:
                asset["reference_paths"] = parse_list(patch.get("reference_paths"))
            if "review_note" in patch:
                asset["review_note"] = compact(patch.get("review_note"))
    batch["updated_at"] = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    save_approval_batch(batch)
    return batch


def _find_confirmation_item(batch: Mapping[str, Any], patch: Mapping[str, Any]) -> dict[str, Any] | None:
    patch_id = compact(patch.get("id"))
    patch_style = compact(patch.get("style_code"))
    patch_item_id = compact(patch.get("item_id"))
    for item in batch.get("items") or []:
        if not isinstance(item, dict):
            continue
        if patch_id and patch_id == compact(item.get("id")):
            return item
        if patch_style and patch_style == compact(item.get("style_code")):
            return item
        if patch_item_id and patch_item_id == compact(item.get("item_id")):
            return item
    return None


def _upsert_reference_asset(item: dict[str, Any], *, kind: str, label: str, path: str) -> None:
    clean_path = compact(path)
    assets = item.setdefault("assets", [])
    existing = next((asset for asset in assets if asset.get("kind") == kind), None)
    if not clean_path:
        item_key = "origin_path" if kind == "origin" else "detail_reference_path"
        item[item_key] = ""
        item["assets"] = [asset for asset in assets if asset.get("kind") != kind]
        return
    asset = _image_asset(
        clean_path,
        style_code=compact(item.get("style_code")),
        kind=kind,
        label=label,
        index=1,
        status="reference",
        extra={"source_label": label},
    )
    if existing:
        existing.update(asset)
    else:
        assets.insert(0 if kind == "origin" else min(1, len(assets)), asset)


def _sync_confirmation_reference_assets(item: dict[str, Any], reference_assets: Iterable[Mapping[str, Any]]) -> None:
    next_assets = [
        asset
        for asset in item.get("assets") or []
        if isinstance(asset, Mapping) and (asset.get("kind") == "origin" or asset.get("kind") == "ai" or asset.get("slot") == "main")
    ]
    reference_paths: list[str] = []
    style_code = compact(item.get("style_code"))
    for index, asset in enumerate(reference_assets or [], start=1):
        if not isinstance(asset, Mapping):
            continue
        path = compact(asset.get("path"))
        if not path or path in reference_paths:
            continue
        reference_paths.append(path)
        next_assets.append(_image_asset(
            path,
            style_code=style_code,
            kind=compact(asset.get("kind")) or "detail_reference",
            label=compact(asset.get("label")) or ("款色参考图" if index == 1 else f"款色参考图 {index}"),
            index=index,
            status="reference",
            extra={
                "source_label": compact(asset.get("source_label")) or "款色参考图",
                "slot": "reference",
                "use_for_generation": bool(asset.get("use_for_generation")),
                "custom_upload": bool(asset.get("custom_upload")),
            },
        ))
    item["assets"] = next_assets
    if reference_paths:
        item["detail_reference_path"] = reference_paths[0]


def update_generation_confirmation(
    batch: dict[str, Any],
    items: Iterable[Mapping[str, Any]],
    *,
    execution_mode: object = "",
) -> dict[str, Any]:
    if compact(execution_mode):
        mode = str(normalize_chain_execution_mode(execution_mode)["mode"])
        batch["execution_mode"] = mode
        run_params = batch.setdefault("run_params", {})
        if not isinstance(run_params, dict):
            run_params = dict(run_params or {})
            batch["run_params"] = run_params
        run_params["execute_mode"] = mode
    for patch in items or []:
        if not isinstance(patch, Mapping):
            continue
        item = _find_confirmation_item(batch, patch)
        if not item:
            continue
        if "origin_path" in patch:
            item["origin_path"] = compact(patch.get("origin_path"))
            _upsert_reference_asset(item, kind="origin", label="原图/主图", path=item["origin_path"])
        if "detail_reference_path" in patch:
            item["detail_reference_path"] = compact(patch.get("detail_reference_path"))
            _upsert_reference_asset(item, kind="detail_reference", label="款色参考图", path=item["detail_reference_path"])
        if "reference_assets" in patch and isinstance(patch.get("reference_assets"), list):
            _sync_confirmation_reference_assets(item, patch.get("reference_assets") or [])
        if "generation_prompts" in patch and isinstance(patch.get("generation_prompts"), list):
            prompts = []
            for index, prompt in enumerate(patch.get("generation_prompts") or [], start=1):
                if not isinstance(prompt, Mapping):
                    continue
                status = compact(prompt.get("status")) or "pending"
                if status.lower() in {"removed", "deleted", "disabled", "skip", "skipped"}:
                    continue
                prompt_text = compact(prompt.get("custom_prompt") or prompt.get("prompt"))
                if not prompt_text:
                    continue
                prompt_index = int(prompt.get("prompt_index") or index)
                image_count = normalize_generation_image_count(
                    prompt.get("image_count") or prompt.get("generation_image_count") or prompt.get("count"),
                    1,
                )
                prompts.append({
                    **json_safe(dict(prompt)),
                    "id": compact(prompt.get("id")) or _generation_prompt_id(compact(item.get("style_code")), prompt_index, prompt_text),
                    "prompt_index": prompt_index,
                    "prompt_name": compact(prompt.get("prompt_name")) or f"新增 Prompt {prompt_index}",
                    "prompt": prompt_text,
                    "custom_prompt": compact(prompt.get("custom_prompt")),
                    "image_count": image_count,
                    "reference_paths": parse_list(prompt.get("reference_paths")),
                    "status": status,
                })
            item["generation_prompts"] = prompts
            item["generation_rows"] = [
                generation_confirmation_prompt_to_row(batch, item, prompt, index=index)
                for index, prompt in enumerate(prompts, start=1)
            ]
    batch["generation_prompt_total"] = generation_confirmation_prompt_total(batch.get("items") or [])
    batch["updated_at"] = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    save_approval_batch(batch)
    return batch


def generation_confirmation_prompt_to_row(
    batch: Mapping[str, Any],
    item: Mapping[str, Any],
    prompt: Mapping[str, Any],
    *,
    index: int,
) -> dict[str, Any]:
    template = prompt.get("generation_row") if isinstance(prompt.get("generation_row"), Mapping) else {}
    if not template:
        rows = item.get("generation_rows") if isinstance(item.get("generation_rows"), list) else []
        template = next((row for row in rows if isinstance(row, Mapping)), {})
    row = dict(template or {})
    workflow = workflow_to_dict(item.get("workflow") or item)
    prompt_index = int(prompt.get("prompt_index") or index)
    prompt_text = compact(prompt.get("custom_prompt") or prompt.get("prompt"))
    reference_paths = parse_list(prompt.get("reference_paths"))
    if not reference_paths:
        reference_paths = [compact(item.get("origin_path"))]
        reference_paths = [path for path in reference_paths if path]
    run_params = batch.get("run_params") if isinstance(batch.get("run_params"), Mapping) else {}
    payload = row.get("__1xm_payload") if isinstance(row.get("__1xm_payload"), Mapping) else {}
    image_count = normalize_generation_image_count(
        prompt.get("image_count")
        or prompt.get("generation_image_count")
        or prompt.get("count")
        or row.get("生成数量")
        or payload.get("n"),
        1,
    )
    row.update({
        "表格行号": workflow.get("row_no", ""),
        "款号": workflow.get("style_code", ""),
        "商品ID": workflow.get("item_id", ""),
        "品类": workflow.get("category", ""),
        "性别": workflow.get("gender", ""),
        "SKC编码": workflow.get("skc_code", ""),
        "提示词分组": compact(prompt.get("prompt_group")) or compact(row.get("提示词分组")) or category_to_prompt_sheet(workflow.get("category", "")),
        "提示词字段名": compact(prompt.get("prompt_name")) or f"新增 Prompt {prompt_index}",
        "提示词序号": prompt_index,
        "最终提示词": prompt_text,
        "完整Prompt": prompt_text,
        "生成数量": image_count,
        "参考图文件": "\n".join(reference_paths),
        "__1xm_generate": True,
        "__1xm_reference_paths": reference_paths,
        "__task_run_uid": approval_batch_run_uid(batch),
    })
    row["尺寸"] = compact(row.get("尺寸")) or compact(run_params.get("image_size")) or "1536x2048"
    row["格式"] = compact(row.get("格式")) or compact(run_params.get("output_format")) or "jpeg"
    row["质量"] = compact(row.get("质量")) or compact(run_params.get("quality")) or "auto"
    row["模型"] = compact(row.get("模型")) or compact(run_params.get("model")) or "gpt-image-2"
    row["比例"] = compact(row.get("比例")) or compact(payload.get("ratio")) or compact(run_params.get("ratio")) or image_ratio_label(row["尺寸"])
    row["__1xm_key_tier"] = compact(row.get("__1xm_key_tier")) or compact(run_params.get("one_xm_key_tier")) or "4k"
    row["__1xm_payload"] = {
        **dict(payload),
        "model": row["模型"],
        "prompt": prompt_text,
        "size": row["尺寸"],
        "ratio": row["比例"],
        "quality": row["质量"],
        "output_format": row["格式"],
        "n": image_count,
    }
    apply_fresh_approval_generation_key(row, batch, "confirm")
    return row


def prepare_generation_confirmation_placeholders(batch: dict[str, Any]) -> dict[str, Any]:
    if compact(batch.get("status")) not in {GENERATION_CONFIRMATION_STATUS, "generating"}:
        raise ValueError("当前批次不在确认提交生图状态")
    run_params = batch.get("run_params") if isinstance(batch.get("run_params"), Mapping) else {}
    prompt_total = 0
    requested_images = 0
    now = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    for item in batch.get("items") or []:
        if not isinstance(item, dict):
            continue
        item["assets"] = [asset for asset in item.get("assets") or [] if asset.get("kind") != "ai"]
        main_asset = next(
            (
                asset for asset in item.get("assets") or []
                if isinstance(asset, Mapping)
                and (compact(asset.get("kind")) == "origin" or compact(asset.get("slot")) == "main")
                and compact(asset.get("path"))
            ),
            {},
        )
        placeholder_preview_path = compact(main_asset.get("path")) or compact(item.get("origin_path"))
        prompts = item.get("generation_prompts") if isinstance(item.get("generation_prompts"), list) else []
        item_rows = []
        for index, prompt in enumerate(prompts, start=1):
            if not isinstance(prompt, dict):
                continue
            status = compact(prompt.get("status")).lower()
            if status in {"removed", "deleted", "disabled", "skip", "skipped"}:
                continue
            row = generation_confirmation_prompt_to_row(batch, item, prompt, index=index)
            item_rows.append(row)
            prompt["status"] = "generating"
            prompt["generation_row"] = json_safe(row)
            prompt_total += 1
            image_count = generation_count_from_row(row, 1)
            requested_images += image_count
            for image_offset in range(1, image_count + 1):
                display_index = index if image_count == 1 else f"{index}-{image_offset}"
                item.setdefault("assets", []).append(_image_asset(
                    "",
                    style_code=compact(item.get("style_code")),
                    kind="ai",
                    label=f"AI 图 {display_index}",
                    index=display_index,
                    status="generating",
                    extra={
                        "prompt_index": index,
                        "prompt_name": compact(prompt.get("prompt_name")),
                        "prompt_group": compact(prompt.get("prompt_group")),
                        "prompt": compact(row.get("完整Prompt") or row.get("最终提示词")),
                        "generation_row": json_safe(row),
                        "reference_paths": parse_list(row.get("参考图文件")),
                        "reference_labels": [Path(ref).name for ref in parse_list(row.get("参考图文件"))],
                        "task_run_uid": compact(row.get("__task_run_uid")) or approval_batch_run_uid(batch),
                        "custom_prompt": "",
                        "review_note": "",
                        "placeholder": True,
                        "placeholder_preview_path": placeholder_preview_path,
                        "created_at": now,
                    },
                ))
        item["generation_rows"] = json_safe(item_rows)
    batch["status"] = "generating"
    batch["generation_summary"] = {
        "requested": requested_images,
        "requested_prompts": prompt_total,
        "generated": 0,
        "failed": 0,
    }
    batch["submit_progress"] = {
        "status": "running",
        "total": prompt_total,
        "completed": 0,
        "attempted": 0,
        "succeeded": 0,
        "failed": 0,
        "image_total": requested_images,
        "current_style": "",
        "message": f"后台生图中：{prompt_total} 条 Prompt / {requested_images} 张 AI 图",
    }
    batch["updated_at"] = now
    batch["approval_message"] = render_approval_message_template(
        str(run_params.get("approval_message_template") or ""),
        batch,
    )
    save_approval_batch(batch)
    return batch


def _same_generation_prompt_asset(asset: Mapping[str, Any], prompt_index: int) -> bool:
    if compact(asset.get("kind")) != "ai":
        return False
    return compact(asset.get("prompt_index")) == compact(prompt_index)


def _prune_generation_ai_assets(item: dict[str, Any], *, keep_placeholders: bool) -> None:
    next_assets = []
    for asset in item.get("assets") or []:
        if not isinstance(asset, Mapping):
            continue
        if compact(asset.get("kind")) != "ai":
            next_assets.append(asset)
            continue
        if keep_placeholders and (asset.get("placeholder") or compact(asset.get("status")) == "generating"):
            next_assets.append(asset)
    item["assets"] = next_assets


def _replace_generation_prompt_assets(item: dict[str, Any], prompt_index: int, assets: list[dict[str, Any]]) -> None:
    next_assets: list[Mapping[str, Any]] = []
    inserted = False
    for asset in item.get("assets") or []:
        if isinstance(asset, Mapping) and _same_generation_prompt_asset(asset, prompt_index):
            if not inserted:
                next_assets.extend(assets)
                inserted = True
            continue
        if isinstance(asset, Mapping):
            next_assets.append(asset)
    if not inserted:
        next_assets.extend(assets)
    item["assets"] = list(next_assets)


def _remove_generation_prompt_placeholders(item: dict[str, Any], prompt_index: int) -> None:
    item["assets"] = [
        asset for asset in item.get("assets") or []
        if not (
            isinstance(asset, Mapping)
            and asset.get("placeholder")
            and _same_generation_prompt_asset(asset, prompt_index)
        )
    ]


def submit_generation_confirmation_batch(batch: dict[str, Any], *, log=print) -> dict[str, Any]:
    if compact(batch.get("status")) not in {GENERATION_CONFIRMATION_STATUS, "generating"}:
        raise ValueError("当前批次不在确认提交生图状态")
    run_params = batch.get("run_params") if isinstance(batch.get("run_params"), Mapping) else {}
    settings = resolve_one_xm_settings()
    require_one_xm_key_for_generation(
        settings,
        model=compact(run_params.get("model")) or "gpt-image-2",
        image_size=compact(run_params.get("image_size")) or "1536x2048",
        key_tier=compact(run_params.get("one_xm_key_tier")) or "4k",
    )
    artifact_dir = Path(compact(batch.get("artifact_dir")) or compact(Path(compact(batch.get("json_path"))).parent) or ".")
    artifact_dir.mkdir(parents=True, exist_ok=True)

    jobs: list[tuple[dict[str, Any], dict[str, Any], int, dict[str, Any]]] = []
    keep_placeholders = compact(batch.get("status")) == "generating"
    for item in batch.get("items") or []:
        if not isinstance(item, dict):
            continue
        _prune_generation_ai_assets(item, keep_placeholders=keep_placeholders)
        prompts = item.get("generation_prompts") if isinstance(item.get("generation_prompts"), list) else []
        item_rows = []
        for index, prompt in enumerate(prompts, start=1):
            if not isinstance(prompt, dict):
                continue
            status = compact(prompt.get("status")).lower()
            if status in {"removed", "deleted", "disabled", "skip", "skipped"}:
                continue
            row = generation_confirmation_prompt_to_row(batch, item, prompt, index=index)
            item_rows.append(row)
            jobs.append((item, prompt, index, row))
        item["generation_rows"] = json_safe(item_rows)

    concurrency = max(1, min(100, int(run_params.get("generation_concurrency") or 100)))
    options = {
        "one_xm_key_tier": compact(run_params.get("one_xm_key_tier")) or "4k",
        "retry_attempts": int(run_params.get("retry_attempts") or 3),
        "compensate_attempts": int(run_params.get("compensate_attempts") or 2),
        "poll_timeout_minutes": float(run_params.get("poll_timeout_minutes") or 12),
    }

    def generate(job: tuple[dict[str, Any], dict[str, Any], int, dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any], int, dict[str, Any], list[str]]:
        item, prompt, index, row = job
        patched = run_one_xm_generation_row(row, options, settings)
        paths = download_generated_images(patched, artifact_dir)
        patched["本地生成图文件"] = "\n".join(paths)
        return item, prompt, index, patched, paths

    generated = 0
    failures = 0
    succeeded_prompts = 0
    requested_images = sum(generation_count_from_row(row, 1) for _item, _prompt, _index, row in jobs)
    if jobs:
        log(f"[approval] 确认提交生图任务 {len(jobs)} 条 Prompt / {requested_images} 张，并发 {concurrency}")

    def persist_generation_progress(*, status: str = "running", current_style: str = "", message: str = "") -> None:
        completed = succeeded_prompts + failures
        now = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
        batch["status"] = "generating" if status == "running" else ("pending_approval" if generated else "generation_failed")
        batch["generation_status"] = (
            "running" if status == "running"
            else "partial_success" if generated and failures
            else "succeeded" if generated
            else "failed"
        )
        if status != "running":
            execution_mode = compact(batch.get("execution_mode") or run_params.get("execute_mode")) or "approval_then_create"
            batch["approval_status"] = "pending" if generated and execution_mode == "approval_then_create" else "not_required"
        batch["generation_summary"] = {
            "requested": requested_images,
            "requested_prompts": len(jobs),
            "generated": generated,
            "failed": failures,
        }
        batch["submit_progress"] = {
            "status": status,
            "total": len(jobs),
            "completed": completed if status == "running" else len(jobs),
            "attempted": completed if status == "running" else len(jobs),
            "succeeded": succeeded_prompts,
            "failed": failures,
            "image_total": requested_images,
            "current_style": current_style,
            "message": message or (
                f"生图进行中：已生成 {generated}/{requested_images} 张，完成 {completed}/{len(jobs)} 条 Prompt"
                if status == "running"
                else (f"生图完成：{generated} 张" if generated else "生图失败，请检查失败原因")
            ),
        }
        batch["updated_at"] = now
        batch["approval_message"] = render_approval_message_template(
            str(run_params.get("approval_message_template") or ""),
            batch,
        )
        save_approval_batch(batch)

    with ThreadPoolExecutor(max_workers=concurrency) as executor:
        future_map = {executor.submit(generate, job): job for job in jobs}
        for future in as_completed(future_map):
            item, prompt, index, row = future_map[future]
            try:
                _item, _prompt, prompt_index, patched, paths = future.result()
            except Exception as exc:
                failures += 1
                prompt["status"] = "failed"
                prompt["error"] = str(exc)
                _remove_generation_prompt_placeholders(item, index)
                persist_generation_progress(
                    current_style=compact(item.get("style_code")),
                    message=f"{compact(item.get('style_code')) or '当前款'} / Prompt {index} 生图失败：{exc}",
                )
                continue
            if not paths:
                failures += 1
                prompt["status"] = "failed"
                prompt["generation_row"] = json_safe(patched)
                prompt["error"] = compact(patched.get("备注")) or "未下载到生成图片"
                _remove_generation_prompt_placeholders(item, prompt_index)
                persist_generation_progress(
                    current_style=compact(item.get("style_code")),
                    message=f"{compact(item.get('style_code')) or '当前款'} / Prompt {prompt_index} 未下载到生成图片",
                )
                continue
            prompt["status"] = "generated"
            prompt["generation_row"] = json_safe(patched)
            generated += len(paths)
            succeeded_prompts += 1
            prompt_assets: list[dict[str, Any]] = []
            for image_offset, path in enumerate(paths, start=1):
                display_index = prompt_index if len(paths) == 1 else f"{prompt_index}-{image_offset}"
                asset = _image_asset(
                    path,
                    style_code=compact(item.get("style_code")),
                    kind="ai",
                    label=f"AI 图 {display_index}",
                    index=display_index,
                    status="pending",
                    extra={
                        "prompt_index": prompt_index,
                        "prompt_name": compact(prompt.get("prompt_name")),
                        "prompt_group": compact(prompt.get("prompt_group")),
                        "prompt": compact(patched.get("完整Prompt") or patched.get("最终提示词")),
                        "generation_row": json_safe(patched),
                        "reference_paths": parse_list(patched.get("参考图文件")),
                        "reference_labels": [Path(ref).name for ref in parse_list(patched.get("参考图文件"))],
                        "task_run_uid": compact(patched.get("__task_run_uid")) or approval_batch_run_uid(batch),
                        "custom_prompt": "",
                        "review_note": "",
                        "generated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
                    },
                )
                prompt_assets.append(asset)
            _replace_generation_prompt_assets(item, prompt_index, prompt_assets)
            persist_generation_progress(current_style=compact(item.get("style_code")))

    persist_generation_progress(status="completed" if generated else "failed")
    notify_channel = compact(run_params.get("approval_notify_channel") or "dingtalk").lower()
    if generated and notify_channel and notify_channel != "none":
        notify_approval_batch(batch, channel=notify_channel, log=log)
    maybe_sync_approval_batch_to_cloud(batch, log=log)
    return {"ok": generated > 0, "generated": generated, "failed": failures, "status": batch.get("status"), "batch": batch}


def approval_asset_matches_batch_run(batch: Mapping[str, Any], asset: Mapping[str, Any]) -> bool:
    batch_run_uid = compact(batch.get("task_run_uid"))
    if not batch_run_uid:
        return True
    asset_run_uid = compact(asset.get("task_run_uid"))
    generation_row = asset.get("generation_row") if isinstance(asset.get("generation_row"), Mapping) else {}
    generation_run_uid = compact(generation_row.get("__task_run_uid"))
    known_run_uid = asset_run_uid or generation_run_uid
    return known_run_uid == batch_run_uid


def approve_generated_assets_for_direct_create(batch: dict[str, Any]) -> int:
    approved = 0
    for item in batch.get("items") or []:
        if not isinstance(item, dict):
            continue
        for asset in item.get("assets") or []:
            if (
                not isinstance(asset, dict)
                or compact(asset.get("kind")) != "ai"
                or not compact(asset.get("path"))
                or not approval_asset_matches_batch_run(batch, asset)
            ):
                continue
            asset["status"] = "approved"
            approved += 1
    batch["approval_status"] = "not_required"
    batch["test_task_status"] = "creating"
    batch["updated_at"] = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    save_approval_batch(batch)
    return approved


def tmall_submit_row_is_online_success(row: Mapping[str, Any]) -> bool:
    if not isinstance(row, Mapping):
        return False
    if compact(row.get("阶段")) != "天猫上传/创建测图任务":
        return False
    result_text = compact(row.get("执行结果"))
    if not result_text or "失败" in result_text or "跳过" in result_text:
        return False
    task_id = compact(row.get("任务ID"))
    if not task_id or task_id == "<experimentTaskId>":
        return False
    online_text = compact(row.get("上线结果"))
    readback_text = compact(row.get("页面回读")).replace(" ", "")
    return online_text == "已上线" or "status=1" in readback_text or "status：1" in readback_text


def latest_submit_rows_by_style(batch: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    latest_rows: dict[str, Mapping[str, Any]] = {}
    for row in batch.get("submit_result_rows") or []:
        if not isinstance(row, Mapping):
            continue
        if compact(row.get("阶段")) != "天猫上传/创建测图任务":
            continue
        style_code = compact(row.get("款号"))
        if not style_code:
            continue
        latest_rows[style_code] = row
    return latest_rows


def successful_submit_rows_by_style(batch: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    latest_rows = latest_submit_rows_by_style(batch)
    return {
        style_code: row
        for style_code, row in latest_rows.items()
        if tmall_submit_row_is_online_success(row)
    }


def submit_image_path_key(path: object) -> str:
    return compact(path).replace("\\", "/")


def successful_submit_ai_paths_by_style(batch: Mapping[str, Any]) -> dict[str, set[str]]:
    submitted: dict[str, set[str]] = {}
    pending_generation_paths: dict[str, list[str]] = {}
    for row in batch.get("submit_result_rows") or []:
        if not isinstance(row, Mapping):
            continue
        style_code = compact(row.get("款号"))
        if not style_code:
            continue
        stage = compact(row.get("阶段"))
        if stage == "1XM生图":
            paths = [submit_image_path_key(path) for path in parse_list(row.get("本地生成图文件") or row.get("本地文件"))]
            paths = [path for path in paths if path]
            if paths:
                pending_generation_paths.setdefault(style_code, []).extend(paths)
            continue
        if stage != "天猫上传/创建测图任务":
            continue
        if not tmall_submit_row_is_online_success(row):
            pending_generation_paths[style_code] = []
            continue
        paths = [submit_image_path_key(path) for path in parse_list(row.get("提交图片文件"))]
        if not paths:
            paths = list(pending_generation_paths.get(style_code) or [])
        paths = [path for path in paths if path]
        if paths:
            submitted.setdefault(style_code, set()).update(paths)
        pending_generation_paths[style_code] = []
    return submitted


def make_selected_upload_plan(
    item: Mapping[str, Any],
    workflow_data: Mapping[str, Any],
    generated_paths: list[str],
    submit_mode: str,
) -> dict[str, Any]:
    return {
        "workflow": workflow_data,
        "origin_path": compact(item.get("origin_path")),
        "detail_reference_path": compact(item.get("detail_reference_path")),
        "generated_paths": generated_paths,
        "item": item,
        "submit_mode": submit_mode,
    }


def selected_upload_plan_from_approval_batch(batch: Mapping[str, Any]) -> list[dict[str, Any]]:
    plans: list[dict[str, Any]] = []
    rerun_plans: list[dict[str, Any]] = []
    run_params = batch.get("run_params") if isinstance(batch.get("run_params"), Mapping) else {}
    force_resubmit_success = is_truthy(run_params.get("approval_resubmit_successful"))
    successful_styles = set(successful_submit_rows_by_style(batch)) if not force_resubmit_success else set()
    submitted_paths_by_style = successful_submit_ai_paths_by_style(batch) if not force_resubmit_success else {}
    for item in batch.get("items") or []:
        workflow_data = workflow_to_dict(item.get("workflow") or item)
        style_code = compact(workflow_data.get("style_code"))
        approved_paths = [
            compact(asset.get("path"))
            for asset in item.get("assets") or []
            if asset.get("kind") == "ai"
            and compact(asset.get("path"))
            and compact(asset.get("status")) in APPROVED_ASSET_STATUSES
            and approval_asset_matches_batch_run(batch, asset)
        ]
        if not approved_paths:
            continue
        submitted_paths = submitted_paths_by_style.get(style_code) or set()
        unsubmitted_paths = [
            path for path in approved_paths
            if submit_image_path_key(path) not in submitted_paths
        ] if submitted_paths else approved_paths
        if unsubmitted_paths:
            mode = (
                "append_unsubmitted" if submitted_paths else
                "rerun_confirmed" if style_code in successful_styles else
                "submit_confirmed"
            )
            plans.append(make_selected_upload_plan(item, workflow_data, unsubmitted_paths, mode))
            continue
        rerun_plans.append(make_selected_upload_plan(item, workflow_data, approved_paths, "rerun_confirmed"))
    return plans or rerun_plans


def approval_result_rows(batch: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for item in batch.get("items") or []:
        ai_total = len([asset for asset in item.get("assets") or [] if asset.get("kind") == "ai"])
        rows.append({
            "表格行号": item.get("row_no", ""),
            "款号": item.get("style_code", ""),
            "商品ID": item.get("item_id", ""),
            "阶段": "图片审批",
            "审批批次ID": batch.get("batch_id", ""),
            "审批状态": batch.get("status", ""),
            "审批看板": batch.get("board_url", ""),
            "云端审批看板": batch.get("cloud_board_url", ""),
            "本地审批看板": batch.get("local_board_url", ""),
            "执行结果": "等待审批",
            "备注": f"AI图={ai_total}；{batch.get('approval_message', '')}",
        })
    return rows


def cloud_approval_sync_warning_row(batch: Mapping[str, Any], message: str) -> dict[str, Any]:
    return {
        "表格行号": "",
        "款号": "",
        "商品ID": "",
        "阶段": "云端审批同步",
        "审批批次ID": batch.get("batch_id", ""),
        "审批状态": batch.get("status", ""),
        "审批看板": batch.get("board_url", ""),
        "云端审批看板": batch.get("cloud_board_url", ""),
        "本地审批看板": batch.get("local_board_url", ""),
        "执行结果": "警告",
        "备注": compact(message),
    }


def maybe_sync_approval_batch_to_cloud(batch: dict[str, Any], *, log=print) -> dict[str, Any] | None:
    from core.cloud_approval_client import CloudApprovalClient
    from core.cloud_batch_sync import sync_local_approval_batch
    from core.config import load_config

    cloud_config = (load_config().get("cloud_approval") or {})
    job_uid = approval_batch_run_uid(batch) or compact(batch.get("batch_id")) or "approval-batch"
    base_url = compact(cloud_config.get("base_url"))
    if not cloud_config.get("machine_enabled"):
        return record_cloud_approval_sync_warning(batch, job_uid, base_url, "云端审批同步跳过：未启用云端审批任务机", log=log)
    if not base_url:
        return record_cloud_approval_sync_warning(batch, job_uid, base_url, "云端审批同步跳过：未配置云端审批地址", log=log)
    credentials = data_sink.get_cloud_machine_credentials() or {}
    machine_token = compact(credentials.get("machine_token"))
    if not machine_token:
        return record_cloud_approval_sync_warning(batch, job_uid, base_url, "云端审批同步跳过：未注册任务机或机器 token 缺失", log=log)
    client = CloudApprovalClient(
        base_url,
        machine_token=machine_token,
        timeout=float(cloud_config.get("poll_timeout_seconds") or 30),
    )
    try:
        result = sync_local_approval_batch(batch, client)
        local_board_url = compact(batch.get("local_board_url") or batch.get("board_url"))
        cloud_board_url = cloud_approval_board_url(base_url, batch.get("batch_id"))
        batch["local_board_url"] = local_board_url
        batch["cloud_board_url"] = cloud_board_url
        batch["board_url"] = cloud_board_url
        run_params = batch.get("run_params") if isinstance(batch.get("run_params"), Mapping) else {}
        batch["approval_message"] = render_approval_message_template(
            str(run_params.get("approval_message_template") or ""),
            batch,
        )
        batch["cloud_sync"] = {
            "status": "synced",
            "base_url": base_url,
            "board_url": cloud_board_url,
            "result": json_safe(result),
            "synced_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        }
        data_sink.record_cloud_job_event(
            job_uid,
            "approval_batch_sync_complete",
            "local approval batch synced to cloud",
            {"batch_id": batch.get("batch_id"), "asset_count": len(result.get("assets") or []) if isinstance(result, Mapping) else 0},
        )
        save_approval_batch(batch)
        return result
    except Exception as exc:
        warning = f"云端审批同步失败：{exc}"
        batch["cloud_sync"] = {
            "status": "warning",
            "base_url": base_url,
            "warning": compact(warning),
            "updated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        }
        data_sink.record_cloud_job_event(
            job_uid,
            "approval_batch_sync_warning",
            "local approval batch cloud sync failed",
            {"batch_id": batch.get("batch_id"), "error": compact(str(exc))},
        )
        save_approval_batch(batch)
        if log:
            log(f"[warn] {warning}")
        return None


def record_cloud_approval_sync_warning(batch: dict[str, Any], job_uid: str, base_url: str, warning: str, *, log=print) -> None:
    batch["cloud_sync"] = {
        "status": "warning",
        "base_url": compact(base_url),
        "warning": compact(warning),
        "updated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
    }
    data_sink.record_cloud_job_event(
        job_uid,
        "approval_batch_sync_warning",
        "local approval batch cloud sync skipped",
        {"batch_id": batch.get("batch_id"), "warning": compact(warning)},
    )
    save_approval_batch(batch)
    if log:
        log(f"[warn] {warning}")
    return None


def notify_approval_batch(batch: Mapping[str, Any], *, channel: str = "dingtalk", log=None) -> None:
    try:
        from core import notifier

        notifier.send(
            channel=channel or "dingtalk",
            title="天猫AI测图审批待处理",
            records=len(batch.get("items") or []),
            adapter_name="天猫运营助手",
            task_name="天猫AI测图全链路",
            sample_rows=[{
                "审批批次ID": batch.get("batch_id", ""),
                "审批看板": batch.get("board_url", ""),
                "状态": batch.get("status", ""),
            }],
            message=compact(batch.get("approval_message")),
        )
        if log:
            log(f"[chain] 审批通知已通过 {channel or 'dingtalk'} 发送")
    except Exception as exc:
        if log:
            log(f"[warn] 审批通知发送失败：{exc}")


def tmall_submission_success_details(rows: Iterable[Mapping[str, Any]]) -> list[dict[str, str]]:
    details: list[dict[str, str]] = []
    seen: set[str] = set()
    for row in rows or []:
        if not isinstance(row, Mapping) or not tmall_submit_row_is_online_success(row):
            continue
        style_code = compact(row.get("款号"))
        item_id = compact(row.get("商品ID"))
        detail_url = compact(row.get("测图详情URL")) or tmall_material_test_detail_url(item_id)
        key = style_code or item_id or compact(row.get("任务ID"))
        if not key or key in seen:
            continue
        seen.add(key)
        details.append({
            "style_code": style_code,
            "item_id": item_id,
            "task_id": compact(row.get("任务ID")),
            "detail_url": detail_url,
        })
    return details


def tmall_submission_failed_details(rows: Iterable[Mapping[str, Any]]) -> list[dict[str, str]]:
    details: list[dict[str, str]] = []
    seen: set[str] = set()
    for row in rows or []:
        if not isinstance(row, Mapping):
            continue
        if compact(row.get("阶段")) != "天猫上传/创建测图任务":
            continue
        if tmall_submit_row_is_online_success(row):
            continue
        result_text = compact(row.get("执行结果"))
        if "失败" not in result_text and "跳过" not in result_text:
            continue
        style_code = compact(row.get("款号"))
        item_id = compact(row.get("商品ID"))
        key = style_code or item_id or compact(row.get("任务ID"))
        if not key or key in seen:
            continue
        seen.add(key)
        details.append({
            "style_code": style_code,
            "item_id": item_id,
            "task_id": compact(row.get("任务ID")),
            "reason": compact(row.get("备注") or result_text),
        })
    return details


def build_tmall_submission_notification_message(
    batch: Mapping[str, Any],
    *,
    status: str,
    attempted: int,
    succeeded: int,
    failed: int,
    rows: Iterable[Mapping[str, Any]],
) -> str:
    success_details = tmall_submission_success_details(rows)
    failed_details = tmall_submission_failed_details(rows)
    lines = [
        "### 天猫AI测图任务创建结果",
        f"- 批次：{compact(batch.get('batch_id')) or '-'}",
        f"- 状态：{compact(status) or '-'}",
        f"- 汇总：尝试 {attempted} 款 / 成功 {succeeded} 款 / 失败 {failed} 款",
        "",
        "#### 成功款号与详情URL",
    ]
    if success_details:
        for item in success_details:
            meta_parts = []
            if item.get("item_id"):
                meta_parts.append(f"商品ID {item['item_id']}")
            if item.get("task_id"):
                meta_parts.append(f"任务ID {item['task_id']}")
            meta = f"（{'，'.join(meta_parts)}）" if meta_parts else ""
            detail_url = item.get("detail_url") or "未生成详情URL"
            lines.append(f"- {item.get('style_code') or item.get('item_id') or '-'}{meta}：{detail_url}")
    else:
        lines.append("- 无")
    if failed_details:
        lines.extend(["", "#### 失败款号"])
        for item in failed_details[:20]:
            item_id = f"，商品ID {item['item_id']}" if item.get("item_id") else ""
            reason = item.get("reason") or "未记录失败原因"
            lines.append(f"- {item.get('style_code') or item.get('item_id') or '-'}{item_id}：{reason}")
    return "\n".join(lines)


def notify_tmall_submission_result(
    batch: Mapping[str, Any],
    *,
    status: str,
    attempted: int,
    succeeded: int,
    failed: int,
    rows: Iterable[Mapping[str, Any]],
    channel: str = "dingtalk",
    log=None,
) -> None:
    notify_channel = compact(channel or "dingtalk").lower()
    if not notify_channel or notify_channel == "none":
        return
    try:
        from core import notifier

        message = build_tmall_submission_notification_message(
            batch,
            status=status,
            attempted=attempted,
            succeeded=succeeded,
            failed=failed,
            rows=rows,
        )
        notifier.send(
            channel=notify_channel,
            title="天猫AI测图任务创建结果",
            records=attempted,
            adapter_name="天猫运营助手",
            task_name="天猫AI测图全链路",
            sample_rows=tmall_submission_success_details(rows),
            message=message,
        )
        if log:
            log(f"[chain] 测图任务创建结果通知已通过 {notify_channel} 发送")
    except Exception as exc:
        if log:
            log(f"[warn] 测图任务创建结果通知发送失败：{exc}")


def find_approval_asset(batch: Mapping[str, Any], asset_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    target = compact(asset_id)
    for item in batch.get("items") or []:
        if not isinstance(item, dict):
            continue
        for asset in item.get("assets") or []:
            if isinstance(asset, dict) and compact(asset.get("id")) == target:
                return item, asset
    raise KeyError(f"未找到审批图片：{asset_id}")


def find_approval_item(batch: Mapping[str, Any], *, item_id: str = "", style_code: str = "") -> dict[str, Any]:
    target_item_id = compact(item_id)
    target_style = compact(style_code)
    for item in batch.get("items") or []:
        if not isinstance(item, dict):
            continue
        if target_item_id and target_item_id in {compact(item.get("id")), compact(item.get("item_id"))}:
            return item
        if target_style and compact(item.get("style_code")) == target_style:
            return item
    raise KeyError(f"未找到审批款号：{target_style or target_item_id}")


def approval_generation_defaults(batch: Mapping[str, Any], item: Mapping[str, Any], generation_row: Mapping[str, Any] | None = None) -> dict[str, Any]:
    run_params = batch.get("run_params") if isinstance(batch.get("run_params"), Mapping) else {}
    source_row = generation_row if isinstance(generation_row, Mapping) else {}
    existing_ai = next((asset for asset in item.get("assets") or [] if isinstance(asset, Mapping) and asset.get("kind") == "ai"), {})
    existing_row = existing_ai.get("generation_row") if isinstance(existing_ai.get("generation_row"), Mapping) else {}
    source_payload = source_row.get("__1xm_payload") if isinstance(source_row.get("__1xm_payload"), Mapping) else {}
    existing_payload = existing_row.get("__1xm_payload") if isinstance(existing_row.get("__1xm_payload"), Mapping) else {}
    count = normalize_generation_image_count(
        source_row.get("生成数量")
        or source_payload.get("n")
        or existing_row.get("生成数量")
        or existing_payload.get("n")
        or run_params.get("generation_image_count")
        or run_params.get("count")
        or run_params.get("n"),
        1,
    )
    image_size = compact(source_row.get("尺寸")) or compact(existing_row.get("尺寸")) or compact(run_params.get("image_size")) or "1536x2048"
    ratio = (
        compact(source_row.get("比例"))
        or compact(source_payload.get("ratio"))
        or compact(existing_row.get("比例"))
        or compact(existing_payload.get("ratio"))
        or compact(run_params.get("ratio"))
        or image_ratio_label(image_size)
    )
    return {
        "model": compact(source_row.get("模型") or existing_row.get("模型") or run_params.get("model")) or "gpt-image-2",
        "image_size": image_size,
        "ratio": ratio,
        "quality": normalize_quality(source_row.get("质量") or existing_row.get("质量") or run_params.get("quality") or "auto"),
        "output_format": normalize_output_format(source_row.get("格式") or existing_row.get("格式") or run_params.get("output_format"), "jpeg"),
        "count": count,
        "one_xm_key_tier": compact(source_row.get("__1xm_key_tier")) or compact(existing_row.get("__1xm_key_tier")) or compact(run_params.get("one_xm_key_tier")) or "4k",
        "retry_attempts": to_int(run_params.get("retry_attempts"), 3),
        "compensate_attempts": to_int(run_params.get("compensate_attempts"), 2),
        "poll_timeout_minutes": float(run_params.get("poll_timeout_minutes") or 12),
    }


def regenerate_approval_asset(
    batch: dict[str, Any],
    asset_id: str,
    *,
    prompt: str = "",
    reference_paths: list[str] | None = None,
) -> dict[str, Any]:
    item, asset = find_approval_asset(batch, asset_id)
    if asset.get("kind") != "ai":
        raise ValueError("只有 AI 图支持重新生成")
    generation_row = dict(asset.get("generation_row") or {})
    if not generation_row:
        raise ValueError("审批图片缺少原始生图参数，无法重试")
    final_prompt = compact(prompt) or compact(asset.get("custom_prompt")) or compact(asset.get("prompt")) or compact(generation_row.get("完整Prompt") or generation_row.get("最终提示词"))
    refs = parse_list(reference_paths if reference_paths is not None else asset.get("reference_paths"))
    if not refs:
        refs = parse_list(generation_row.get("参考图文件")) or [compact(item.get("origin_path"))]
    generation_row["最终提示词"] = final_prompt
    generation_row["完整Prompt"] = final_prompt
    generation_row["参考图文件"] = "\n".join(refs)
    generation_row["主参考图文件"] = refs[0] if refs else ""
    generation_row["细节参考图文件"] = "\n".join(refs[1:])
    generation_row["__1xm_reference_paths"] = refs
    defaults = approval_generation_defaults(batch, item, generation_row)
    generation_row["比例"] = compact(generation_row.get("比例")) or defaults["ratio"]
    generation_row["生成数量"] = normalize_generation_image_count(generation_row.get("生成数量") or defaults["count"], 1)
    if isinstance(generation_row.get("__1xm_payload"), dict):
        generation_row["__1xm_payload"]["prompt"] = final_prompt
        generation_row["__1xm_payload"]["ratio"] = compact(generation_row["__1xm_payload"].get("ratio")) or generation_row["比例"]
        generation_row["__1xm_payload"]["n"] = normalize_generation_image_count(generation_row["__1xm_payload"].get("n") or generation_row["生成数量"], 1)
    else:
        generation_row["__1xm_payload"] = {
            "model": defaults["model"],
            "prompt": final_prompt,
            "size": defaults["image_size"],
            "ratio": defaults["ratio"],
            "quality": defaults["quality"],
            "output_format": defaults["output_format"],
            "n": defaults["count"],
        }
    nonce = apply_fresh_approval_generation_key(generation_row, batch, "retry")
    settings = resolve_one_xm_settings()
    defaults = approval_generation_defaults(batch, item, generation_row)
    require_one_xm_key_for_generation(
        settings,
        model=defaults["model"],
        image_size=defaults["image_size"],
        key_tier=defaults["one_xm_key_tier"],
    )
    patched = run_one_xm_generation_row(generation_row, {
        "one_xm_key_tier": defaults["one_xm_key_tier"],
        "retry_attempts": defaults["retry_attempts"],
        "compensate_attempts": defaults["compensate_attempts"],
        "poll_timeout_minutes": defaults["poll_timeout_minutes"],
    }, settings)
    artifact_dir = Path(compact(batch.get("artifact_dir")) or ".")
    paths = download_generated_images(patched, artifact_dir)
    if not paths:
        raise RuntimeError(compact(patched.get("备注")) or "重新生图没有下载到本地图片")
    regenerated_at = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    asset.update({
        "path": paths[0],
        "filename": Path(paths[0]).name,
        "status": "pending",
        "prompt": final_prompt,
        "custom_prompt": final_prompt,
        "reference_paths": refs,
        "generation_row": json_safe({**patched, "本地生成图文件": "\n".join(paths)}),
        "task_run_uid": approval_batch_run_uid(batch),
        "regeneration_nonce": nonce,
        "regenerated_at": regenerated_at,
        "updated_at": regenerated_at,
    })
    batch["updated_at"] = regenerated_at
    save_approval_batch(batch)
    return asset


def generate_approval_asset_for_item(
    batch: dict[str, Any],
    *,
    item_id: str = "",
    style_code: str = "",
    prompt: str = "",
    main_image_path: str = "",
    reference_paths: list[str] | None = None,
) -> dict[str, Any]:
    item = find_approval_item(batch, item_id=item_id, style_code=style_code)
    final_prompt = compact(prompt)
    if not final_prompt:
        raise ValueError("新增生图缺少 Prompt")
    main_path = compact(main_image_path) or compact(item.get("origin_path"))
    if not main_path:
        raise ValueError("新增生图缺少主图")
    refs = [main_path, *parse_list(reference_paths)]
    refs = list(dict.fromkeys([ref for ref in refs if compact(ref)]))
    workflow = workflow_from_dict(item.get("workflow") or item)
    defaults = approval_generation_defaults(batch, item)
    ai_assets = [asset for asset in item.get("assets") or [] if isinstance(asset, Mapping) and asset.get("kind") == "ai"]
    prompt_index = len(ai_assets) + 1
    generation_row = {
        "表格行号": workflow.row_no,
        "款号": workflow.style_code,
        "SKC编码": workflow.skc_code,
        "商品ID": workflow.item_id,
        "品类": workflow.category,
        "性别": workflow.gender,
        "提示词分组": "审批新增",
        "提示词字段名": f"审批新增 {prompt_index}",
        "提示词序号": prompt_index,
        "尺寸": defaults["image_size"],
        "比例": defaults["ratio"],
        "格式": defaults["output_format"],
        "质量": defaults["quality"],
        "生成数量": defaults["count"],
        "参考图文件": "\n".join(refs),
        "主参考图文件": refs[0],
        "细节参考图文件": "\n".join(refs[1:]),
        "最终提示词": final_prompt,
        "完整Prompt": final_prompt,
        "1XM任务ID": "",
        "1XM轮询URL": "",
        "生成图URL": "",
        "生成图数量": "",
        "执行结果": "待生成",
        "备注": "",
        "__1xm_generate": True,
        "__1xm_key_tier": defaults["one_xm_key_tier"],
        "__task_run_uid": approval_batch_run_uid(batch),
        "__1xm_reference_paths": refs,
        "__1xm_payload": {
            "model": defaults["model"],
            "prompt": final_prompt,
            "size": defaults["image_size"],
            "ratio": defaults["ratio"],
            "quality": defaults["quality"],
            "output_format": defaults["output_format"],
            "n": defaults["count"],
        },
    }
    nonce = apply_fresh_approval_generation_key(generation_row, batch, "manual")
    settings = resolve_one_xm_settings()
    require_one_xm_key_for_generation(
        settings,
        model=defaults["model"],
        image_size=defaults["image_size"],
        key_tier=defaults["one_xm_key_tier"],
    )
    patched = run_one_xm_generation_row(generation_row, {
        "one_xm_key_tier": defaults["one_xm_key_tier"],
        "retry_attempts": defaults["retry_attempts"],
        "compensate_attempts": defaults["compensate_attempts"],
        "poll_timeout_minutes": defaults["poll_timeout_minutes"],
    }, settings)
    artifact_dir = Path(compact(batch.get("artifact_dir")) or ".")
    paths = download_generated_images(patched, artifact_dir)
    if not paths:
        raise RuntimeError(compact(patched.get("备注")) or "新增生图没有下载到本地图片")
    created_at = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    asset = _image_asset(
        paths[0],
        style_code=workflow.style_code,
        kind="ai",
        label=f"AI 图 {prompt_index}",
        index=f"{prompt_index}-{safe_token(nonce)}",
        status="pending",
        extra={
            "prompt_index": prompt_index,
            "prompt_name": f"审批新增 {prompt_index}",
            "prompt_group": "审批新增",
            "prompt": final_prompt,
            "custom_prompt": final_prompt,
            "reference_paths": refs,
            "generation_row": json_safe({**patched, "本地生成图文件": "\n".join(paths)}),
            "task_run_uid": approval_batch_run_uid(batch),
            "manual_added": True,
            "generation_nonce": nonce,
            "created_at": created_at,
            "updated_at": created_at,
            "review_note": "",
        },
    )
    item.setdefault("assets", []).append(asset)
    batch["updated_at"] = created_at
    save_approval_batch(batch)
    return asset


def normalize_quality(value: object) -> str:
    text = compact(value).lower()
    return text if text in {"auto", "standard", "low", "medium", "high"} else "auto"


def read_workbook_table(path: str, sheet: str | None = None, header_row: int = 1) -> dict[str, Any]:
    from core.api_server import _read_local_excel

    result = _read_local_excel(str(Path(path).expanduser()), sheet=sheet, header_row=header_row)
    if result.get("error"):
        raise RuntimeError(str(result["error"]))
    return result


def resolve_one_xm_settings() -> dict:
    from core.api_server import _resolve_one_xm_settings

    return _resolve_one_xm_settings()


def require_one_xm_key_for_generation(settings: Mapping[str, Any], *, model: str, image_size: str, key_tier: str = "auto") -> str:
    from core import ai_image_service

    try:
        config_id, _api_key = ai_image_service.select_model_key({
            "model_key": compact(model) or "gpt-image-2",
            "size": compact(image_size),
            "model_key_tier": compact(key_tier).lower() or "auto",
        }, settings)
    except ai_image_service.MissingModelKeyError as exc:
        raise RuntimeError(f"未配置 1XM Key：{exc}") from exc
    return config_id


def run_one_xm_generation_row(row: dict, run_params: dict, one_xm_settings: dict) -> dict:
    from core.api_server import _run_one_xm_generation_row

    return _run_one_xm_generation_row(row, run_params, one_xm_settings)


def sync_generation_plan_to_ai_image_workbench(
    plan: Mapping[str, Any],
    *,
    task_run_uid: str,
    model: str,
    image_size: str,
    ratio: str,
    output_format: str,
    quality: str,
) -> dict[str, Any]:
    workflow = plan.get("workflow")
    style_code = compact(getattr(workflow, "style_code", ""))
    generation_rows = list(plan.get("generation_rows") or [])
    generated_paths_by_index = list(plan.get("generated_paths_by_index") or [])
    output_files: list[str] = []
    runs: list[dict[str, Any]] = []
    now = datetime.now(timezone.utc).isoformat()
    for index, row in enumerate(generation_rows):
        row = row if isinstance(row, Mapping) else {}
        paths = generated_paths_by_index[index] if index < len(generated_paths_by_index) else []
        paths = [compact(path) for path in (paths or []) if compact(path)]
        output_files.extend(path for path in paths if path not in output_files)
        prompt = compact(row.get("完整Prompt") or row.get("最终提示词"))
        task_id = compact(row.get("1XM任务ID"))
        runs.append({
            "run_uid": f"{compact(task_run_uid) or 'tmall-chain'}-{style_code or 'style'}-{index + 1}",
            "created_at": now,
            "prompt": prompt,
            "model_key": compact(model) or "gpt-image-2",
            "size": compact(image_size),
            "ratio": compact(ratio),
            "status": "completed" if paths else "failed",
            "task_id": task_id,
            "poll_url": compact(row.get("1XM轮询URL")),
            "image_urls": [
                compact(value)
                for value in str(row.get("生成图URL") or "").splitlines()
                if compact(value)
            ],
            "output_files": paths,
            "input_params": {
                "main_image_path": compact(plan.get("main_origin_path")),
                "reference_image_paths": [
                    compact(path)
                    for path in (row.get("参考图文件") or [])
                    if compact(path)
                ] if isinstance(row.get("参考图文件"), (list, tuple)) else [],
            },
            "error": "" if paths else compact(row.get("备注") or "未生成本地结果"),
        })
    completed = any(run["status"] == "completed" for run in runs)
    first_prompt = next((run["prompt"] for run in runs if run["prompt"]), "")
    output_dir = str(Path(output_files[0]).parent) if output_files else ""
    job = data_sink.create_ai_image_job({
        "title": f"天猫 AI 测图 · {style_code or '未命名款号'}",
        "prompt": first_prompt,
        "model_key": compact(model) or "gpt-image-2",
        "status": "completed" if completed else "failed",
        "output_dir": output_dir,
        "params": {
            "workflow": "tmall_ai_image_test_chain",
            "style_code": style_code,
            "task_run_uid": compact(task_run_uid),
            "size": compact(image_size),
            "ratio": compact(ratio),
            "output_format": compact(output_format),
            "quality": compact(quality),
            "main_image_path": compact(plan.get("main_origin_path")),
        },
        "summary": {
            "ok": completed,
            "runs": runs,
            "output_files": output_files,
            "image_urls": [
                url
                for run in runs
                for url in run.get("image_urls") or []
            ],
            "source": "tmall_ai_image_test_chain",
            "source_task_run_uid": compact(task_run_uid),
        },
    })
    return job


@dataclass
class WorkflowItem:
    row_no: int
    style_code: str
    item_id: str
    category: str
    gender: str
    prompt_name: str = ""
    skc_code: str = ""


@dataclass
class PromptItem:
    sheet_name: str
    field_name: str
    field_order: int
    size_label: str
    output_format: str
    prompt: str
    female_priority: int
    neutral_priority: int


def normalize_workflow_rows(table: Mapping[str, Any], limit: int | None = None) -> tuple[list[WorkflowItem], list[dict[str, Any]]]:
    rows: list[WorkflowItem] = []
    invalid: list[dict[str, Any]] = []
    for index, row in enumerate(table.get("rows") or [], start=2):
        style_code = row_value(row, ["款号", "编码", "货号", "款号/编码", "styleCode", "spu"])
        item_id = row_value(row, ["ID（用于测图的ID）", "测图ID", "天猫商品ID", "商品ID", "宝贝ID", "itemId", "item_id"])
        if not style_code:
            invalid.append({"表格行号": index, "执行结果": "跳过", "备注": "缺少款号"})
            continue
        rows.append(WorkflowItem(
            row_no=index,
            style_code=style_code,
            item_id=item_id,
            category=row_value(row, ["品类（后期匹配）", "品类", "类目", "提示词分组", "promptSheet"]),
            gender=row_value(row, ["模特性别", "性别", "男女", "gender"]) or "中性",
            prompt_name=row_value(row, ["提示词字段名", "提示词名称", "字段名", "promptName"]),
            skc_code=row_value(row, ["SKC编码", "SKC", "款色编码", "款色号", "最优销售色号", "销售色号", "色号"]),
        ))
        if limit and limit > 0 and len(rows) >= limit:
            break
    return rows, invalid


def table_to_raw_rows(table: Mapping[str, Any]) -> list[list[str]]:
    headers = [compact(header) for header in table.get("headers") or []]
    raw_rows = [headers]
    for row in table.get("rows") or []:
        raw_rows.append([compact((row or {}).get(header)) for header in headers])
    return raw_rows


def normalize_header_row(raw_rows: list[list[str]]) -> tuple[list[str], list[dict[str, str]]]:
    header_index = -1
    for index, row in enumerate(raw_rows):
        keys = [normalize_key(value) for value in row]
        if normalize_key("字段名") in keys and normalize_key("描述内容") in keys:
            header_index = index
            break
    if header_index < 0:
        return [], []
    headers = [compact(value) or f"列{index + 1}" for index, value in enumerate(raw_rows[header_index])]
    rows: list[dict[str, str]] = []
    for raw in raw_rows[header_index + 1:]:
        if not any(compact(value) for value in raw):
            continue
        rows.append({header: compact(raw[index] if index < len(raw) else "") for index, header in enumerate(headers)})
    return headers, rows


def is_truthy(value: object) -> bool:
    return compact(value).lower() in {"1", "true", "yes", "y", "是", "启用", "生成"}


def to_int(value: object, fallback: int = 0) -> int:
    try:
        return int(float(compact(value)))
    except Exception:
        return fallback


def normalize_output_format(value: object, fallback: str = "jpeg") -> str:
    text = compact(value).lower()
    if text in {"jpg", "jpeg"}:
        return "jpeg"
    if text in {"png", "webp"}:
        return text
    return fallback


def read_prompt_library(path: str) -> list[PromptItem]:
    import openpyxl

    workbook_path = Path(path).expanduser()
    if not workbook_path.is_file():
        raise RuntimeError(f"File not found: {workbook_path}")
    wb = openpyxl.load_workbook(workbook_path, read_only=False, data_only=True)
    prompts: list[PromptItem] = []
    try:
        for ws in wb.worksheets:
            raw_rows = [
                [compact(value) for value in row]
                for row in ws.iter_rows(values_only=True)
            ]
            _, parsed_rows = normalize_header_row(raw_rows)

            for index, row in enumerate(parsed_rows, start=1):
                field_name = row_value(row, ["字段名"])
                prompt = row_value(row, ["描述内容", "提示词", "prompt"])
                if not field_name or not prompt:
                    continue
                in_view = row_value(row, ["在当前视图"])
                if in_view and not is_truthy(in_view):
                    continue
                prompts.append(PromptItem(
                    sheet_name=compact(ws.title),
                    field_name=field_name,
                    field_order=to_int(row_value(row, ["字段顺序"]), index),
                    size_label=row_value(row, ["尺寸"]),
                    output_format=normalize_output_format(row_value(row, ["格式"]), "jpeg"),
                    prompt=prompt,
                    female_priority=to_int(row_value(row, ["女性优先度"]), 0),
                    neutral_priority=to_int(row_value(row, ["男性/中性优先度", "男性优先度", "中性优先度"]), 0),
                ))
    finally:
        wb.close()
    return prompts


def prompt_items_from_cloud_templates(templates: object) -> list[PromptItem]:
    source = templates
    if isinstance(source, Mapping):
        source = source.get("templates") or []
    if not isinstance(source, list):
        return []

    prompts: list[PromptItem] = []
    for index, item in enumerate(source, start=1):
        if not isinstance(item, Mapping):
            continue
        field_name = compact(item.get("field_name") or item.get("prompt_name") or item.get("name"))
        prompt_text = compact(item.get("prompt_text") or item.get("prompt") or item.get("content"))
        if not field_name or not prompt_text:
            continue
        priority = to_int(item.get("priority"), index)
        prompts.append(PromptItem(
            sheet_name=compact(item.get("group_name") or item.get("sheet_name") or item.get("category")) or "上装",
            field_name=field_name,
            field_order=to_int(item.get("field_order"), index),
            size_label=compact(item.get("size_label") or item.get("size") or ""),
            output_format=normalize_output_format(item.get("output_format") or item.get("format"), "jpeg"),
            prompt=prompt_text,
            female_priority=to_int(item.get("female_priority"), priority),
            neutral_priority=to_int(item.get("male_neutral_priority") or item.get("neutral_priority"), priority),
        ))
    return prompts


def load_prompt_library_for_args(args: argparse.Namespace) -> list[PromptItem]:
    source = compact(getattr(args, "prompt_source", "local_excel")).lower() or "local_excel"
    if source == "cloud_prompt_library":
        payload_text = compact(getattr(args, "cloud_prompt_templates_json", ""))
        if not payload_text:
            raise RuntimeError("云端 Prompt 库未返回可用模板")
        try:
            payload = json.loads(payload_text)
        except Exception as exc:
            raise RuntimeError("云端 Prompt 库模板 JSON 解析失败") from exc
        prompts = prompt_items_from_cloud_templates(payload)
        if not prompts:
            raise RuntimeError("云端 Prompt 库未解析到可用提示词")
        return prompts
    return read_prompt_library(getattr(args, "prompt_file", ""))


def category_to_prompt_sheet(category: str) -> str:
    text = compact(category)
    if text in PROMPT_SHEETS:
        return text
    if re.search(r"裤|短裤|长裤|打底裤|牛仔裤", text):
        return "下装"
    if re.search(r"裙", text):
        return "短裙"
    if re.search(r"鞋|靴", text):
        return "鞋品"
    if re.search(r"连衣", text):
        return "连衣裙"
    return "上装"


def prompt_priority(prompt: PromptItem, workflow: WorkflowItem) -> tuple[int, int, str]:
    gender = workflow.gender.lower()
    priority = prompt.female_priority if re.search(r"女|female|girl", gender) else prompt.neutral_priority
    normalized = priority if priority > 0 else 9999
    return normalized, prompt.field_order or 9999, prompt.field_name


def is_risky_default_prompt(prompt: PromptItem) -> bool:
    return bool(RISKY_DEFAULT_PROMPT_RE.search(f"{prompt.field_name}\n{prompt.prompt}"))


def select_prompts(workflow: WorkflowItem, prompts: list[PromptItem], limit: int = 1) -> list[PromptItem]:
    prompt_names = set(parse_list(workflow.prompt_name))
    group = category_to_prompt_sheet(workflow.category)
    matched = [
        prompt
        for prompt in prompts
        if prompt.sheet_name == group and (not prompt_names or prompt.field_name in prompt_names)
    ]
    if not matched and prompt_names:
        matched = [prompt for prompt in prompts if prompt.field_name in prompt_names]
    if not matched:
        return []
    if not prompt_names:
        safe_matched = [prompt for prompt in matched if not is_risky_default_prompt(prompt)]
        if safe_matched:
            matched = safe_matched
    return sorted(matched, key=lambda item: prompt_priority(item, workflow))[:max(1, limit)]


def select_prompt(workflow: WorkflowItem, prompts: list[PromptItem]) -> PromptItem | None:
    selected = select_prompts(workflow, prompts, 1)
    return selected[0] if selected else None


def build_prompt_text(prompt: PromptItem, workflow: WorkflowItem) -> str:
    metadata = [
        f"款号={workflow.style_code}",
        f"天猫商品ID={workflow.item_id}" if workflow.item_id else "",
        f"品类={workflow.category}" if workflow.category else "",
        f"性别={workflow.gender}" if workflow.gender else "",
    ]
    suffix = "；".join(item for item in metadata if item)
    text = prompt.prompt
    for key, value in {
        "{{款号}}": workflow.style_code,
        "{{商品ID}}": workflow.item_id,
        "{{品类}}": workflow.category,
        "{{性别}}": workflow.gender,
    }.items():
        text = text.replace(key, value)
    guard = (
        "生成约束：只以参考图中的当前主商品为主体；严格保持主商品颜色、图案、版型、面料不变；"
        "禁止把参考图里的其他颜色、条纹叠放衣物、配件或多款组合当成主商品；禁止换色、混款、增加不存在的款式。"
    )
    body = f"{text}\n\n{guard}"
    return f"{body}\n\n商品属性：{suffix}" if suffix else body


def make_generation_row(
    workflow: WorkflowItem,
    prompt: PromptItem,
    reference_paths: list[str] | tuple[str, ...] | str,
    *,
    image_size: str,
    quality: str,
    output_format: str,
    key_tier: str,
    model: str = "gpt-image-2",
    ratio: str = "",
    prompt_index: int = 1,
    image_count: int = 1,
    run_nonce: str = "",
) -> dict[str, Any]:
    final_prompt = build_prompt_text(prompt, workflow)
    normalized_quality = normalize_quality(quality)
    normalized_image_count = normalize_generation_image_count(image_count, 1)
    ratio = compact(ratio) or image_ratio_label(image_size)
    references = [compact(item) for item in (reference_paths if isinstance(reference_paths, (list, tuple)) else [reference_paths]) if compact(item)]
    key_parts = [
        "tmall_ai_chain",
        safe_token(run_nonce, "run"),
        safe_token(workflow.style_code),
        safe_token(prompt.sheet_name),
        safe_token(prompt.field_name),
        safe_token(prompt_index),
        safe_token(image_size),
        stable_hash(final_prompt),
    ]
    idempotency_key = "_".join(key_parts)[:180]
    return {
        "表格行号": workflow.row_no,
        "款号": workflow.style_code,
        "SKC编码": workflow.skc_code,
        "商品ID": workflow.item_id,
        "品类": workflow.category,
        "性别": workflow.gender,
        "提示词分组": prompt.sheet_name,
        "提示词字段名": prompt.field_name,
        "提示词序号": prompt_index,
        "尺寸": image_size,
        "比例": ratio,
        "格式": output_format,
        "质量": normalized_quality,
        "生成数量": normalized_image_count,
        "参考图文件": "\n".join(references),
        "主参考图文件": references[0] if references else "",
        "细节参考图文件": "\n".join(references[1:]),
        "最终提示词": final_prompt,
        "完整Prompt": final_prompt,
        "1XM任务ID": "",
        "1XM轮询URL": "",
        "生成图URL": "",
        "生成图数量": "",
        "执行结果": "待生成",
        "备注": "",
        "__1xm_generate": True,
        "__1xm_key_tier": key_tier,
        "__1xm_idempotency_key": idempotency_key,
        "__task_run_uid": compact(run_nonce),
        "__1xm_reference_paths": references,
        "__1xm_payload": {
            "model": compact(model) or "gpt-image-2",
            "prompt": final_prompt,
            "size": image_size,
            "ratio": ratio,
            "quality": normalized_quality,
            "output_format": output_format,
            "n": normalized_image_count,
        },
    }


async def get_runner(bridge: CDPBridge, prefix: str, entry_url: str, artifact_dir: Path, timeout: int = 120) -> JSRunner:
    tab = bridge.find_tab(prefix)
    if not tab:
        tab = bridge.new_tab(entry_url)
        await asyncio.sleep(3)
    ws_url = bridge.get_tab_ws_url(tab)
    if not ws_url:
        raise RuntimeError(f"Chrome tab 缺少 WebSocket URL: {prefix}")
    runner = JSRunner(
        ws_url,
        timeout=timeout,
        tab_id=str(tab.get("id") or ""),
        tab_url=str(tab.get("url") or entry_url),
        artifact_dir=str(artifact_dir),
    )
    url = str(tab.get("url") or "")
    if not url.startswith(prefix):
        nav = await runner.navigate(entry_url, wait_seconds=4)
        if not nav.success:
            raise RuntimeError(nav.error or f"无法打开页面: {entry_url}")
    return runner


async def wait_for_semir_api_ready(runner: JSRunner, *, log=None, timeout_ms: int = 20000) -> None:
    result = await runner.evaluate_with_reconnect(
        js_call(SEMIR_READY_JS, {"timeout_ms": timeout_ms}),
        allow_navigation_retry=True,
    )
    if not result.success:
        raise RuntimeError(result.error or "森马云盘 API 未就绪")
    if log:
        meta = (result.data[0] if isinstance(result.data, list) and result.data else {}) or {}
        log(f"[chain] 森马云盘 API 已就绪：挂载点 {meta.get('mountCount') or 0} 个")


async def wait_for_tmall_api_ready(
    runner: JSRunner,
    *,
    log=None,
    timeout_ms: int = TMALL_LOGIN_READY_TIMEOUT_MS,
    poll_interval_seconds: float = TMALL_READY_POLL_INTERVAL_SECONDS,
) -> None:
    timeout_seconds = max(0.0, float(timeout_ms or 0) / 1000.0)
    deadline = time.monotonic() + timeout_seconds
    original_timeout = getattr(runner, "timeout", None)
    last_status: Mapping[str, Any] = {}
    last_error = ""
    attempt = 0
    navigated_after_login = False
    probe_timeout = 20
    try:
        if original_timeout is not None:
            try:
                probe_timeout = min(max(int(float(original_timeout)), 10), 30)
            except Exception:
                probe_timeout = 20
        while True:
            attempt += 1
            if original_timeout is not None:
                try:
                    runner.timeout = probe_timeout
                except Exception:
                    pass
            result = await runner.evaluate_with_reconnect(
                js_call(TMALL_READY_PROBE_JS, {}),
                allow_navigation_retry=True,
            )
            if result.success:
                status = (result.data[0] if isinstance(result.data, list) and result.data else {}) or {}
                if isinstance(status, Mapping):
                    last_status = status
                if bool(status.get("ready")):
                    if log:
                        title = compact(status.get("title"))
                        href = compact(status.get("href"))
                        log(f"[chain] 天猫测图页面已就绪：{title or href or 'material-test'}")
                    return

                if (
                    not navigated_after_login
                    and bool(status.get("hasMtop"))
                    and not bool(status.get("isLoginUrl"))
                    and not bool(status.get("isMaterialUrl"))
                    and hasattr(runner, "navigate")
                ):
                    navigated_after_login = True
                    nav = await runner.navigate(TMALL_URL, wait_seconds=4)
                    if not nav.success:
                        last_error = compact(nav.error) or "切换到天猫测图页面失败"
                    elif log:
                        log("[chain] 天猫已登录，正在切换到测图页面")

                if log and (attempt == 1 or attempt % 10 == 0):
                    href = compact(status.get("href"))
                    title = compact(status.get("title"))
                    if bool(status.get("isLoginUrl")):
                        log(f"[chain] 等待天猫登录（最多 {int(timeout_seconds)} 秒）：当前页面 {title or href or '登录页'}")
                    elif not bool(status.get("hasMtop")):
                        log(f"[chain] 等待天猫测图页面加载 mtop：当前页面 {title or href or '未知页面'}")
                    else:
                        log(f"[chain] 等待天猫测图页面就绪：当前页面 {title or href or '未知页面'}")
            else:
                last_error = compact(result.error) or "天猫页面探针执行失败"
                if log and (attempt == 1 or attempt % 10 == 0):
                    log(f"[chain] 等待天猫测图页面就绪：{last_error}")

            remaining = deadline - time.monotonic()
            if remaining <= 0:
                href = compact(last_status.get("href")) if isinstance(last_status, Mapping) else ""
                title = compact(last_status.get("title")) if isinstance(last_status, Mapping) else ""
                if isinstance(last_status, Mapping) and bool(last_status.get("isLoginUrl")):
                    state = "当前仍停留在天猫登录页"
                elif isinstance(last_status, Mapping) and not bool(last_status.get("hasMtop")):
                    state = "当前页面还未加载天猫 mtop API"
                else:
                    state = "当前天猫测图页面还未就绪"
                detail = title or href or last_error or "未获取到页面状态"
                raise RuntimeError(
                    f"等待天猫登录超时（{int(timeout_seconds)} 秒）：{state}。"
                    f"当前页面：{detail}。请在 Chrome 9222 的天猫页面完成登录后重试。"
                )
            sleep_seconds = min(max(0.0, float(poll_interval_seconds or 0)), remaining)
            await asyncio.sleep(sleep_seconds)
    finally:
        if original_timeout is not None:
            try:
                runner.timeout = original_timeout
            except Exception:
                pass


def js_call(body: str, payload: Mapping[str, Any]) -> str:
    return f"""
(async () => {{
  const __payload = {json.dumps(payload, ensure_ascii=False)};
{body}
}})()
"""


SEMIR_FIND_JS = r"""
  const SEARCH_SCOPE = '["filename", "tag"]';
  const IMAGE_EXTS = new Set(['jpg', 'jpeg', 'png', 'webp', 'bmp', 'gif', 'tif', 'tiff']);
  const BLOCKED_EXTS = new Set(['psd', 'pdf', 'ai', 'cdr', 'zip', 'rar', '7z']);
  const compact = (value) => String(value || '').replace(/\s+/g, ' ').trim();
  const normalizeUrl = (value) => {
    const url = compact(value);
    if (!url) return '';
    return url.startsWith('//') ? `https:${url}` : url;
  };
  const semirApiUrl = (value) => {
    const url = compact(value);
    if (!url || /^https?:\/\//i.test(url)) return url;
    if (url.startsWith('/')) return `https://fmp.semirapp.com${url}`;
    return url;
  };
  const sleep = (ms) => new Promise((resolve) => setTimeout(resolve, Math.max(0, Number(ms || 0))));
  const extOf = (item) => {
    const explicit = compact(item?.ext).toLowerCase().replace(/^\./, '');
    if (explicit) return explicit;
    const name = compact(item?.filename || item?.name || item?.fullpath || '');
    const index = name.lastIndexOf('.');
    return index >= 0 ? name.slice(index + 1).toLowerCase() : '';
  };
  const isDir = (item) => item?.dir === 1 || item?.dir === '1' || item?.dir === true || item?.type === 'folder';
  const isImage = (item) => !isDir(item) && IMAGE_EXTS.has(extOf(item)) && !BLOCKED_EXTS.has(extOf(item));
  const naturalCompare = (a, b) => String(a || '').localeCompare(String(b || ''), 'zh-Hans-CN', { numeric: true, sensitivity: 'base' });
  async function fetchJson(url, init = {}) {
    const requestUrl = semirApiUrl(url);
    const errors = [];
    for (let attempt = 1; attempt <= 5; attempt += 1) {
      try {
        const response = await fetch(requestUrl, { credentials: 'include', ...init });
        const text = await response.text();
        let payload = null;
        try { payload = text ? JSON.parse(text) : {}; } catch (error) { payload = null; }
        if (!response.ok) throw new Error(`HTTP ${response.status}: ${text.slice(0, 240)}`);
        if (!payload) throw new Error(`接口未返回 JSON`);
        return payload;
      } catch (error) {
        errors.push(String(error?.message || error));
        if (attempt >= 5) break;
        await sleep(350 * attempt);
      }
    }
    throw new Error(`森马云盘接口请求失败：${requestUrl || url}；${errors.join('；')}`);
  }
  function parseCloudPath(raw) {
    const value = compact(raw);
    if (!value) return null;
    const divider = value.indexOf('//');
    if (divider < 0) throw new Error('云盘路径格式需要“挂载点//目录/子目录”');
    return {
      mountName: compact(value.slice(0, divider)),
      relativePath: value.slice(divider + 2).replace(/\\/g, '/').split('/').map(compact).filter(Boolean).join('/'),
    };
  }
  async function fetchMounts() {
    const payload = await fetchJson('/fengcloud/1/account/mount');
    return Array.isArray(payload) ? payload : (payload.list || payload?.data?.list || []);
  }
  async function searchFiles(mountId, keyword) {
    const all = [];
    let start = 0;
    while (true) {
      const body = new URLSearchParams({
        size: '100',
        start: String(start),
        keyword: String(keyword || ''),
        mount_id: String(mountId || ''),
        scope: SEARCH_SCOPE,
      });
      const payload = await fetchJson('/fengcloud/2/file/search', {
        method: 'POST',
        headers: { 'Content-Type': 'application/x-www-form-urlencoded' },
        body: body.toString(),
      });
      const items = Array.isArray(payload?.list) ? payload.list : (payload?.data?.list || []);
      const total = Number(payload?.total || payload?.data?.total || items.length || 0);
      all.push(...items);
      start += items.length;
      if (!items.length || start >= total) break;
    }
    return all;
  }
  function extractFolderItems(payload) {
    if (Array.isArray(payload)) return payload;
    if (Array.isArray(payload?.list)) return payload.list;
    if (Array.isArray(payload?.items)) return payload.items;
    if (Array.isArray(payload?.files)) return payload.files;
    if (Array.isArray(payload?.data)) return payload.data;
    if (Array.isArray(payload?.data?.list)) return payload.data.list;
    if (Array.isArray(payload?.data?.items)) return payload.data.items;
    return null;
  }
  function extractFolderTotal(payload, fallbackCount) {
    const candidates = [payload?.total, payload?.count, payload?.data?.total, payload?.data?.count];
    const total = candidates.map(Number).find((value) => Number.isFinite(value) && value >= 0);
    return total == null ? fallbackCount : total;
  }
  function normalizeListedItem(item, parentFullpath) {
    if (!item || typeof item !== 'object') return item;
    const filename = compact(item.filename || item.name || '');
    const fullpath = compact(item.fullpath || item.path || '');
    if (fullpath || !filename || !parentFullpath) return item;
    return { ...item, filename, fullpath: `${String(parentFullpath || '').replace(/\/+$/, '')}/${filename}` };
  }
  async function fetchFolderPage(mountId, fullpath, start, method, endpoint) {
    const query = new URLSearchParams({
      order: 'filename asc',
      size: '100',
      start: String(start),
      mount_id: String(mountId || ''),
      fullpath: String(fullpath || ''),
      path: String(fullpath || ''),
      current: '1',
    });
    if (method === 'GET') return fetchJson(`${endpoint}?${query.toString()}`);
    return fetchJson(endpoint, {
      method,
      headers: { 'Content-Type': 'application/x-www-form-urlencoded' },
      body: query.toString(),
    });
  }
  async function listFolderItems(mountId, fullpath) {
    const attempts = [
      { method: 'GET', endpoint: '/fengcloud/1/file/ls' },
      { method: 'GET', endpoint: '/fengcloud/2/file/list' },
      { method: 'POST', endpoint: '/fengcloud/2/file/list' },
      { method: 'GET', endpoint: '/fengcloud/1/file/list' },
      { method: 'POST', endpoint: '/fengcloud/1/file/list' },
    ];
    const errors = [];
    for (const attempt of attempts) {
      try {
        const all = [];
        let start = 0;
        let total = null;
        while (true) {
          const payload = await fetchFolderPage(mountId, fullpath, start, attempt.method, attempt.endpoint);
          const itemsRaw = extractFolderItems(payload);
          if (!Array.isArray(itemsRaw)) throw new Error(`${attempt.method} ${attempt.endpoint} 未返回列表字段`);
          const items = itemsRaw.map((item) => normalizeListedItem(item, fullpath));
          all.push(...items);
          const pageTotal = extractFolderTotal(payload, start + items.length);
          if (total == null) total = pageTotal;
          if (!items.length) break;
          start += items.length;
          if (start >= pageTotal) break;
        }
        return { ok: true, items: all };
      } catch (error) {
        errors.push(String(error?.message || error));
      }
    }
    return { ok: false, items: [], error: errors[0] || '列目录失败' };
  }
  async function collectDescendantImages(mountId, folderPath, maxDepth, remainingBudget = { value: 600 }) {
    if (!folderPath || maxDepth < 0 || remainingBudget.value <= 0) return { assets: [], errors: [] };
    const listed = await listFolderItems(mountId, folderPath);
    if (!listed.ok) return { assets: [], errors: [`${folderPath}: ${listed.error}`] };
    const assets = [];
    const errors = [];
    for (const item of listed.items) {
      if (remainingBudget.value <= 0) break;
      if (isDir(item)) {
        if (maxDepth <= 0) continue;
        const child = await collectDescendantImages(mountId, item?.fullpath || item?.filename || '', maxDepth - 1, remainingBudget);
        assets.push(...child.assets);
        errors.push(...child.errors);
        continue;
      }
      if (!isImage(item)) continue;
      assets.push(item);
      remainingBudget.value -= 1;
    }
    return { assets, errors };
  }
  function pathSegments(fullpath) {
    return String(fullpath || '').replace(/\\/g, '/').split('/').map(compact).filter(Boolean);
  }
  function parentFullpath(fullpath) {
    const segments = pathSegments(fullpath);
    segments.pop();
    return segments.join('/');
  }
  function sameParentFullpath(left, right) {
    const leftParent = parentFullpath(left?.fullpath || left?.path || '');
    const rightParent = parentFullpath(right?.fullpath || right?.path || '');
    return Boolean(leftParent && rightParent && leftParent === rightParent);
  }
  function folderPathsFromSearchItem(item, styleCode) {
    const style = compact(styleCode);
    if (!style) return [];
    const fullpath = String(item?.fullpath || item?.filename || '').replace(/\\/g, '/');
    const segments = pathSegments(fullpath);
    const paths = [];
    const addPath = (value) => {
      const path = compact(value);
      if (path && !paths.includes(path)) paths.push(path);
    };
    if (isDir(item) && fullpath.includes(style)) {
      addPath(fullpath);
    } else if (isImage(item) && parentFullpath(fullpath).includes(style)) {
      addPath(parentFullpath(fullpath));
    }
    for (let index = segments.length - 1; index >= 0; index -= 1) {
      if (!segments[index].includes(style)) continue;
      if (isImage(item) && index === segments.length - 1) {
        const parent = parentFullpath(fullpath);
        if (parent.includes(style)) addPath(parent);
      } else {
        addPath(segments.slice(0, index + 1).join('/'));
      }
    }
    return paths;
  }
  async function fileInfo(mountId, fullpath) {
    const query = new URLSearchParams({ mount_id: String(mountId || ''), fullpath: String(fullpath || '') });
    return fetchJson(`/fengcloud/2/file/info?${query.toString()}`);
  }
  function findFirstRemoteUrl(value, seen = new Set()) {
    if (!value) return '';
    if (typeof value === 'string') {
      const direct = normalizeUrl(value);
      return /^https?:\/\//i.test(direct) || /^\/\//.test(value) ? direct : '';
    }
    if (typeof value !== 'object' || seen.has(value)) return '';
    seen.add(value);
    if (Array.isArray(value)) {
      for (const item of value) {
        const found = findFirstRemoteUrl(item, seen);
        if (found) return found;
      }
      return '';
    }
    for (const key of ['url', 'picUrl', 'imageUrl', 'materialUrl', 'downloadUrl', 'fileUrl', 'src']) {
      const found = findFirstRemoteUrl(value[key], seen);
      if (found) return found;
    }
    for (const item of Object.values(value)) {
      const found = findFirstRemoteUrl(item, seen);
      if (found) return found;
    }
    return '';
  }
  function basenameOf(value) {
    return compact(value).replace(/\\/g, '/').split('/').filter(Boolean).pop() || compact(value);
  }
  function startsWithCodeToken(value, code) {
    const text = compact(value).toLowerCase();
    const target = compact(code).toLowerCase();
    if (!text || !target) return false;
    const escaped = target.replace(/[.*+?^${}()|[\]\\]/g, '\\$&');
    return new RegExp(`^${escaped}(?:$|[^0-9a-z])`, 'i').test(text);
  }
  function isShowcaseOneName(filename) {
    return /橱窗\s*0?1(?!\d)/i.test(compact(filename));
  }
  function isYzOneName(filename) {
    const name = compact(filename);
    return isYzOneBracketName(name) || isYzOneBareName(name);
  }
  function isYzOneBracketName(filename) {
    const name = compact(filename);
    return /(^|[^0-9a-z])yz\s*[\(（]\s*0?1\s*[\)）]/i.test(name);
  }
  function isYzOneBareName(filename) {
    const name = compact(filename);
    return /(^|[^0-9a-z])yz\s*0?1(?!\d)/i.test(name);
  }
  function isCombinedYzOneAiOriginalName(filename) {
    const name = basenameOf(filename);
    return isYzOneBracketName(name)
      && /&\s*o\s*[\(（]\s*0?1\s*[\)）]\s*[-_ ]?\s*ai(?:$|[-_ .])/i.test(name);
  }
  function isDerivedMainReferenceName(filename) {
    const name = basenameOf(filename);
    if (isCombinedYzOneAiOriginalName(name)) return false;
    return /颜色\s*ai|ai\s*[-_ ]?\d|[-_&]ai(?:$|[-_ .])|效果图|创意图|生图|生成图|生成款|改色|换色|变色|智能图/i.test(name);
  }
  function isDerivedReferenceName(filename) {
    const name = basenameOf(filename);
    return /AI\s*参考|颜色\s*ai|ai\s*[-_ ]?\d|[-_&]ai(?:$|[-_ .])|效果图|创意图|生图|生成图|生成款|改色|换色|变色|智能图/i.test(name);
  }
  function isMainReferenceCandidate(filename) {
    const name = basenameOf(filename);
    return !isDerivedMainReferenceName(name) && (isShowcaseOneName(name) || isYzOneName(name));
  }
  function isSelectedModelOriginalPath(fullpath) {
    const full = compact(fullpath);
    return /模拍原图/.test(full) && /已选/.test(full);
  }
  function folderPathScanRank(fullpath) {
    const full = compact(fullpath);
    return (isSelectedModelOriginalPath(full) ? 100 : 0) + (/已选/.test(full) ? 20 : 0) + (/模拍原图/.test(full) ? 10 : 0);
  }
  function isStyleColorReferenceName(filename, styleCode) {
    const style = compact(styleCode);
    if (!style) return false;
    const escaped = style.replace(/[.*+?^${}()|[\]\\]/g, '\\$&');
    return new RegExp(`^${escaped}[-_][0-9]{4,6}(?:[-_][0-9]{1,4})?(?:\\.[^.]+)?$`, 'i').test(basenameOf(filename));
  }
  function skcCodeFromStyleName(filename, styleCode) {
    const style = compact(styleCode);
    if (!style) return '';
    const escaped = style.replace(/[.*+?^${}()|[\]\\]/g, '\\$&');
    const match = basenameOf(filename).match(new RegExp(`^(${escaped}[-_][0-9]{5})(?=$|[^0-9])`, 'i'));
    return match ? match[1] : '';
  }
  function isSkcFlatCandidate(filename, skcCode, styleCode) {
    const skc = compact(skcCode);
    if (skc && startsWithCodeToken(basenameOf(filename), skc)) return true;
    return Boolean(skcCodeFromStyleName(filename, styleCode));
  }
  function referenceCandidate(item, styleCode, skcCode) {
    const filename = compact(item?.filename || item?.name);
    const fullpath = compact(item?.fullpath || item?.path);
    const full = `${fullpath}/${filename}`.replace(/\\/g, '/');
    if (!isImage({ ...item, filename })) return null;
    if (/尺码表|尺寸表|规格表|吊牌|制单|pdf|psd|源文件|视觉推荐|视频|模特卡/i.test(full)) return null;
    const baseName = basenameOf(filename || fullpath);
    const derivedReference = isDerivedReferenceName(baseName);
    const stylePathOk = !compact(styleCode) || full.includes(compact(styleCode));
    const flags = {
      showcase1: isShowcaseOneName(baseName),
      yz1Bracket: isYzOneBracketName(baseName),
      yz1Bare: isYzOneBareName(baseName),
      yz1: isYzOneName(baseName),
      flat: isSkcFlatCandidate(baseName, skcCode, styleCode),
      skcExact: compact(skcCode) ? startsWithCodeToken(baseName, skcCode) : false,
      styleColor: isStyleColorReferenceName(baseName, styleCode),
      flatPath: /平拍|平铺|白底|静物|正面|front/i.test(full),
      selectedPath: /已选/.test(full),
      modelPath: /模拍/.test(full),
      selectedModelOriginalPath: isSelectedModelOriginalPath(full),
      front: /正面|前面|front/i.test(baseName),
      back: /背面|后面|back/i.test(baseName),
    };
    flags.selectedModelStyleColor = flags.selectedModelOriginalPath && flags.styleColor && stylePathOk;
    const derivedMain = isDerivedMainReferenceName(baseName);
    const isMain = (flags.showcase1 || flags.yz1) && stylePathOk && !derivedMain;
    const isDetail = (flags.selectedModelStyleColor || flags.flat) && !derivedReference;
    if (!isMain && !isDetail) return null;
    const mainRank = (
      flags.yz1Bracket ? 30
        : flags.yz1Bare ? 20
          : flags.showcase1 ? 10
            : 0
    ) + (flags.selectedPath ? 2 : 0) + (flags.modelPath ? 1 : 0);
    const detailRank = (
      (flags.selectedModelStyleColor ? 100 : 0)
      + (flags.skcExact ? 30 : 0)
      + (flags.styleColor ? 40 : 0)
      + (flags.selectedModelOriginalPath ? 20 : 0)
      + (flags.flatPath ? 8 : 0)
      + (flags.front ? 2 : flags.back ? 0 : 1)
    );
    const role = flags.yz1Bracket ? 'origin_yz1_bracket'
        : flags.yz1Bare ? 'origin_yz1_bare'
          : flags.showcase1 ? 'origin_showcase1'
            : flags.selectedModelStyleColor ? 'detail_selected_model_style_color'
              : 'detail_skc_flat';
    const score = isMain ? 300 + mainRank * 20 : 200 + detailRank * 10;
    const resolvedSkcCode = flags.flat ? (compact(skcCode) || skcCodeFromStyleName(baseName, styleCode)) : '';
    return { ...item, filename, fullpath, ext: extOf({ ...item, filename }), score, mainRank, detailRank, role, flags: { ...flags, derivedMain, derivedReference }, skcCode: resolvedSkcCode };
  }
  try {
    const styleCode = compact(__payload.style_code);
    const configured = parseCloudPath(__payload.cloud_path || '');
    const mounts = await fetchMounts();
    const mountPlans = [];
    if (configured) {
      const mount = mounts.find((item) => compact(item.org_name || item.name) === configured.mountName);
      if (!mount) throw new Error(`未找到挂载点：${configured.mountName}`);
      mountPlans.push({
        mountId: String(mount.mount_id || mount.id || ''),
        mountName: compact(mount.org_name || mount.name),
        relativePath: configured.relativePath,
        source: 'configured',
      });
    } else {
      for (const mount of mounts) {
        mountPlans.push({
          mountId: String(mount.mount_id || mount.id || ''),
          mountName: compact(mount.org_name || mount.name),
          relativePath: '',
          source: 'mounted-global',
        });
      }
    }
    const searches = [];
    const allItems = [];
    const skcCode = compact(__payload.skc_code || '');
    const keywords = [...new Set([styleCode, skcCode].map(compact).filter(Boolean))];
    const scanErrors = [];
    for (const mount of mountPlans.filter((item) => item.mountId)) {
      for (const keyword of keywords) {
        const found = await searchFiles(mount.mountId, keyword);
        const scoped = mount.relativePath ? found.filter((item) => String(item.fullpath || '').replace(/\\/g, '/').startsWith(mount.relativePath)) : found;
        searches.push({ mountId: mount.mountId, mountName: mount.mountName, source: mount.source, keyword, found: found.length, scoped: scoped.length, relativePath: mount.relativePath });
        allItems.push(...scoped.map((item) => ({ ...item, mountId: mount.mountId, mountName: mount.mountName })));
        const folderPaths = [];
        for (const item of scoped) {
          for (const folderPath of folderPathsFromSearchItem(item, styleCode)) {
            if (folderPath && !folderPaths.includes(folderPath)) folderPaths.push(folderPath);
          }
        }
        folderPaths.sort((a, b) => folderPathScanRank(b) === folderPathScanRank(a) ? naturalCompare(a, b) : folderPathScanRank(b) - folderPathScanRank(a));
        for (const folderPath of folderPaths.slice(0, 8)) {
          const descendants = await collectDescendantImages(mount.mountId, folderPath, Number(__payload.folder_scan_depth || 3));
          scanErrors.push(...descendants.errors);
          allItems.push(...descendants.assets.map((item) => ({ ...item, mountId: mount.mountId, mountName: mount.mountName })));
        }
      }
    }
    const seen = new Set();
    const ranked = [];
    for (const item of allItems) {
      const candidate = referenceCandidate(item, styleCode, skcCode);
      const key = compact(candidate?.fullpath || candidate?.filename);
      if (!candidate || !key || seen.has(key)) continue;
      seen.add(key);
      ranked.push(candidate);
    }
    ranked.sort((a, b) => b.score === a.score ? naturalCompare(a.fullpath || a.filename, b.fullpath || b.filename) : b.score - a.score);
    const mainRanked = ranked
      .filter((item) => item.flags.showcase1 || item.flags.yz1)
      .sort((a, b) => b.mainRank === a.mainRank ? naturalCompare(a.fullpath || a.filename, b.fullpath || b.filename) : b.mainRank - a.mainRank);
    const chosenMain = mainRanked[0] || null;
    const detailRanked = ranked
      .filter((item) => chosenMain && item.flags.styleColor && sameParentFullpath(item, chosenMain))
      .filter((item) => !chosenMain || compact(item.fullpath || item.filename) !== compact(chosenMain.fullpath || chosenMain.filename))
      .sort((a, b) => b.detailRank === a.detailRank ? naturalCompare(a.fullpath || a.filename, b.fullpath || b.filename) : b.detailRank - a.detailRank);
    const chosenDetail = detailRanked[0] || null;
    const candidates = ranked.slice(0, Number(__payload.limit || 8));
    const infoTargets = [];
    for (const candidate of [chosenMain, ...detailRanked, ...candidates.slice(0, 4)]) {
      if (!candidate) continue;
      const key = compact(candidate.fullpath || candidate.filename);
      if (!key || infoTargets.some((item) => compact(item.fullpath || item.filename) === key)) continue;
      infoTargets.push(candidate);
    }
    for (const candidate of infoTargets) {
      try {
        candidate.materialUrl = findFirstRemoteUrl(await fileInfo(candidate.mountId, candidate.fullpath));
      } catch (error) {
        candidate.materialUrl = '';
        candidate.infoError = String(error?.message || error);
      }
    }
    return { success: true, data: [{ styleCode, skcCode, searches, scanErrors: scanErrors.slice(0, 10), candidates, mainCandidates: mainRanked.slice(0, 5), detailCandidates: detailRanked, chosen: chosenMain, chosenMain, chosenDetail }], meta: { has_more: false } };
  } catch (error) {
    return { success: false, error: String(error?.message || error) };
  }
"""


SEMIR_READY_JS = r"""
  const sleep = (ms) => new Promise((resolve) => setTimeout(resolve, Math.max(0, Number(ms || 0))));
  const timeoutMs = Math.max(1000, Number(__payload.timeout_ms || 20000));
  const deadline = Date.now() + timeoutMs;
  const errors = [];
  for (let attempt = 1; ; attempt += 1) {
    try {
      const response = await fetch('https://fmp.semirapp.com/fengcloud/1/account/mount', { credentials: 'include' });
      const text = await response.text();
      let payload = null;
      try { payload = text ? JSON.parse(text) : {}; } catch (error) { payload = null; }
      if (!response.ok) throw new Error(`HTTP ${response.status}: ${text.slice(0, 160)}`);
      const mounts = Array.isArray(payload) ? payload : (payload?.list || payload?.data?.list || []);
      if (!Array.isArray(mounts) || !mounts.length) throw new Error('mount API 未返回挂载点');
      return {
        success: true,
        data: [{ url: location.href, origin: location.origin, mountCount: mounts.length }],
        meta: { has_more: false },
      };
    } catch (error) {
      errors.push(String(error?.message || error));
      if (Date.now() >= deadline) break;
      await sleep(Math.min(1200, 250 * attempt));
    }
  }
  return {
    success: false,
    error: `森马云盘 API 未就绪：${errors.slice(-3).join('；')}`,
  };
"""


TMALL_READY_PROBE_JS = r"""
  const compact = (value) => String(value || '').replace(/\s+/g, ' ').trim();
  const href = compact(location.href);
  const host = compact(location.hostname).toLowerCase();
  const title = compact(document.title);
  const bodyText = compact(document.body?.innerText || '').slice(0, 500);
  const hasMtop = !!(window.lib && window.lib.mtop && typeof window.lib.mtop.request === 'function');
  const isLoginUrl = (
    host === 'login.taobao.com'
    || host === 'loginmyseller.taobao.com'
    || host.endsWith('.login.taobao.com')
    || href.includes('login.taobao.com')
    || href.includes('loginmyseller.taobao.com')
  );
  const isMaterialUrl = (
    host === 'myseller.taobao.com'
    && href.includes('/material-center/material-test')
  );
  const hasLoginText = /登录|扫码登录|密码登录|验证码|安全验证/.test(bodyText);
  return {
    success: true,
    data: [{
      ready: Boolean(hasMtop && isMaterialUrl && !isLoginUrl),
      hasMtop,
      isLoginUrl: Boolean(isLoginUrl || (hasLoginText && !hasMtop)),
      isMaterialUrl,
      href,
      title,
      bodySample: bodyText,
    }],
    meta: { has_more: false },
  };
"""


TMALL_UPLOAD_CREATE_JS = r"""
  const compact = (value) => String(value || '').replace(/\s+/g, ' ').trim();
  function describeError(error) {
    if (!error) return '未知错误';
    if (typeof error === 'string') return error;
    if (error.message) return String(error.message);
    if (Array.isArray(error.ret)) return error.ret.join('；');
    try { return JSON.stringify(error).slice(0, 1000); } catch (jsonError) {}
    return String(error);
  }
  const normalizeUrl = (value) => {
    const url = compact(value);
    if (!url) return '';
    return url.startsWith('//') ? `https:${url}` : url;
  };
  function getCookieValue(name) {
    const key = `${String(name || '')}=`;
    for (const raw of String(document?.cookie || '').split(';')) {
      const item = raw.trim();
      if (item.startsWith(key)) return decodeURIComponent(item.slice(key.length));
    }
    return '';
  }
  async function dataUrlToBlob(dataUrl) {
    if (typeof fetch !== 'function') throw new Error('当前页面不支持 data URL 转 Blob');
    return (await fetch(String(dataUrl || ''))).blob();
  }
  function findFirstRemoteUrl(value, seen = new Set()) {
    if (!value) return '';
    if (typeof value === 'string') {
      const direct = normalizeUrl(value);
      return /^https?:\/\//i.test(direct) || /^\/\//.test(value) ? direct : '';
    }
    if (typeof value !== 'object' || seen.has(value)) return '';
    seen.add(value);
    if (Array.isArray(value)) {
      for (const item of value) {
        const found = findFirstRemoteUrl(item, seen);
        if (found) return found;
      }
      return '';
    }
    for (const key of ['url', 'picUrl', 'imageUrl', 'materialUrl', 'downloadUrl', 'fileUrl', 'src']) {
      const found = findFirstRemoteUrl(value[key], seen);
      if (found) return found;
    }
    for (const item of Object.values(value)) {
      const found = findFirstRemoteUrl(item, seen);
      if (found) return found;
    }
    return '';
  }
  function inferMaterialSize(item) {
    const raw = compact(item?.materialSize || item?.size || item?.pixel || item?.raw?.object?.pix);
    const match = raw.match(/(\d+)\D+(\d+)/);
    if (match) {
      const width = Number(match[1]);
      const height = Number(match[2]);
      if (width > 0 && height > 0) {
        const ratio = width / height;
        if (Math.abs(ratio - 1) <= 0.08) return '1:1';
        if (Math.abs(ratio - 0.75) <= 0.10 || Math.abs(ratio - (2 / 3)) <= 0.10) return '3:4';
        return ratio > 1 ? '1:1' : '3:4';
      }
    }
    return '3:4';
  }
  const sleep = (ms) => new Promise((resolve) => setTimeout(resolve, Math.max(0, Number(ms || 0))));
  function describeMtopError(error) {
    const parts = [];
    const ret = Array.isArray(error?.ret) ? error.ret.join('|') : '';
    const dataMsg = compact(error?.data?.errorMsg || error?.data?.message || error?.data?.msg);
    const message = compact(error?.message || error);
    if (ret) parts.push(ret);
    if (dataMsg && !parts.includes(dataMsg)) parts.push(dataMsg);
    if (message && message !== '[object Object]' && !parts.includes(message)) parts.push(message);
    return parts.join('|') || String(error || 'MTop 调用失败');
  }
  function extractRows(value) {
    for (const rows of [
      value?.result?.list,
      value?.list,
      value?.records,
      value?.data?.result?.list,
      value?.data?.list,
      value?.data?.records,
      value?.modelDataList,
    ]) {
      if (Array.isArray(rows)) return rows;
    }
    return Array.isArray(value) ? value : [];
  }
  function metricItems(value) {
    const result = [];
    for (const list of Object.values(value || {})) {
      if (Array.isArray(list)) result.push(...list);
    }
    return result;
  }
  function materialIdsFromMetrics(value) {
    const ids = [];
    const seen = new Set();
    const pushId = (id) => {
      const text = compact(id);
      if (!text || seen.has(text)) return;
      seen.add(text);
      ids.push(text);
    };
    for (const item of metricItems(value)) {
      pushId(item?.imageId);
      for (const id of Array.isArray(item?.imageIds) ? item.imageIds : []) pushId(id);
      for (const op of Array.isArray(item?.materialOpParams) ? item.materialOpParams : []) {
        for (const material of Array.isArray(op?.materials) ? op.materials : []) pushId(material?.imageId);
      }
    }
    return ids;
  }
  function taskDataItems(readback, targetSource = 'common_search') {
    const rows = Array.isArray(readback?.rows) ? readback.rows : [];
    const target = compact(targetSource).toLowerCase();
    const result = [];
    for (const row of rows) {
      const itemId = compact(row?.domainId || row?.itemId || row?.id);
      const dataList = row?.columns?.test_data?.dataList || row?.test_data?.dataList || row?.testData?.dataList || [];
      for (const item of Array.isArray(dataList) ? dataList : []) {
        const source = compact(item?.imageTestSource || item?.source || item?.testChannel).toLowerCase();
        if (source !== target) continue;
        const experimentTaskId = compact(item?.experimentTaskId || item?.taskId || item?.id);
        if (!experimentTaskId) continue;
        result.push({
          itemId,
          experimentTaskId,
          source: target,
          status: compact(item?.testStatus ?? item?.status),
          originMaterialIds: materialIdsFromMetrics(item?.originImageMetrics || {}),
          testMaterialIds: materialIdsFromMetrics(item?.testImageMetrics || {}),
          raw: item,
        });
      }
    }
    return result;
  }
  async function uploadDataUrl(dataUrl, name) {
    const query = new URLSearchParams({
      appkey: 'tu',
      folderId: String(__payload.folder_id || '0'),
      watermark: 'false',
      picCompress: 'true',
      _input_charset: 'utf-8',
    });
    const form = new FormData();
    form.append('water', 'false');
    form.append('name', name);
    form.append('_tb_token_', getCookieValue('_tb_token_'));
    form.append('file', await dataUrlToBlob(dataUrl), name);
    const response = await fetch(`https://stream-upload.taobao.com/api/upload.api?${query.toString()}`, {
      method: 'POST',
      credentials: 'include',
      body: form,
    });
    const text = await response.text();
    let payload = null;
    try { payload = text ? JSON.parse(text) : {}; } catch (error) { payload = null; }
    if (!response.ok) throw new Error(`图片上传 HTTP ${response.status}: ${text.slice(0, 240)}`);
    if (!payload || payload.success === false) throw new Error(payload?.message || payload?.msg || `图片上传失败：${name}`);
    const url = normalizeUrl(payload?.object?.url || findFirstRemoteUrl(payload));
    if (!url) throw new Error(`图片上传未返回 URL：${name}`);
    return {
      name,
      url,
      fileId: compact(payload?.object?.fileId),
      folderId: compact(payload?.object?.folderId),
      pixel: compact(payload?.object?.pix),
      raw: payload,
    };
  }
  async function callMtop(api, data, options = {}) {
    const client = window.lib?.mtop || window.mtop;
    if (!client || typeof client.request !== 'function') {
      throw new Error('未找到千牛 MTop 客户端，请确认当前 tab 是天猫素材测试页');
    }
    let payload = null;
    try {
      payload = await client.request({
        api,
        v: options.v || '1.0',
        type: options.type || 'POST',
        dataType: 'json',
        H5Request: true,
        preventFallback: true,
        data,
      });
    } catch (error) {
      throw new Error(`${api} 返回失败：${describeMtopError(error)}`);
    }
    if (Array.isArray(payload?.ret)) {
      const failed = payload.ret.find((item) => !/^SUCCESS/i.test(String(item || '')));
      if (failed) throw new Error(`${api} 返回失败：${payload.ret.join('；')}`);
    }
    return payload?.data !== undefined ? payload.data : payload;
  }
  function extractTaskStatusList(value, targetSource = 'common_search') {
    const result = [];
    const seen = new Set();
    const visit = (node) => {
      if (!node || typeof node !== 'object' || seen.has(node)) return;
      seen.add(node);
      if (Array.isArray(node)) {
        node.forEach(visit);
        return;
      }
      const taskId = compact(node.experimentTaskId || node.taskId || node.id);
      if (taskId) {
        const source = compact(node.source || node.imageTestSource || node.testChannel || targetSource).toLowerCase();
        if (!source || source === targetSource || source === 'common_search' || source === 'commonsearch') {
          result.push({ experimentTaskId: taskId, source: targetSource });
        }
      }
      Object.values(node).forEach(visit);
    };
    visit(value);
    const unique = [];
    const uniqueIds = new Set();
    for (const item of result) {
      if (uniqueIds.has(item.experimentTaskId)) continue;
      uniqueIds.add(item.experimentTaskId);
      unique.push(item);
    }
    return unique;
  }
  function summarizeTask(readback, taskId, targetSource = 'common_search') {
    const rows = Array.isArray(readback?.rows) ? readback.rows : [];
    const targetTaskId = compact(taskId);
    const target = compact(targetSource).toLowerCase();
    const countMetrics = (metrics) => {
      let total = 0;
      const bySize = {};
      for (const [size, list] of Object.entries(metrics || {})) {
        const count = Array.isArray(list) ? list.length : 0;
        bySize[size] = count;
        total += count;
      }
      return { total, bySize };
    };
    for (const row of rows) {
      const dataList = row?.columns?.test_data?.dataList || row?.test_data?.dataList || row?.testData?.dataList || [];
      for (const item of Array.isArray(dataList) ? dataList : []) {
        const source = compact(item?.imageTestSource || item?.source || item?.testChannel).toLowerCase();
        const currentTaskId = compact(item?.experimentTaskId || item?.taskId || item?.id);
        if (source !== target) continue;
        if (targetTaskId && currentTaskId && targetTaskId !== currentTaskId) continue;
        const origin = countMetrics(item?.originImageMetrics || {});
        const test = countMetrics(item?.testImageMetrics || {});
        return {
          taskId: currentTaskId,
          source,
          status: compact(item?.testStatus ?? item?.status),
          originImageCount: origin.total,
          originImageCountBySize: origin.bySize,
          testImageCount: test.total,
          testImageCountBySize: test.bySize,
          online: String(item?.testStatus ?? item?.status) === '1',
        };
      }
    }
    return null;
  }
  async function searchTasks(itemId, options = {}) {
    const params = {
      tabCode: options.tabCode || 'all',
      itemIdOrName: String(itemId || ''),
    };
    if (options.testChannel) params.testChannel = options.testChannel;
    if (options.testStatus !== undefined && options.testStatus !== null && options.testStatus !== '') {
      params.testStatus = String(options.testStatus);
    }
    const runSearch = async (extraParams = {}) => {
      const requestParams = { ...params, ...extraParams };
      const payload = await callMtop('mtop.taobao.qn.copilot.framework.listmodel.data.search', {
        modelCode: 'image_test_mgr',
        params: JSON.stringify(requestParams),
        currentPage: 1,
        pageSize: 10,
      });
      const rows = extractRows(payload);
      const total = Number(payload?.result?.total || payload?.total || payload?.count || rows.length || 0);
      return {
        total: Number.isFinite(total) ? total : rows.length,
        rows,
        requestParams,
        raw: payload,
      };
    };
    if (options.testStatus !== undefined && options.testStatus !== null && options.testStatus !== '') {
      return runSearch();
    }
    const primary = await runSearch();
    if (primary.rows.length || options.skipStatusFallback) return primary;
    const mergedRows = [];
    const seen = new Set();
    const readbacks = [primary];
    for (const status of ['0', '1']) {
      const readback = await runSearch({ testStatus: status });
      readbacks.push(readback);
      for (const row of readback.rows) {
        const rowKey = `${compact(row?.domainId || row?.itemId || row?.id)}:${JSON.stringify((row?.columns?.test_data?.dataList || row?.test_data?.dataList || row?.testData?.dataList || []).map((item) => compact(item?.experimentTaskId || item?.taskId || item?.id)))}`;
        if (seen.has(rowKey)) continue;
        seen.add(rowKey);
        mergedRows.push(row);
      }
    }
    return {
      total: mergedRows.length,
      rows: mergedRows,
      requestParams: params,
      raw: primary.raw,
      statusReadbacks: readbacks.map((item) => ({ total: item.total, requestParams: item.requestParams })),
    };
  }
  async function findReusableStoppedTask(itemId, targetSource = 'common_search') {
    const readback = await searchTasks(itemId, {
      testChannel: targetSource,
      testStatus: '0',
      skipStatusFallback: true,
    });
    const tasks = taskDataItems(readback, targetSource)
      .filter((task) => compact(task.itemId) === compact(itemId) || !task.itemId);
    const reusable = tasks.find((task) => task.status === '-1') || tasks[0] || null;
    if (!reusable) return null;
    return {
      ...reusable,
      readbackTotal: readback.total,
    };
  }
  async function deleteTaskMaterial(experimentTaskId, itemId, materialId, source = 'common_search') {
    return callMtop('mtop.taobao.qn.copilot.test.image.batch.delete', {
      experimentTaskId,
      itemId: String(itemId || ''),
      materialIds: String(materialId || ''),
      source,
    });
  }
  async function clearExistingTestMaterials(task, options = {}) {
    const deleted = [];
    const errors = [];
    const ids = Array.isArray(task?.testMaterialIds) ? task.testMaterialIds : [];
    const delayMs = Number(options.delayMs || 0);
    for (const materialId of ids) {
      try {
        await deleteTaskMaterial(task.experimentTaskId, options.itemId || task.itemId, materialId, task.source || 'common_search');
        deleted.push(materialId);
      } catch (error) {
        errors.push({ materialId, error: describeError(error) });
        break;
      }
      if (delayMs > 0 && deleted.length < ids.length) await sleep(delayMs);
    }
    return { deleted, errors };
  }
  const result = {
    uploaded: [],
    materials: [],
    createPayload: {
      source: 'qn',
      itemId: String(__payload.item_id || ''),
      imageTestSources: JSON.stringify(['COMMON_SEARCH']),
    },
    batchPayload: null,
    batchPayloads: [],
    batchAddErrors: [],
    batchAddWarnings: [],
    reusedExistingTask: null,
    clearedMaterialIds: [],
    clearErrors: [],
    liveOnlineRequested: Boolean(__payload.live_online),
    onlinePayload: null,
    createResult: null,
    batchAddResult: null,
    batchAddResults: [],
    onlineResult: null,
    readback: null,
    error: '',
  };
  try {
    const uploadDelayMs = Number(__payload.upload_delay_ms || 0);
    const createDelayMs = Number(__payload.create_delay_ms || 0);
    const batchDelayMs = Number(__payload.batch_delay_ms || 0);
    const readbackDelayMs = Number(__payload.readback_delay_ms || 0);
    for (const file of (__payload.files || [])) {
      if (!__payload.live_upload) {
        result.uploaded.push({ name: file.name, role: file.role, localPath: file.localPath, url: '', skipped: true });
        continue;
      }
      const item = await uploadDataUrl(file.dataUrl, file.name);
      result.uploaded.push({ ...item, role: file.role, localPath: file.localPath, skipped: false });
      if (uploadDelayMs > 0) await sleep(uploadDelayMs);
    }
    result.materials = result.uploaded
      .filter((item) => item.url)
      .map((item) => ({ sourceType: 4, picUrl: item.url, size: inferMaterialSize(item) }));
    if (!__payload.live_create) {
      result.batchPayload = {
        experimentTaskId: '<experimentTaskId>',
        itemId: String(__payload.item_id || ''),
        source: 'common_search',
        materials: JSON.stringify(result.materials),
      };
      return { success: true, data: [result], meta: { has_more: false } };
    }
    if (!result.materials.length) throw new Error('没有可添加的天猫素材 URL，请先启用并完成 live_upload');
    let taskStatusList = [];
    let task = null;
    const reusableTask = __payload.allow_reuse_stopped_task !== false
      ? await findReusableStoppedTask(__payload.item_id, 'common_search')
      : null;
    if (reusableTask) {
      result.reusedExistingTask = {
        experimentTaskId: reusableTask.experimentTaskId,
        source: reusableTask.source,
        status: reusableTask.status,
        testMaterialIds: reusableTask.testMaterialIds,
      };
      const clearResult = await clearExistingTestMaterials(reusableTask, {
        itemId: String(__payload.item_id || ''),
        delayMs: batchDelayMs,
      });
      result.clearedMaterialIds = clearResult.deleted;
      result.clearErrors = clearResult.errors;
      if (result.clearErrors.length) {
        throw new Error(`清空旧测图素材失败：${result.clearErrors[0].materialId} ${result.clearErrors[0].error}`);
      }
      task = { experimentTaskId: reusableTask.experimentTaskId, source: reusableTask.source || 'common_search' };
      taskStatusList = [task];
    } else {
      result.createResult = await callMtop('mtop.taobao.qn.copilot.test.image.task.create', result.createPayload);
      taskStatusList = extractTaskStatusList(result.createResult, 'common_search');
      if (!taskStatusList.length) throw new Error('创建测图任务后未解析到 COMMON_SEARCH 任务 ID');
      task = taskStatusList[0];
    }
    if (createDelayMs > 0) await sleep(createDelayMs);
    for (const material of result.materials) {
      const batchPayload = {
        experimentTaskId: task.experimentTaskId,
        itemId: String(__payload.item_id || ''),
        source: 'common_search',
        materials: JSON.stringify([material]),
      };
      if (!result.batchPayload) result.batchPayload = batchPayload;
      result.batchPayloads.push(batchPayload);
      try {
        const batchResult = await callMtop('mtop.taobao.qn.copilot.test.image.batch.add', batchPayload);
        result.batchAddResults.push(batchResult);
      } catch (batchError) {
        result.batchAddErrors.push({
          payload: batchPayload,
          error: describeError(batchError),
        });
        break;
      }
      if (batchDelayMs > 0) await sleep(batchDelayMs);
    }
    result.batchAddResult = result.batchAddResults[0] || null;
    result.readback = await searchTasks(__payload.item_id);
    result.readbackTaskSummary = summarizeTask(result.readback, task.experimentTaskId, 'common_search');
    const expectedTestImages = result.materials.length;
    const actualTestImages = Number(result.readbackTaskSummary?.testImageCount || 0);
    if (result.batchAddErrors.length && actualTestImages < expectedTestImages) {
      throw new Error(`天猫 batch.add 返回失败且回读素材不足：${actualTestImages}/${expectedTestImages}；${result.batchAddErrors[0].error}`);
    }
    if (result.batchAddErrors.length) {
      result.batchAddWarnings.push(`batch.add 返回失败，但回读已写入 ${actualTestImages}/${expectedTestImages} 张素材，继续上线`);
    }
    if (__payload.live_online) {
      result.onlinePayload = {
        source: 'qn',
        itemId: String(__payload.item_id || ''),
        taskStatusList: JSON.stringify(taskStatusList),
      };
      try {
        result.onlineResult = await callMtop('mtop.taobao.qn.copilot.test.image.task.online', result.onlinePayload);
      } catch (onlineError) {
        result.onlineError = describeError(onlineError);
      }
      if (readbackDelayMs > 0) await sleep(readbackDelayMs);
      result.readback = await searchTasks(__payload.item_id);
      result.readbackTaskSummary = summarizeTask(result.readback, task.experimentTaskId, 'common_search');
      if (!result.onlineResult && result.readbackTaskSummary?.online) {
        result.onlineResult = { recoveredByReadback: true, warning: result.onlineError || '' };
      }
      if (!result.onlineResult) {
        throw new Error(result.onlineError || '天猫上线失败，且回读未确认任务已上线');
      }
    }
    return { success: true, data: [result], meta: { has_more: false } };
  } catch (error) {
    result.error = describeError(error);
    try {
      result.readback = await searchTasks(__payload.item_id);
    } catch (readbackError) {
      result.readbackError = describeError(readbackError);
    }
    return { success: true, data: [result], meta: { has_more: false } };
  }
"""


TMALL_SUBMISSION_READBACK_JS = r"""
  const compact = (value) => String(value || '').replace(/\s+/g, ' ').trim();
  const normalizeUrl = (value) => {
    const url = compact(value);
    if (!url) return '';
    return url.startsWith('//') ? `https:${url}` : url;
  };
  function describeError(error) {
    if (!error) return '未知错误';
    if (typeof error === 'string') return error;
    if (error.message) return String(error.message);
    if (Array.isArray(error.ret)) return error.ret.join('；');
    try { return JSON.stringify(error).slice(0, 1000); } catch (jsonError) {}
    return String(error);
  }
  function describeMtopError(error) {
    const parts = [];
    const ret = Array.isArray(error?.ret) ? error.ret.join('|') : '';
    const dataMsg = compact(error?.data?.errorMsg || error?.data?.message || error?.data?.msg);
    const message = compact(error?.message || error);
    if (ret) parts.push(ret);
    if (dataMsg && !parts.includes(dataMsg)) parts.push(dataMsg);
    if (message && message !== '[object Object]' && !parts.includes(message)) parts.push(message);
    return parts.join('|') || String(error || 'MTop 调用失败');
  }
  async function callMtop(api, data, options = {}) {
    const client = window.lib?.mtop || window.mtop;
    if (!client || typeof client.request !== 'function') {
      throw new Error('未找到千牛 MTop 客户端，请确认当前 tab 是天猫素材测试页');
    }
    let payload = null;
    try {
      payload = await client.request({
        api,
        v: options.v || '1.0',
        type: options.type || 'POST',
        dataType: 'json',
        H5Request: true,
        preventFallback: true,
        data,
      });
    } catch (error) {
      throw new Error(`${api} 返回失败：${describeMtopError(error)}`);
    }
    if (Array.isArray(payload?.ret)) {
      const failed = payload.ret.find((item) => !/^SUCCESS/i.test(String(item || '')));
      if (failed) throw new Error(`${api} 返回失败：${payload.ret.join('；')}`);
    }
    return payload?.data !== undefined ? payload.data : payload;
  }
  function extractRows(value) {
    for (const rows of [
      value?.result?.list,
      value?.list,
      value?.records,
      value?.data?.result?.list,
      value?.data?.list,
      value?.data?.records,
      value?.modelDataList,
    ]) {
      if (Array.isArray(rows)) return rows;
    }
    return Array.isArray(value) ? value : [];
  }
  function metricItems(value) {
    const result = [];
    for (const list of Object.values(value || {})) {
      if (Array.isArray(list)) result.push(...list);
    }
    return result;
  }
  function materialRefsFromMetrics(value) {
    const refs = [];
    const seenIds = new Set();
    const seenUrls = new Set();
    const push = (id, url) => {
      const materialId = compact(id);
      const imageUrl = normalizeUrl(url);
      if (materialId && seenIds.has(materialId)) return;
      if (imageUrl && seenUrls.has(imageUrl)) return;
      if (!materialId && !imageUrl) return;
      if (materialId) seenIds.add(materialId);
      if (imageUrl) seenUrls.add(imageUrl);
      refs.push({ id: materialId, url: imageUrl || (materialId ? `material:${materialId}` : '') });
    };
    for (const item of metricItems(value)) {
      push(item?.imageId, item?.imageUrl || item?.picUrl || item?.url || item?.materialUrl);
      for (const id of Array.isArray(item?.imageIds) ? item.imageIds : []) push(id, '');
      for (const op of Array.isArray(item?.materialOpParams) ? item.materialOpParams : []) {
        for (const material of Array.isArray(op?.materials) ? op.materials : []) {
          push(material?.imageId, material?.imageUrl || material?.picUrl || material?.url || material?.materialUrl);
        }
      }
    }
    return refs;
  }
  function taskDataItems(readback, targetSource = 'common_search') {
    const rows = Array.isArray(readback?.rows) ? readback.rows : [];
    const target = compact(targetSource).toLowerCase();
    const result = [];
    for (const row of rows) {
      const itemId = compact(row?.domainId || row?.itemId || row?.id);
      const dataList = row?.columns?.test_data?.dataList || row?.test_data?.dataList || row?.testData?.dataList || [];
      for (const item of Array.isArray(dataList) ? dataList : []) {
        const source = compact(item?.imageTestSource || item?.source || item?.testChannel).toLowerCase();
        if (source !== target) continue;
        const experimentTaskId = compact(item?.experimentTaskId || item?.taskId || item?.id);
        if (!experimentTaskId) continue;
        const originRefs = materialRefsFromMetrics(item?.originImageMetrics || {});
        const testRefs = materialRefsFromMetrics(item?.testImageMetrics || {});
        result.push({
          itemId,
          experimentTaskId,
          source: target,
          status: compact(item?.testStatus ?? item?.status),
          originMaterialIds: originRefs.map((ref) => ref.id).filter(Boolean),
          testMaterialIds: testRefs.map((ref) => ref.id).filter(Boolean),
          testMaterialUrls: testRefs.map((ref) => ref.url).filter(Boolean),
          testImageCount: testRefs.length,
          raw: item,
        });
      }
    }
    return result;
  }
  function summarizeTask(task) {
    if (!task) return null;
    return {
      taskId: task.experimentTaskId,
      source: task.source || 'common_search',
      status: task.status,
      originImageCount: task.originMaterialIds?.length || 0,
      testImageCount: task.testImageCount || task.testMaterialUrls?.length || task.testMaterialIds?.length || 0,
      online: String(task.status) === '1',
    };
  }
  async function searchTasks(itemId, options = {}) {
    const params = {
      tabCode: options.tabCode || 'all',
      itemIdOrName: String(itemId || ''),
    };
    if (options.testChannel) params.testChannel = options.testChannel;
    if (options.testStatus !== undefined && options.testStatus !== null && options.testStatus !== '') {
      params.testStatus = String(options.testStatus);
    }
    const runSearch = async (extraParams = {}) => {
      const requestParams = { ...params, ...extraParams };
      const payload = await callMtop('mtop.taobao.qn.copilot.framework.listmodel.data.search', {
        modelCode: 'image_test_mgr',
        params: JSON.stringify(requestParams),
        currentPage: 1,
        pageSize: 10,
      });
      const rows = extractRows(payload);
      const total = Number(payload?.result?.total || payload?.total || payload?.count || rows.length || 0);
      return {
        total: Number.isFinite(total) ? total : rows.length,
        rows,
        requestParams,
        raw: payload,
      };
    };
    if (options.testStatus !== undefined && options.testStatus !== null && options.testStatus !== '') {
      return runSearch();
    }
    const mergedRows = [];
    const seen = new Set();
    const readbacks = [];
    for (const status of ['1', '0', '']) {
      const readback = await runSearch(status ? { testStatus: status } : {});
      readbacks.push(readback);
      for (const row of readback.rows) {
        const rowKey = `${compact(row?.domainId || row?.itemId || row?.id)}:${JSON.stringify((row?.columns?.test_data?.dataList || row?.test_data?.dataList || row?.testData?.dataList || []).map((item) => compact(item?.experimentTaskId || item?.taskId || item?.id)))}`;
        if (seen.has(rowKey)) continue;
        seen.add(rowKey);
        mergedRows.push(row);
      }
    }
    return {
      total: mergedRows.length,
      rows: mergedRows,
      requestParams: params,
      raw: readbacks[0]?.raw || {},
      statusReadbacks: readbacks.map((item) => ({ total: item.total, requestParams: item.requestParams })),
    };
  }
  const result = {
    uploaded: [],
    materials: [],
    batchPayload: null,
    onlineResult: null,
    readback: null,
    readbackTaskSummary: null,
    batchAddWarnings: [],
    recoveredByReadback: false,
    error: '',
  };
  try {
    result.readback = await searchTasks(__payload.item_id, { testChannel: 'common_search' });
    const expected = Number(__payload.expected_test_images || 0);
    const tasks = taskDataItems(result.readback, 'common_search')
      .filter((task) => compact(task.itemId) === compact(__payload.item_id) || !task.itemId)
      .sort((left, right) => {
        const leftOnline = String(left.status) === '1' ? 1 : 0;
        const rightOnline = String(right.status) === '1' ? 1 : 0;
        if (leftOnline !== rightOnline) return rightOnline - leftOnline;
        return Number(right.testImageCount || 0) - Number(left.testImageCount || 0);
      });
    const task = tasks.find((item) =>
      String(item.status) === '1' && (!expected || Number(item.testImageCount || 0) >= Math.min(expected, 1))
    ) || tasks[0] || null;
    if (!task) throw new Error(`未回读到商品 ${__payload.item_id} 的搜索测图任务`);
    result.batchPayload = {
      experimentTaskId: task.experimentTaskId,
      itemId: String(__payload.item_id || ''),
      source: 'common_search',
    };
    result.uploaded = (task.testMaterialUrls || []).map((url, index) => ({ name: `readback-${index + 1}`, role: 'ai', url }));
    result.materials = (task.testMaterialUrls || []).map((url) => ({ sourceType: 4, picUrl: url, size: '3:4' }));
    result.readbackTaskSummary = summarizeTask(task);
    if (String(task.status) === '1') {
      result.onlineResult = { recoveredByReadback: true };
      result.recoveredByReadback = true;
      result.batchAddWarnings.push(compact(__payload.reason) || '通过天猫后台回读确认测图任务已上线');
    } else {
      result.error = `回读任务未上线：status=${task.status || ''}`;
    }
    return { success: true, data: [result], meta: { has_more: false } };
  } catch (error) {
    result.error = describeError(error);
    return { success: true, data: [result], meta: { has_more: false } };
  }
"""


SEMIR_FETCH_FILE_JS = r"""
  const compact = (value) => String(value || '').replace(/\s+/g, ' ').trim();
  const url = compact(__payload.url);
  if (!url) return { success: false, error: '缺少下载 URL' };
  try {
    const response = await fetch(url, { credentials: 'include' });
    if (!response.ok) {
      const text = await response.text().catch(() => '');
      throw new Error(`HTTP ${response.status}: ${text.slice(0, 240) || response.statusText}`);
    }
    const blob = await response.blob();
    const buffer = await blob.arrayBuffer();
    const bytes = new Uint8Array(buffer);
    let binary = '';
    const chunkSize = 0x8000;
    for (let index = 0; index < bytes.length; index += chunkSize) {
      binary += String.fromCharCode.apply(null, bytes.subarray(index, index + chunkSize));
    }
    return {
      success: true,
      data: [{
        url,
        bytes: bytes.length,
        mime: blob.type || '',
        base64: btoa(binary),
      }],
      meta: { has_more: false },
    };
  } catch (error) {
    return { success: false, error: String(error?.message || error) };
  }
"""


async def download_semir_candidate(
    runner: JSRunner,
    workflow: WorkflowItem,
    candidate: Mapping[str, Any],
    artifact_dir: Path,
    *,
    role: str,
) -> str:
    material_url = compact((candidate or {}).get("materialUrl"))
    if not material_url:
        return ""
    suffix = ".jpg"
    match = IMAGE_EXT_RE.search(material_url) or IMAGE_EXT_RE.search(compact((candidate or {}).get("filename")))
    if match:
        suffix = f".{match.group(1).lower().replace('jpeg', 'jpg')}"
    filename = f"{workflow.style_code}-{role}{suffix}"
    fallback = artifact_dir / filename
    page_result = await runner.evaluate_with_reconnect(js_call(SEMIR_FETCH_FILE_JS, {"url": material_url}), allow_navigation_retry=True)
    if page_result.success and isinstance(page_result.data, list) and page_result.data:
        payload = page_result.data[0] or {}
        encoded = compact(payload.get("base64"))
        if encoded:
            fallback.parent.mkdir(parents=True, exist_ok=True)
            fallback.write_bytes(base64.b64decode(encoded))
            return optimize_reference_image(str(fallback))
    fallback_result = download_url_to_file(material_url, fallback, timeout=45)
    if fallback_result.get("success"):
        return optimize_reference_image(str(fallback_result["path"]))
    return ""


async def find_semir_references(
    runner: JSRunner,
    workflow: WorkflowItem,
    cloud_path: str,
    artifact_dir: Path,
    *,
    allow_global_fallback: bool = False,
    download_files: bool = True,
    download_detail: bool = True,
) -> tuple[dict[str, Any], str, str]:
    payload = {
        "style_code": workflow.style_code,
        "skc_code": workflow.skc_code,
        "cloud_path": cloud_path,
        "limit": 8,
    }
    result = await runner.evaluate_with_reconnect(js_call(SEMIR_FIND_JS, payload), allow_navigation_retry=True)
    if not result.success:
        raise RuntimeError(result.error or "森马云盘找图失败")
    data = (result.data[0] if isinstance(result.data, list) and result.data else {}) if result.data is not None else {}
    main_missing = not (data.get("chosenMain") or data.get("chosen"))
    detail_missing = download_detail and not data.get("chosenDetail")
    if allow_global_fallback and cloud_path and (main_missing or detail_missing):
        fallback_data, _fallback_main_path, _fallback_detail_path = await find_semir_references(
            runner,
            workflow,
            "",
            artifact_dir,
            allow_global_fallback=False,
            download_files=False,
        )
        if not (data.get("chosenMain") or data.get("chosen")) and (fallback_data.get("chosenMain") or fallback_data.get("chosen")):
            data["chosenMain"] = fallback_data.get("chosenMain") or fallback_data.get("chosen")
            data["chosen"] = data["chosenMain"]
        if detail_missing and fallback_data.get("chosenDetail"):
            data["chosenDetail"] = fallback_data.get("chosenDetail")
        data["fallbackFromCloudPath"] = cloud_path
        data["fallbackSearches"] = fallback_data.get("searches") or []
    if not download_files:
        return data, "", ""
    main = data.get("chosenMain") or data.get("chosen") or {}
    detail = data.get("chosenDetail") or {}
    main_path = await download_semir_candidate(runner, workflow, main, artifact_dir, role="main-origin") if main else ""
    detail_path = ""
    detail_paths: list[str] = []
    if download_detail:
        raw_detail_candidates = data.get("detailCandidates") if isinstance(data.get("detailCandidates"), list) else []
        detail_candidates = [candidate for candidate in raw_detail_candidates if isinstance(candidate, Mapping)]
        if detail and not any(compact(candidate.get("fullpath") or candidate.get("filename")) == compact(detail.get("fullpath") or detail.get("filename")) for candidate in detail_candidates):
            detail_candidates.insert(0, detail)
        hydrated_candidates = []
        for index, candidate in enumerate(detail_candidates, start=1):
            path = await download_semir_candidate(runner, workflow, candidate, artifact_dir, role=f"detail-flat-{index}")
            next_candidate = dict(candidate)
            if path:
                next_candidate["localPath"] = path
                detail_paths.append(path)
                if not detail_path:
                    detail_path = path
            hydrated_candidates.append(next_candidate)
        if hydrated_candidates:
            data["detailCandidates"] = hydrated_candidates
        if detail_paths:
            data["detailReferencePaths"] = detail_paths
    if not main_path:
        data["downloadError"] = "未下载到橱窗1/yz(1)主图原图"
    if download_detail and not detail_path:
        data["detailDownloadError"] = "未下载到 SKC 编码平铺参考图"
    return data, main_path, detail_path


async def find_semir_origin(
    runner: JSRunner,
    workflow: WorkflowItem,
    cloud_path: str,
    artifact_dir: Path,
    *,
    allow_global_fallback: bool = False,
) -> tuple[dict[str, Any], str]:
    data, main_path, _detail_path = await find_semir_references(
        runner,
        workflow,
        cloud_path,
        artifact_dir,
        allow_global_fallback=allow_global_fallback,
    )
    return data, main_path


def download_url_to_file(url: str, target: Path, timeout: int = 60) -> dict[str, Any]:
    target.parent.mkdir(parents=True, exist_ok=True)
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    partial = target.with_name(f"{target.name}.part")
    try:
        request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with opener.open(request, timeout=timeout) as response:
            with partial.open("wb") as handle:
                while True:
                    chunk = response.read(1024 * 1024)
                    if not chunk:
                        break
                    handle.write(chunk)
        partial.replace(target)
        return {"success": True, "path": str(target), "bytes": target.stat().st_size}
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        try:
            partial.unlink(missing_ok=True)
        except Exception:
            pass
        return {"success": False, "path": str(target), "error": str(exc)}


def optimize_reference_image(path: str, *, max_bytes: int = REFERENCE_IMAGE_MAX_BYTES) -> str:
    source = Path(path)
    if not source.is_file() or source.stat().st_size <= max_bytes:
        return str(source)
    try:
        from PIL import Image
    except Exception:
        return str(source)

    target = source.with_name(f"{source.stem}-optimized.jpg")
    try:
        with Image.open(source) as image:
            image = image.convert("RGB")
            width, height = image.size
            quality = 92
            scale = 1.0
            while True:
                candidate = image
                if scale < 0.999:
                    candidate = image.resize(
                        (max(1, int(width * scale)), max(1, int(height * scale))),
                        Image.Resampling.LANCZOS,
                    )
                candidate.save(target, format="JPEG", quality=quality, optimize=True)
                if target.stat().st_size <= max_bytes:
                    return str(target)
                if quality > 72:
                    quality -= 8
                    continue
                if scale > 0.55:
                    scale *= 0.88
                    quality = 88
                    continue
                return str(target)
    except Exception:
        try:
            target.unlink(missing_ok=True)
        except Exception:
            pass
        return str(source)


def normalize_tmall_upload_image(path: str, *, size: tuple[int, int] = TMALL_UPLOAD_IMAGE_SIZE) -> str:
    source = Path(path)
    if not source.is_file():
        return str(path)
    try:
        from PIL import Image, ImageOps
    except Exception:
        return str(source)

    target = source.with_name(f"{source.stem}-tmall-3x4.jpg")
    if source.resolve() == target.resolve():
        return str(source)
    target_ratio = size[0] / size[1]
    try:
        with Image.open(source) as image:
            image = ImageOps.exif_transpose(image).convert("RGB")
            width, height = image.size
            if width <= 0 or height <= 0:
                return str(source)
            ratio = width / height
            if ratio > target_ratio:
                crop_width = max(1, int(height * target_ratio))
                left = max(0, (width - crop_width) // 2)
                image = image.crop((left, 0, left + crop_width, height))
            elif ratio < target_ratio:
                crop_height = max(1, int(width / target_ratio))
                top = max(0, (height - crop_height) // 2)
                image = image.crop((0, top, width, top + crop_height))
            image = image.resize(size, Image.Resampling.LANCZOS)
            image.save(target, format="JPEG", quality=92, optimize=True)
            return str(target)
    except Exception:
        try:
            target.unlink(missing_ok=True)
        except Exception:
            pass
        return str(source)


def guess_image_suffix(url: str, fallback: str = ".jpg") -> str:
    match = IMAGE_EXT_RE.search(url)
    if not match:
        return fallback
    ext = match.group(1).lower()
    return f".{'jpg' if ext == 'jpeg' else ext}"


def download_generated_images(row: Mapping[str, Any], artifact_dir: Path) -> list[str]:
    urls = parse_list(row.get("生成图URL"))
    paths: list[str] = []
    prompt_index = int(row.get("提示词序号") or 1)
    output_token = safe_token(row.get("__1xm_output_token"), "")
    for index, url in enumerate(urls, start=1):
        suffix = guess_image_suffix(url, ".jpg")
        image_index = prompt_index if len(urls) == 1 else f"{prompt_index}-{index}"
        token_suffix = f"-{output_token}" if output_token else ""
        target = artifact_dir / f"{compact(row.get('款号'))}-ai-{image_index}{token_suffix}{suffix}"
        result = download_url_to_file(url, target, timeout=90)
        if result.get("success"):
            paths.append(str(result["path"]))
    return paths


def image_file_payload(path: str, *, role: str, name: str) -> dict[str, str]:
    data_url = file_to_data_url(path)
    return {
        "name": name,
        "role": role,
        "localPath": str(path),
        "dataUrl": data_url,
    }


async def upload_and_create_tmall_task(
    runner: JSRunner,
    workflow: WorkflowItem,
    origin_path: str,
    generated_paths: list[str],
    *,
    live_upload: bool,
    live_create: bool,
    live_online: bool,
    upload_delay_seconds: float = 1.5,
    create_delay_seconds: float = 5.0,
    batch_delay_seconds: float = 12.0,
    readback_delay_seconds: float = 5.0,
) -> dict[str, Any]:
    files = []
    if origin_path:
        files.append(image_file_payload(normalize_tmall_upload_image(origin_path), role="origin", name=f"{workflow.style_code}-origin.jpg"))
    for index, path in enumerate(generated_paths, start=1):
        files.append(image_file_payload(normalize_tmall_upload_image(path), role="ai", name=f"{workflow.style_code}-ai-{index}.jpg"))
    payload = {
        "style_code": workflow.style_code,
        "item_id": workflow.item_id,
        "files": files,
        "live_upload": live_upload,
        "live_create": live_create,
        "live_online": live_online,
        "folder_id": "0",
        "allow_reuse_stopped_task": True,
        "upload_delay_ms": max(0, int(float(upload_delay_seconds or 0) * 1000)),
        "create_delay_ms": max(0, int(float(create_delay_seconds or 0) * 1000)),
        "batch_delay_ms": max(0, int(float(batch_delay_seconds or 0) * 1000)),
        "readback_delay_ms": max(0, int(float(readback_delay_seconds or 0) * 1000)),
    }
    result = await runner.evaluate_with_reconnect(js_call(TMALL_UPLOAD_CREATE_JS, payload), allow_navigation_retry=True)
    if not result.success:
        raise RuntimeError(result.error or "天猫上传/创建测图任务失败")
    return (result.data[0] if isinstance(result.data, list) and result.data else {}) if result.data is not None else {}


async def recover_tmall_submission_by_readback(
    runner: JSRunner,
    workflow: WorkflowItem,
    *,
    expected_test_images: int = 0,
    reason: str = "",
    attempts: int = 4,
    delay_seconds: float = 2.0,
) -> dict[str, Any] | None:
    payload = {
        "item_id": workflow.item_id,
        "style_code": workflow.style_code,
        "expected_test_images": max(0, int(expected_test_images or 0)),
        "reason": compact(reason),
    }
    latest: dict[str, Any] | None = None
    total_attempts = max(1, min(10, int(attempts or 1)))
    for attempt in range(total_attempts):
        result = await runner.evaluate_with_reconnect(
            js_call(TMALL_SUBMISSION_READBACK_JS, payload),
            allow_navigation_retry=True,
        )
        if result.success:
            data = (result.data[0] if isinstance(result.data, list) and result.data else {}) if result.data is not None else {}
            latest = dict(data) if isinstance(data, Mapping) else None
            if tmall_submission_success(latest)[0]:
                return latest
        if attempt + 1 < total_attempts and float(delay_seconds or 0) > 0:
            await asyncio.sleep(float(delay_seconds))
    return latest


def should_recover_tmall_submit_error(error: object, previous_row: Mapping[str, Any] | None = None) -> bool:
    if previous_row is not None:
        return True
    text = compact(error).lower()
    return any(marker in text for marker in ("timeout", "超时", "batch.add", "上线", "readback", "回读"))


def tmall_submission_success(tmall_result: Mapping[str, Any] | None) -> tuple[bool, str, str]:
    if not isinstance(tmall_result, Mapping):
        return False, "", "天猫提交未返回结果"
    task_id = ""
    if isinstance(tmall_result.get("batchPayload"), Mapping):
        task_id = compact(tmall_result["batchPayload"].get("experimentTaskId"))
    tmall_error = compact(tmall_result.get("error") or tmall_result.get("uploadError") or tmall_result.get("备注"))
    if not task_id or task_id == "<experimentTaskId>":
        return False, task_id, tmall_error or "未解析到测图任务 ID"
    if tmall_error:
        return False, task_id, tmall_error
    task_summary = tmall_result.get("readbackTaskSummary") if isinstance(tmall_result.get("readbackTaskSummary"), Mapping) else {}
    online_ok = bool(tmall_result.get("onlineResult")) or bool(task_summary.get("online")) or compact(task_summary.get("status")) == "1"
    if not online_ok:
        return False, task_id, "任务未上线，未开始测试"
    return True, task_id, ""


def write_submit_result_artifacts(
    batch: dict[str, Any],
    result_rows: list[dict[str, Any]],
    *,
    status: str,
    attempted: int,
    succeeded: int,
    failed: int,
    artifact_dir: Path,
) -> tuple[str, str]:
    excel_path = data_sink.export_excel(
        result_rows,
        ADAPTER_ID,
        TASK_ID,
        "天猫AI测图全链路执行证据_{timestamp}.xlsx",
        column_order=TMALL_AI_IMAGE_RESULT_COLUMN_ORDER,
    )
    audit_path = artifact_dir / f"tmall-ai-image-approval-submit-{batch.get('batch_id', 'batch')}.json"
    audit_path.write_text(json.dumps({
        "status": status,
        "attempted": attempted,
        "succeeded": succeeded,
        "failed": failed,
        "submitted": succeeded,
        "rows": result_rows,
        "excel_path": excel_path,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    batch["submit_result_path"] = str(audit_path)
    batch["submit_result_excel_path"] = str(excel_path)
    return str(audit_path), str(excel_path)


def submit_mode_label(mode: object) -> str:
    text = compact(mode)
    if text == "append_unsubmitted":
        return "追加未提交图片"
    if text == "rerun_confirmed":
        return "重新提交已确认图片"
    return "提交已确认图片"


def submit_progress_message_for_plan(plan: Mapping[str, Any], workflow: WorkflowItem) -> str:
    image_count = len(plan.get("generated_paths") or [])
    label = submit_mode_label(plan.get("submit_mode"))
    if image_count > 0:
        return f"{label}：{workflow.style_code}（AI图 {image_count} 张）"
    return f"{label}：{workflow.style_code}"


async def upload_approved_tmall_batch(batch: dict[str, Any], log=None) -> dict[str, Any]:
    log = log or (lambda _message: None)
    batch["test_task_status"] = "creating"
    if compact(batch.get("execution_mode")) == "approval_then_create":
        batch["approval_status"] = "completed"
    previous_latest_rows = latest_submit_rows_by_style(batch)
    previous_success_rows = successful_submit_rows_by_style(batch)
    plans = selected_upload_plan_from_approval_batch(batch)
    if not plans and not previous_success_rows:
        raise RuntimeError("审批批次中没有已确认的 AI 图")

    run_params = batch.get("run_params") if isinstance(batch.get("run_params"), Mapping) else {}
    artifact_dir = Path(compact(batch.get("artifact_dir")) or ".")
    result_rows: list[dict[str, Any]] = [dict(row) for row in previous_success_rows.values()]
    attempted = len(previous_success_rows)
    succeeded = len(previous_success_rows)
    failed = 0
    started_at = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    submit_total = len(plans) + len(previous_success_rows)

    def persist_submit_progress(status: str = "running", *, current_style: str = "", message: str = "") -> None:
        batch["submit_progress"] = {
            "status": status,
            "total": submit_total,
            "completed": attempted,
            "attempted": attempted,
            "succeeded": succeeded,
            "failed": failed,
            "current_style": current_style,
            "message": message,
            "started_at": started_at,
            "updated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        }
        batch["submit_summary"] = {
            "attempted": attempted,
            "succeeded": succeeded,
            "failed": failed,
        }
        batch["submit_result_rows"] = json_safe(result_rows)
        save_approval_batch(batch)

    if not plans:
        status = "created"
        batch["status"] = status
        batch["test_task_status"] = "succeeded"
        batch["submitted_at"] = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
        audit_path, excel_path = write_submit_result_artifacts(
            batch,
            result_rows,
            status=status,
            attempted=attempted,
            succeeded=succeeded,
            failed=failed,
            artifact_dir=artifact_dir,
        )
        persist_submit_progress("completed", message="全部测图任务已通过历史回读确认")
        save_approval_batch(batch)
        notify_tmall_submission_result(
            batch,
            status=status,
            attempted=attempted,
            succeeded=succeeded,
            failed=failed,
            rows=result_rows,
            channel=compact(run_params.get("approval_notify_channel") or "dingtalk").lower(),
            log=log,
        )
        return {
            "ok": True,
            "status": batch["status"],
            "attempted": attempted,
            "succeeded": succeeded,
            "failed": failed,
            "submitted": succeeded,
            "rows": result_rows,
            "result_path": excel_path,
            "excel_path": excel_path,
            "audit_path": audit_path,
        }

    cdp_url = compact(run_params.get("cdp_url")) or "http://127.0.0.1:9222"
    bridge = CDPBridge(cdp_url)
    if not bridge.is_available(timeout=5):
        raise RuntimeError(f"无法连接 Chrome CDP：{cdp_url}")
    runner = await get_runner(
        bridge,
        "https://myseller.taobao.com/home.htm/material-center/material-test",
        TMALL_URL,
        artifact_dir,
    )
    try:
        runner.timeout = max(int(getattr(runner, "timeout", 60) or 60), 120)
    except Exception:
        pass
    await wait_for_tmall_api_ready(runner, log=log, timeout_ms=TMALL_LOGIN_READY_TIMEOUT_MS)
    batch["status"] = "submitting"
    persist_submit_progress("running", message="准备提交已确认图片")
    for plan in plans:
        workflow = workflow_from_dict(plan.get("workflow") or {})
        existing_row = previous_latest_rows.get(workflow.style_code)
        plan_message = submit_progress_message_for_plan(plan, workflow)
        if not workflow.item_id:
            attempted += 1
            failed += 1
            result_rows.append({
                "表格行号": workflow.row_no,
                "款号": workflow.style_code,
                "商品ID": workflow.item_id,
                "阶段": "审批后上传/创建测图任务",
                "执行结果": "跳过",
                "备注": "缺少商品 ID",
                "提交模式": submit_mode_label(plan.get("submit_mode")),
                "提交图片文件": "\n".join(plan.get("generated_paths") or []),
                "提交图片数量": len(plan.get("generated_paths") or []),
            })
            persist_submit_progress("running", current_style=workflow.style_code, message="缺少商品 ID，已跳过")
            continue
        if existing_row and not tmall_submit_row_is_online_success(existing_row):
            persist_submit_progress("running", current_style=workflow.style_code, message=f"正在回读 {workflow.style_code} 的后台最新状态")
            recovered_result = await recover_tmall_submission_by_readback(
                runner,
                workflow,
                expected_test_images=len(plan.get("generated_paths") or []),
                reason="历史提交显示 timeout，已通过天猫后台回读纠偏，未重复上传",
            )
            recovered_ok, _recovered_task_id, _recovered_error = tmall_submission_success(recovered_result)
            if recovered_ok:
                attempted += 1
                succeeded += 1
                log(f"[approval] {workflow.style_code} 历史 timeout 已通过天猫后台回读确认上线")
                result_rows.extend(result_rows_for_workflow(
                    workflow,
                    (plan.get("item") or {}).get("semir_data") or {},
                    compact(plan.get("origin_path")),
                    compact(plan.get("detail_reference_path")),
                    [
                        asset.get("generation_row") or {}
                        for asset in (plan.get("item") or {}).get("assets", [])
                        if asset.get("kind") == "ai" and compact(asset.get("path")) in set(plan.get("generated_paths") or [])
                    ],
                    list(plan.get("generated_paths") or []),
                    recovered_result,
                    reference_mode=compact((plan.get("item") or {}).get("reference_mode")) or "main_only",
                    submit_mode=plan.get("submit_mode"),
                ))
                persist_submit_progress("running", current_style=workflow.style_code, message="后台回读确认已上线，跳过重复提交")
                continue
        persist_submit_progress("running", current_style=workflow.style_code, message=plan_message)
        log(f"[approval] {workflow.style_code} 上传 {len(plan.get('generated_paths') or [])} 张已确认 AI 图并创建测图任务")
        try:
            tmall_result = await upload_and_create_tmall_task(
                runner,
                workflow,
                compact(plan.get("origin_path")),
                list(plan.get("generated_paths") or []),
                live_upload=True,
                live_create=True,
                live_online=True,
                upload_delay_seconds=float(run_params.get("tmall_upload_delay_seconds") or 1.5),
                create_delay_seconds=float(run_params.get("tmall_create_delay_seconds") or 5.0),
                batch_delay_seconds=float(run_params.get("tmall_batch_delay_seconds") or 12.0),
                readback_delay_seconds=float(run_params.get("tmall_readback_delay_seconds") or 5.0),
            )
        except Exception as exc:
            submit_exception = str(exc)
            recovered_result = None
            if should_recover_tmall_submit_error(submit_exception, existing_row):
                recovered_result = await recover_tmall_submission_by_readback(
                    runner,
                    workflow,
                    expected_test_images=len(plan.get("generated_paths") or []),
                    reason=f"提交调用返回 {submit_exception}，已通过天猫后台回读确认",
                )
            recovered_ok, _recovered_task_id, _recovered_error = tmall_submission_success(recovered_result)
            if recovered_ok:
                tmall_result = recovered_result or {}
            else:
                tmall_result = {"uploaded": [], "materials": [], "batchPayload": {}, "onlineResult": None, "error": submit_exception}
        attempted += 1
        success, _task_id, submit_error = tmall_submission_success(tmall_result)
        if success:
            succeeded += 1
            log(f"[approval] {workflow.style_code} 测图任务已提交成功")
        else:
            failed += 1
            log(f"[approval] {workflow.style_code} 测图任务提交失败：{submit_error}")
        result_rows.extend(result_rows_for_workflow(
            workflow,
            (plan.get("item") or {}).get("semir_data") or {},
            compact(plan.get("origin_path")),
            compact(plan.get("detail_reference_path")),
            [
                asset.get("generation_row") or {}
                for asset in (plan.get("item") or {}).get("assets", [])
                if asset.get("kind") == "ai" and compact(asset.get("path")) in set(plan.get("generated_paths") or [])
            ],
            list(plan.get("generated_paths") or []),
            tmall_result,
            reference_mode=compact((plan.get("item") or {}).get("reference_mode")) or "main_only",
            submit_mode=plan.get("submit_mode"),
        ))
        persist_submit_progress("running", current_style=workflow.style_code, message="已完成本款提交" if success else submit_error)

    if attempted > 0 and succeeded == attempted and failed == 0:
        status = "created"
    elif succeeded > 0:
        status = "partial_failed"
    else:
        status = "create_failed"
    batch["status"] = status
    batch["test_task_status"] = "succeeded" if status == "created" else (
        "partial_success" if status == "partial_failed" else "failed"
    )
    batch["submitted_at"] = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    batch["submit_result_rows"] = json_safe(result_rows)
    audit_path, excel_path = write_submit_result_artifacts(
        batch,
        result_rows,
        status=status,
        attempted=attempted,
        succeeded=succeeded,
        failed=failed,
        artifact_dir=artifact_dir,
    )
    batch["submit_summary"] = {
        "attempted": attempted,
        "succeeded": succeeded,
        "failed": failed,
    }
    batch["submit_progress"] = {
        "status": "completed" if failed == 0 and attempted > 0 else status,
        "total": submit_total,
        "completed": attempted,
        "attempted": attempted,
        "succeeded": succeeded,
        "failed": failed,
        "current_style": "",
        "message": "提交完成" if failed == 0 and attempted > 0 else "提交完成，存在失败款",
        "started_at": started_at,
        "updated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
    }
    save_approval_batch(batch)
    notify_tmall_submission_result(
        batch,
        status=status,
        attempted=attempted,
        succeeded=succeeded,
        failed=failed,
        rows=result_rows,
        channel=compact(run_params.get("approval_notify_channel") or "dingtalk").lower(),
        log=log,
    )
    return {
        "ok": failed == 0 and attempted > 0,
        "status": batch["status"],
        "attempted": attempted,
        "succeeded": succeeded,
        "failed": failed,
        "submitted": succeeded,
        "rows": result_rows,
        "result_path": excel_path,
        "excel_path": excel_path,
        "audit_path": audit_path,
    }


def result_rows_for_workflow(
    workflow: WorkflowItem,
    semir_data: Mapping[str, Any],
    main_origin_path: str,
    detail_reference_path: str,
    generation_rows: list[Mapping[str, Any]],
    generated_paths: list[str],
    tmall_result: Mapping[str, Any] | None,
    error: str = "",
    reference_mode: str = "main_only",
    submit_mode: object = "",
) -> list[dict[str, Any]]:
    chosen = (semir_data or {}).get("chosenMain") or (semir_data or {}).get("chosen") or {}
    detail = (semir_data or {}).get("chosenDetail") or {}
    main_candidate_paths = "\n".join(
        compact(item.get("fullpath") or item.get("filename"))
        for item in ((semir_data or {}).get("mainCandidates") or [])
        if isinstance(item, Mapping) and compact(item.get("fullpath") or item.get("filename"))
    )
    detail_candidate_paths = "\n".join(
        compact(item.get("fullpath") or item.get("filename"))
        for item in ((semir_data or {}).get("detailCandidates") or [])
        if isinstance(item, Mapping) and compact(item.get("fullpath") or item.get("filename"))
    )
    detail_reference_files = parse_list((semir_data or {}).get("detailReferencePaths")) or parse_list(detail_reference_path)
    resolved_skc_code = workflow.skc_code or compact(detail.get("skcCode"))
    requires_detail = compact(reference_mode) == "main_and_detail"
    reference_ok = bool(main_origin_path and (detail_reference_path or not requires_detail))
    if reference_ok and detail_reference_path:
        reference_result = "已下载双参考图"
    elif reference_ok:
        reference_result = "已下载主图参考图"
    else:
        reference_result = "参考图不完整"
    rows = [{
        "表格行号": workflow.row_no,
        "款号": workflow.style_code,
        "商品ID": workflow.item_id,
        "阶段": "森马云盘找图",
        "参考图模式": reference_mode,
        "品类": workflow.category,
        "性别": workflow.gender,
        "SKC编码": resolved_skc_code,
        "云盘路径": compact(chosen.get("fullpath")),
        "文件名": compact(chosen.get("filename")),
        "主图候选Top5": main_candidate_paths,
        "款色图候选": detail_candidate_paths,
        "素材URL": compact(chosen.get("materialUrl")),
        "主图原图文件": main_origin_path,
        "平铺参考图路径": compact(detail.get("fullpath")),
        "平铺参考图文件名": compact(detail.get("filename")),
        "平铺参考图URL": compact(detail.get("materialUrl")),
        "细节参考图文件": "\n".join(detail_reference_files),
        "本地文件": "\n".join(item for item in [main_origin_path, *detail_reference_files] if item),
        "执行结果": reference_result,
        "备注": error or "；".join(item for item in [
            compact((semir_data or {}).get("downloadError")),
            compact((semir_data or {}).get("detailDownloadError")),
        ] if item),
    }]
    for index, generation_row in enumerate(generation_rows or [], start=1):
        row_paths = parse_list(generation_row.get("本地生成图文件"))
        rows.append({
            "表格行号": workflow.row_no,
            "款号": workflow.style_code,
            "SKC编码": resolved_skc_code,
            "商品ID": workflow.item_id,
            "阶段": "1XM生图",
            "参考图模式": reference_mode,
            "提示词分组": generation_row.get("提示词分组", ""),
            "提示词字段名": generation_row.get("提示词字段名", ""),
            "提示词序号": generation_row.get("提示词序号", index),
            "尺寸": generation_row.get("尺寸", ""),
            "格式": generation_row.get("格式", ""),
            "质量": generation_row.get("质量", ""),
            "参考图文件": generation_row.get("参考图文件", ""),
            "主参考图文件": generation_row.get("主参考图文件", ""),
            "细节参考图文件": generation_row.get("细节参考图文件", ""),
            "最终提示词": generation_row.get("最终提示词", ""),
            "完整Prompt": generation_row.get("完整Prompt") or generation_row.get("最终提示词", ""),
            "1XM任务ID": generation_row.get("1XM任务ID", ""),
            "1XM轮询URL": generation_row.get("1XM轮询URL", ""),
            "生成图URL": generation_row.get("生成图URL", ""),
            "生成图数量": generation_row.get("生成图数量", ""),
            "本地文件": "\n".join(row_paths),
            "执行结果": generation_row.get("执行结果", ""),
            "备注": generation_row.get("备注", ""),
        })
    if tmall_result is not None:
        uploaded = tmall_result.get("uploaded") or []
        materials = tmall_result.get("materials") or []
        readback = tmall_result.get("readback") if isinstance(tmall_result.get("readback"), dict) else {}
        readback_total = readback.get("total") or readback.get("count") or ""
        task_id = ""
        if isinstance(tmall_result.get("batchPayload"), dict):
            task_id = compact(tmall_result["batchPayload"].get("experimentTaskId"))
        online_ok = bool(tmall_result.get("onlineResult"))
        tmall_error = compact(tmall_result.get("error") or tmall_result.get("uploadError") or tmall_result.get("备注"))
        tmall_warning = "；".join(compact(item) for item in (tmall_result.get("batchAddWarnings") or []) if compact(item))
        task_summary = tmall_result.get("readbackTaskSummary") if isinstance(tmall_result.get("readbackTaskSummary"), dict) else {}
        readback_note = f"total={readback_total}" if readback_total != "" else ""
        if task_summary:
            readback_note = "；".join(item for item in [
                readback_note,
                f"status={task_summary.get('status', '')}",
                f"testImages={task_summary.get('testImageCount', '')}",
            ] if item)
        submit_ok, _submit_task_id, submit_error = tmall_submission_success(tmall_result)
        if not tmall_error and not submit_ok:
            tmall_error = submit_error
        first_generation = generation_rows[0] if generation_rows else {}
        rows.append({
            "表格行号": workflow.row_no,
            "款号": workflow.style_code,
            "SKC编码": resolved_skc_code,
            "商品ID": workflow.item_id,
            "阶段": "天猫上传/创建测图任务",
            "测图详情URL": tmall_material_test_detail_url(workflow.item_id),
            "参考图模式": reference_mode,
            "提示词分组": first_generation.get("提示词分组", ""),
            "提示词字段名": "\n".join(compact(row.get("提示词字段名")) for row in generation_rows if compact(row.get("提示词字段名"))),
            "完整Prompt": "\n\n---\n\n".join(compact(row.get("完整Prompt") or row.get("最终提示词")) for row in generation_rows if compact(row.get("完整Prompt") or row.get("最终提示词"))),
            "任务ID": task_id,
            "上传图数量": len([item for item in uploaded if item.get("url")]),
            "提交图片数量": len(generated_paths),
            "提交图片文件": "\n".join(generated_paths),
            "提交模式": submit_mode_label(submit_mode),
            "上线结果": "已上线" if online_ok else "未上线",
            "页面回读": readback_note,
            "素材URL": "\n".join(item.get("picUrl", "") for item in materials),
            "执行结果": (
                "天猫上传/创建失败" if tmall_error else
                "已创建/加素材/已上线"
                if task_id and task_id != "<experimentTaskId>" and online_ok
                else "已生成计划"
            ),
            "备注": tmall_error or tmall_warning or f"uploaded={len(uploaded)} materials={len(materials)} online={online_ok}",
        })
    return rows


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run Tmall AI image-test chain through CDP page APIs.")
    parser.add_argument("--execute-mode", default="approval_then_create", choices=CHAIN_EXECUTION_MODES)
    parser.add_argument("--workflow-file", default=DEFAULT_WORKFLOW_FILE)
    parser.add_argument("--prompt-source", default="local_excel", choices=["local_excel", "cloud_prompt_library"])
    parser.add_argument("--prompt-file", default=DEFAULT_PROMPT_FILE)
    parser.add_argument("--cloud-prompt-library-id", default="")
    parser.add_argument("--cloud-prompt-templates-json", default="")
    parser.add_argument("--limit", type=int, default=0, help=argparse.SUPPRESS)
    parser.add_argument("--cloud-path", default=DEFAULT_CLOUD_PATH)
    parser.add_argument("--cdp-url", default="http://127.0.0.1:9222")
    parser.add_argument("--tmall-upload-delay-seconds", type=float, default=1.5, help="Delay between Tmall picture uploads.")
    parser.add_argument("--tmall-create-delay-seconds", type=float, default=5.0, help="Delay after task.create before batch.add.")
    parser.add_argument("--tmall-batch-delay-seconds", type=float, default=12.0, help="Delay between Tmall batch.add calls.")
    parser.add_argument("--tmall-readback-delay-seconds", type=float, default=5.0, help="Delay after task.online before readback.")
    parser.add_argument("--ratio", default="")
    parser.add_argument("--image-size", default="1536x2048")
    parser.add_argument("--model", default="gpt-image-2")
    parser.add_argument("--ai-image-count", type=int, default=4)
    parser.add_argument("--generation-concurrency", type=int, default=100)
    parser.add_argument("--reference-mode", default="main_only", choices=["main_only", "main_and_detail"])
    parser.add_argument("--allow-global-semir-fallback", action="store_true", help="Allow searching all Semir mounts when the configured cloud path misses required images.")
    parser.add_argument("--quality", default="auto", choices=["auto", "low", "medium", "high"])
    parser.add_argument("--output-format", default="jpeg")
    parser.add_argument("--one-xm-key-tier", default="4k", choices=["2k", "4k", "auto"])
    parser.add_argument("--retry-attempts", type=int, default=3)
    parser.add_argument("--compensate-attempts", type=int, default=2)
    parser.add_argument("--poll-timeout-minutes", type=float, default=12)
    parser.add_argument("--approval-message-template", default="")
    parser.add_argument("--approval-notify-channel", default="dingtalk", choices=["dingtalk", "feishu", "webhook", "none"])
    parser.add_argument("--approval-base-url", default="")
    parser.add_argument("--task-instance-uid", default="")
    parser.add_argument("--task-run-uid", default="")
    return parser


async def run_chain_rows(args: argparse.Namespace, artifact_dir: Path, log=None, wait_for_control=None) -> list[dict[str, Any]]:
    args = apply_chain_execution_mode(args)
    if not compact(getattr(args, "task_run_uid", "")):
        args.task_run_uid = f"manual-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{secrets.token_hex(3)}"
    all_rows: list[dict[str, Any]] = []
    approval_items: list[dict[str, Any]] = []
    workflow_results: list[dict[str, Any]] = []
    generation_plans: list[dict[str, Any]] = []
    log = log or print

    workflow_table = read_workbook_table(args.workflow_file)
    workflow_rows, invalid_rows = normalize_workflow_rows(workflow_table, int(getattr(args, "limit", 0) or 0))
    prompts = load_prompt_library_for_args(args)
    if not workflow_rows:
        raise RuntimeError("工作流表未解析到款号")
    if not prompts:
        raise RuntimeError("提示词库未解析到可用提示词")

    bridge = CDPBridge(args.cdp_url)
    if not bridge.is_available(timeout=5):
        raise RuntimeError(f"无法连接 Chrome CDP：{args.cdp_url}")
    semir_runner = await get_runner(bridge, "https://fmp.semirapp.com/", SEMIR_URL, artifact_dir)
    await wait_for_semir_api_ready(semir_runner, log=log)
    tmall_runner: JSRunner | None = None

    settings = resolve_one_xm_settings()
    if args.generate:
        require_one_xm_key_for_generation(
            settings,
            model=getattr(args, "model", "gpt-image-2"),
            image_size=args.image_size,
            key_tier=args.one_xm_key_tier,
        )

    total = len(workflow_rows)
    reference_mode = compact(getattr(args, "reference_mode", "main_only")) or "main_only"
    requires_detail_reference = reference_mode == "main_and_detail"
    should_download_detail_reference = requires_detail_reference or bool(getattr(args, "confirm_generation", False))
    for completed, workflow in enumerate(workflow_rows, start=1):
        if wait_for_control:
            await wait_for_control({
                "records": len(all_rows),
                "shared": {
                    "total_rows": total,
                    "current_exec_no": completed,
                    "current_row_no": workflow.row_no,
                    "current_buyer_id": workflow.style_code,
                    "current_store": "森马云盘找图",
                    "search_total_codes": total,
                    "search_completed_codes": completed - 1,
                    "generation_total_jobs": 0,
                    "generation_submitted_jobs": 0,
                    "generation_completed_jobs": 0,
                },
                "phase": "tmall_ai_chain_semir",
            })
        log(f"[chain] 森马云盘找图 {completed}/{total} {workflow.style_code}")
        try:
            semir_data, main_origin_path, detail_reference_path = await find_semir_references(
                semir_runner,
                workflow,
                args.cloud_path,
                artifact_dir,
                allow_global_fallback=bool(getattr(args, "allow_global_semir_fallback", False)),
                download_detail=should_download_detail_reference,
            )
        except Exception as exc:
            workflow_results.append({"kind": "rows", "rows": [{
                "表格行号": workflow.row_no,
                "款号": workflow.style_code,
                "商品ID": workflow.item_id,
                "阶段": "森马云盘找图",
                "参考图模式": reference_mode,
                "品类": workflow.category,
                "性别": workflow.gender,
                "SKC编码": workflow.skc_code,
                "执行结果": "找图失败",
                "备注": str(exc),
            }]})
            if wait_for_control:
                await wait_for_control({
                    "records": len(all_rows),
                    "shared": {
                        "total_rows": total,
                        "current_exec_no": completed,
                        "current_row_no": workflow.row_no,
                        "current_buyer_id": workflow.style_code,
                        "current_store": "森马云盘找图失败",
                        "search_total_codes": total,
                        "search_completed_codes": completed,
                        "generation_total_jobs": 0,
                        "generation_submitted_jobs": 0,
                        "generation_completed_jobs": 0,
                    },
                    "phase": "tmall_ai_chain_semir",
                })
            continue

        detail_candidate = semir_data.get("chosenDetail") if isinstance(semir_data, dict) else {}
        if not workflow.skc_code and isinstance(detail_candidate, Mapping) and compact(detail_candidate.get("skcCode")):
            workflow.skc_code = compact(detail_candidate.get("skcCode"))

        selected_prompts = select_prompts(workflow, prompts, max(1, int(args.ai_image_count or 4)))
        if not selected_prompts:
            workflow_results.append({"kind": "rows", "rows": [{
                "表格行号": workflow.row_no,
                "款号": workflow.style_code,
                "商品ID": workflow.item_id,
                "阶段": "提示词匹配",
                "参考图模式": reference_mode,
                "品类": workflow.category,
                "性别": workflow.gender,
                "SKC编码": workflow.skc_code,
                "执行结果": "未匹配到提示词",
                "备注": f"映射分组={category_to_prompt_sheet(workflow.category)}",
            }]})
            if wait_for_control:
                await wait_for_control({
                    "records": len(all_rows),
                    "shared": {
                        "total_rows": total,
                        "current_exec_no": completed,
                        "current_row_no": workflow.row_no,
                        "current_buyer_id": workflow.style_code,
                        "current_store": "提示词未匹配",
                        "search_total_codes": total,
                        "search_completed_codes": completed,
                        "generation_total_jobs": 0,
                        "generation_submitted_jobs": 0,
                        "generation_completed_jobs": 0,
                    },
                    "phase": "tmall_ai_chain_semir",
                })
            continue

        if not main_origin_path or (requires_detail_reference and not detail_reference_path):
            missing_parts = []
            if not main_origin_path:
                missing_parts.append("橱窗1/yz(1)主图原图")
            if requires_detail_reference and not detail_reference_path:
                missing_parts.append("SKC编码平铺参考图")
            workflow_results.append({"kind": "rows", "rows": result_rows_for_workflow(
                workflow,
                semir_data,
                main_origin_path,
                detail_reference_path,
                [],
                [],
                None,
                error=f"缺少{'、'.join(missing_parts)}，已跳过 1XM 生图和天猫上传",
                reference_mode=reference_mode,
            )})
            if wait_for_control:
                await wait_for_control({
                    "records": len(all_rows),
                    "shared": {
                        "total_rows": total,
                        "current_exec_no": completed,
                        "current_row_no": workflow.row_no,
                        "current_buyer_id": workflow.style_code,
                        "current_store": "森马云盘找图完成",
                        "search_total_codes": total,
                        "search_completed_codes": completed,
                        "generation_total_jobs": 0,
                        "generation_submitted_jobs": 0,
                        "generation_completed_jobs": 0,
                    },
                    "phase": "tmall_ai_chain_semir",
                })
            continue
        reference_paths = [main_origin_path]
        if requires_detail_reference and not getattr(args, "confirm_generation", False) and detail_reference_path:
            reference_paths.append(detail_reference_path)
        generation_rows: list[dict[str, Any]] = [
            make_generation_row(
                workflow,
                prompt,
                reference_paths,
                image_size=args.image_size,
                quality=args.quality,
                output_format=args.output_format,
                key_tier=args.one_xm_key_tier,
                model=getattr(args, "model", "gpt-image-2"),
                ratio=getattr(args, "ratio", ""),
                prompt_index=prompt_index,
                run_nonce=args.task_run_uid,
            )
            for prompt_index, prompt in enumerate(selected_prompts, start=1)
        ]
        plan = {
            "workflow": workflow,
            "semir_data": semir_data,
            "main_origin_path": main_origin_path,
            "detail_reference_path": detail_reference_path,
            "selected_prompts": selected_prompts,
            "generation_rows": generation_rows,
            "generated_paths_by_index": [[] for _ in generation_rows],
            "reference_mode": reference_mode,
        }
        generation_plans.append(plan)
        workflow_results.append({"kind": "plan", "plan": plan})
        if wait_for_control:
            await wait_for_control({
                "records": len(all_rows),
                "shared": {
                    "total_rows": total,
                    "current_exec_no": completed,
                    "current_row_no": workflow.row_no,
                    "current_buyer_id": workflow.style_code,
                    "current_store": "森马云盘找图完成",
                    "search_total_codes": total,
                    "search_completed_codes": completed,
                    "generation_total_jobs": 0,
                    "generation_submitted_jobs": 0,
                    "generation_completed_jobs": 0,
                },
                "phase": "tmall_ai_chain_semir",
            })

    generation_jobs: list[tuple[int, int, dict[str, Any]]] = []
    for plan_index, plan in enumerate(generation_plans):
        for row_index, generation_row in enumerate(plan["generation_rows"]):
            generation_jobs.append((plan_index, row_index, generation_row))

    if getattr(args, "confirm_generation", False) and generation_plans:
        for entry in workflow_results:
            if entry.get("kind") == "rows":
                all_rows.extend(entry.get("rows") or [])
        confirmation_items = [
            build_generation_confirmation_item(
                plan["workflow"],
                plan["semir_data"],
                plan["main_origin_path"],
                plan["detail_reference_path"],
                plan["generation_rows"],
                reference_mode=plan.get("reference_mode") or reference_mode,
            )
            for plan in generation_plans
        ]
        batch = write_generation_confirmation_batch(
            artifact_dir,
            confirmation_items,
            run_params=vars(args),
            base_url=getattr(args, "approval_base_url", ""),
        )
        all_rows.extend(generation_confirmation_result_rows(batch))
        if invalid_rows:
            all_rows.extend(invalid_rows)
        return all_rows

    generation_total = len(generation_jobs)
    if args.generate and generation_jobs:
        concurrency = max(1, min(100, int(getattr(args, "generation_concurrency", 100) or 100)))
        log(f"[chain] 1XM 生图批量提交 {generation_total} 张，并发 {concurrency}")
        if wait_for_control:
            await wait_for_control({
                "records": len(all_rows),
                "shared": {
                    "total_rows": total,
                    "current_exec_no": total,
                    "search_total_codes": total,
                    "search_completed_codes": total,
                    "generation_total_jobs": generation_total,
                    "generation_submitted_jobs": 0,
                    "generation_completed_jobs": 0,
                    "current_store": f"1XM 生图批量提交，并发 {concurrency}",
                },
                "phase": "tmall_ai_chain_generate",
            })
        semaphore = asyncio.Semaphore(concurrency)
        submitted_count = 0
        completed_count = 0
        progress_lock = asyncio.Lock()

        async def generate_one(plan_index: int, row_index: int, generation_row: dict[str, Any]) -> tuple[int, int, dict[str, Any], list[str]]:
            nonlocal submitted_count, completed_count
            async with semaphore:
                plan = generation_plans[plan_index]
                workflow = plan["workflow"]
                prompt_no = int(generation_row.get("提示词序号") or row_index + 1)
                prompt_total = len(plan["generation_rows"])
                async with progress_lock:
                    submitted_count += 1
                    current_submitted = submitted_count
                    current_completed = completed_count
                log(f"[chain] {workflow.style_code} 1XM 生图提交 {prompt_no}/{prompt_total}（批量队列）")
                if wait_for_control:
                    await wait_for_control({
                        "records": len(all_rows),
                        "shared": {
                            "total_rows": total,
                            "current_exec_no": total,
                            "current_row_no": workflow.row_no,
                            "current_buyer_id": workflow.style_code,
                            "current_source_filename": compact(generation_row.get("提示词字段名")),
                            "current_store": f"1XM 生图提交 {current_submitted}/{generation_total}",
                            "search_total_codes": total,
                            "search_completed_codes": total,
                            "generation_total_jobs": generation_total,
                            "generation_submitted_jobs": current_submitted,
                            "generation_completed_jobs": current_completed,
                        },
                        "phase": "tmall_ai_chain_generate",
                    })
                try:
                    patched = await asyncio.to_thread(run_one_xm_generation_row, generation_row, {
                        "one_xm_key_tier": args.one_xm_key_tier,
                        "retry_attempts": args.retry_attempts,
                        "compensate_attempts": args.compensate_attempts,
                        "poll_timeout_minutes": args.poll_timeout_minutes,
                    }, settings)
                    row_generated_paths = download_generated_images(patched, artifact_dir)
                    patched["本地生成图文件"] = "\n".join(row_generated_paths)
                    return plan_index, row_index, patched, row_generated_paths
                except Exception as exc:
                    failed = dict(generation_row)
                    failed["执行结果"] = "生图失败"
                    failed["备注"] = str(exc)
                    return plan_index, row_index, failed, []

        tasks = [
            asyncio.create_task(generate_one(plan_index, row_index, row))
            for plan_index, row_index, row in generation_jobs
        ]
        for done_count, task in enumerate(asyncio.as_completed(tasks), start=1):
            plan_index, row_index, patched_row, row_paths = await task
            plan = generation_plans[plan_index]
            workflow = plan["workflow"]
            plan["generation_rows"][row_index] = patched_row
            plan["generated_paths_by_index"][row_index] = row_paths
            async with progress_lock:
                completed_count = done_count
                current_submitted = submitted_count
            if wait_for_control:
                await wait_for_control({
                    "records": len(all_rows),
                    "shared": {
                        "total_rows": total,
                        "current_exec_no": total,
                        "current_row_no": workflow.row_no,
                        "current_buyer_id": workflow.style_code,
                        "current_source_filename": compact(patched_row.get("提示词字段名")),
                        "current_store": f"1XM 生图完成 {done_count}/{generation_total}",
                        "search_total_codes": total,
                        "search_completed_codes": total,
                        "generation_total_jobs": generation_total,
                        "generation_submitted_jobs": current_submitted,
                        "generation_completed_jobs": done_count,
                    },
                    "phase": "tmall_ai_chain_generate",
                })
    elif not args.generate:
        for plan in generation_plans:
            for generation_row in plan["generation_rows"]:
                generation_row["执行结果"] = "已生成计划"
                generation_row["备注"] = "未启用 generate，未调用 1XM"

    if args.generate:
        for plan in generation_plans:
            sync_generation_plan_to_ai_image_workbench(
                plan,
                task_run_uid=compact(getattr(args, "task_run_uid", "")),
                model=compact(getattr(args, "model", "gpt-image-2")),
                image_size=compact(getattr(args, "image_size", "")),
                ratio=compact(getattr(args, "ratio", "")),
                output_format=compact(getattr(args, "output_format", "")),
                quality=compact(getattr(args, "quality", "")),
            )

    for entry in workflow_results:
        if entry.get("kind") == "rows":
            all_rows.extend(entry.get("rows") or [])
            continue

        plan = entry.get("plan") or {}
        workflow = plan["workflow"]
        semir_data = plan["semir_data"]
        main_origin_path = plan["main_origin_path"]
        detail_reference_path = plan["detail_reference_path"]
        generation_rows = plan["generation_rows"]
        generated_paths_by_index = plan["generated_paths_by_index"]
        generated_paths = [row_paths[0] for row_paths in generated_paths_by_index if row_paths]
        expected_ai_count = len(plan["selected_prompts"])
        plan_reference_mode = plan.get("reference_mode") or reference_mode

        if args.live_create and len(generated_paths) < expected_ai_count:
            all_rows.extend(result_rows_for_workflow(
                workflow,
                semir_data,
                main_origin_path,
                detail_reference_path,
                generation_rows,
                generated_paths,
                {"error": f"{workflow.style_code} 仅生成 {len(generated_paths)}/{expected_ai_count} 张 AI 图，未创建天猫测图任务"},
                reference_mode=plan_reference_mode,
            ))
            continue

        if getattr(args, "approval_required", False):
            approval_items.append(build_approval_item(
                workflow,
                semir_data,
                main_origin_path,
                detail_reference_path,
                generation_rows,
                generated_paths[:expected_ai_count],
                reference_mode=plan_reference_mode,
            ))
            all_rows.extend(result_rows_for_workflow(
                workflow,
                semir_data,
                main_origin_path,
                detail_reference_path,
                generation_rows,
                generated_paths[:expected_ai_count],
                None,
                reference_mode=plan_reference_mode,
            ))
            continue

        tmall_result = None
        if args.live_upload or args.live_create:
            if not workflow.item_id:
                tmall_result = {
                    "uploaded": [],
                    "materials": [],
                    "batchPayload": {},
                    "onlineResult": None,
                    "error": f"{workflow.style_code} 缺少商品 ID，无法创建天猫测图任务",
                }
            elif not tmall_runner:
                tmall_runner = await get_runner(
                    bridge,
                    "https://myseller.taobao.com/home.htm/material-center/material-test",
                    TMALL_URL,
                    artifact_dir,
                )
                await wait_for_tmall_api_ready(tmall_runner, log=log, timeout_ms=TMALL_LOGIN_READY_TIMEOUT_MS)
            if workflow.item_id and wait_for_control:
                await wait_for_control({
                    "records": len(all_rows),
                    "shared": {
                        "total_rows": total,
                        "current_exec_no": total,
                        "current_row_no": workflow.row_no,
                        "current_buyer_id": workflow.style_code,
                        "current_store": "天猫上传/创建/上线",
                        "search_total_codes": total,
                        "search_completed_codes": total,
                        "generation_total_jobs": generation_total,
                        "generation_submitted_jobs": generation_total,
                        "generation_completed_jobs": generation_total,
                    },
                    "phase": "tmall_ai_chain_tmall",
                })
            if workflow.item_id:
                log(f"[chain] {workflow.style_code} 天猫上传/创建测图任务")
                try:
                    tmall_result = await upload_and_create_tmall_task(
                        tmall_runner,
                        workflow,
                        main_origin_path,
                        generated_paths[:expected_ai_count],
                        live_upload=args.live_upload,
                        live_create=args.live_create,
                        live_online=args.live_online,
                        upload_delay_seconds=args.tmall_upload_delay_seconds,
                        create_delay_seconds=args.tmall_create_delay_seconds,
                        batch_delay_seconds=args.tmall_batch_delay_seconds,
                        readback_delay_seconds=args.tmall_readback_delay_seconds,
                    )
                except Exception as exc:
                    tmall_result = {
                        "uploaded": [],
                        "materials": [],
                        "batchPayload": {},
                        "onlineResult": None,
                        "error": str(exc),
                    }
        elif generated_paths or main_origin_path:
            tmall_result = {
                "uploaded": [],
                "materials": [],
                "batchPayload": {"experimentTaskId": "<experimentTaskId>"},
            }

        all_rows.extend(result_rows_for_workflow(
            workflow,
            semir_data,
            main_origin_path,
            detail_reference_path,
            generation_rows,
            generated_paths[:expected_ai_count],
            tmall_result,
            reference_mode=plan_reference_mode,
        ))

    if invalid_rows:
        all_rows.extend(invalid_rows)

    if approval_items:
        batch = write_approval_batch(
            artifact_dir,
            approval_items,
            run_params=vars(args),
            base_url=getattr(args, "approval_base_url", ""),
        )
        maybe_sync_approval_batch_to_cloud(batch, log=log)
        notify_channel = compact(getattr(args, "approval_notify_channel", "dingtalk")).lower()
        if notify_channel and notify_channel != "none":
            notify_approval_batch(batch, channel=notify_channel, log=log)
        all_rows.extend(approval_result_rows(batch))
        if (batch.get("cloud_sync") or {}).get("status") == "warning":
            all_rows.append(cloud_approval_sync_warning_row(batch, (batch.get("cloud_sync") or {}).get("warning", "")))

    return all_rows


async def run_chain(args: argparse.Namespace) -> dict[str, Any]:
    data_sink.init_db()
    run_id = data_sink.begin_run(ADAPTER_ID, TASK_ID)
    artifact_dir = Path(data_sink.prepare_artifact_dir(ADAPTER_ID, TASK_ID, run_id, "artifacts"))
    output_files: list[str] = []
    all_rows: list[dict[str, Any]] = []

    try:
        all_rows = await run_chain_rows(args, artifact_dir)

        json_path = data_sink.export_json(
            all_rows,
            ADAPTER_ID,
            TASK_ID,
            "天猫AI测图全链路执行证据_{timestamp}.json",
        )
        excel_path = data_sink.export_excel(
            all_rows,
            ADAPTER_ID,
            TASK_ID,
            "天猫AI测图全链路执行证据_{timestamp}.xlsx",
            column_order=TMALL_AI_IMAGE_RESULT_COLUMN_ORDER,
        )
        output_files.extend([excel_path, json_path])
        data_sink.finish_run(run_id, len(all_rows), output_files)
        return {
            "run_id": run_id,
            "artifact_dir": str(artifact_dir),
            "rows": len(all_rows),
            "output_files": output_files,
        }
    except Exception as exc:
        if all_rows:
            try:
                output_files.append(data_sink.export_json(
                    all_rows,
                    ADAPTER_ID,
                    TASK_ID,
                    "天猫AI测图全链路执行证据_失败_{timestamp}.json",
                ))
            except Exception:
                pass
        data_sink.fail_run(run_id, str(exc))
        raise


def main() -> None:
    args = build_parser().parse_args()
    started = time.monotonic()
    result = asyncio.run(run_chain(args))
    elapsed = time.monotonic() - started
    print(json.dumps({**result, "elapsed_seconds": round(elapsed, 2)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
